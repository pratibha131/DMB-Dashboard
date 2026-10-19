import os
import json
import time
from pathlib import Path
from datetime import date, datetime
from io import BytesIO
import re
import threading
from zipfile import BadZipFile

import dash
from dash import ALL, Dash, Input, Output, State, ctx, dcc, html, no_update
from dash.exceptions import PreventUpdate
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import pandas as pd
import plotly.graph_objects as go

from data_loader import load_dashboard_data, load_strategic_rca_actions, clean_text, is_value_lever
from kpi_calculations import (
    get_default_reporting_month,
    get_month_summary,
    prepare_kpi_data,
)


def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).replace("\xa0", " ").strip()


# =========================================================
# SETTINGS AND PATHS
# =========================================================

GAUGE_TARGET = 85
INSIGHT_DISPLAY_LIMIT = 3

BASE_DIR = Path(__file__).resolve().parent

# This also supports running the downloaded app.py before it is copied
# into the main project folder.
if not (BASE_DIR / "data").exists() and (BASE_DIR.parent / "data").exists():
    BASE_DIR = BASE_DIR.parent

ASSETS_DIR = BASE_DIR / "assets"
CSS_FILE = ASSETS_DIR / "style.css"


def find_masterfile():
    preferred_file = BASE_DIR / "data" / "Masterfile_DMB_Dashboard.xlsx"

    candidates = sorted(
        file_path
        for file_path in (BASE_DIR / "data").glob("*.xlsx")
        if not file_path.name.startswith("~$")
    )

    if preferred_file.exists():
        candidates = [preferred_file] + [
            file_path
            for file_path in candidates
            if file_path != preferred_file
        ]

    fallback_file = None

    # Identify the masterfile by its required source sheets, so an additional
    # workbook such as Book1.xlsx is never selected by mistake.
    for file_path in candidates:
        try:
            sheet_names = set(pd.ExcelFile(
                file_path,
                engine="openpyxl",
            ).sheet_names)

            if {
                "MPR Masterfile",
                "DMB Masterfile",
            }.issubset(sheet_names):
                if "RCA Actions" in sheet_names:
                    return file_path

                if fallback_file is None:
                    fallback_file = file_path
        except (OSError, ValueError):
            continue

    return fallback_file or preferred_file


MASTERFILE_PATH = find_masterfile()


# =========================================================
# LOAD AND PREPARE DATA (DYNAMIC DATA STORE)
# =========================================================

import sharepoint_sync
import data_loader

_data_store = data_loader.get_data_store()
_data_store.reload()


def get_active_mpr_data():
    df = _data_store.get_mpr_data()
    if not df.empty and "kpi_name" in df.columns:
        df = df[~df["kpi_name"].map(is_value_lever)].copy()
        if "strategic_imperative" in df.columns:
            df.loc[
                df["kpi_name"].str.strip().str.lower().str.startswith("win rate"),
                "strategic_imperative",
            ] = "Drive growth through Commercial Excellence"
        return df
    return df


def get_active_dmb_data():
    return _data_store.get_dmb_data()


def get_active_strat_rca():
    df = _data_store.get_strat_rca()
    if not df.empty and "kpi_name" in df.columns:
        df = df[~df["kpi_name"].map(is_value_lever)].copy()
        if "strategic_imperative" in df.columns:
            df.loc[
                df["kpi_name"].str.strip().str.lower().str.startswith("win rate"),
                "strategic_imperative",
            ] = "Drive growth through Commercial Excellence"
        return df
    return df


def load_rca_actions():
    strat_actions = load_strategic_rca_actions()
    if not strat_actions.empty:
        if "kpi_name" in strat_actions.columns:
            strat_actions = strat_actions[~strat_actions["kpi_name"].map(is_value_lever)].copy()
        return strat_actions

    required_columns = [
        "reporting_month",
        "strategic_imperative",
        "kpi_name",
        "cause",
        "action",
    ]

    masterfile = find_masterfile()
    if not masterfile or not masterfile.exists():
        return pd.DataFrame(columns=required_columns)

    try:
        actions = pd.read_excel(
            masterfile,
            sheet_name="RCA Actions",
            engine="openpyxl",
        )
    except Exception:
        return pd.DataFrame(columns=required_columns)

    column_lookup = {
        str(column).strip().lower().replace(" ", "_"): column
        for column in actions.columns
    }

    rename_map = {}
    for required_column in required_columns:
        source_column = column_lookup.get(required_column)
        if source_column is not None:
            rename_map[source_column] = required_column

    actions = actions.rename(columns=rename_map)

    for required_column in required_columns:
        if required_column not in actions.columns:
            actions[required_column] = ""

    actions = actions[required_columns].copy()
    actions["reporting_month"] = (
        pd.to_datetime(actions["reporting_month"], errors="coerce")
        .dt.to_period("M")
        .dt.to_timestamp()
    )

    return actions


def get_active_rca_actions():
    return load_rca_actions()


mpr_data = get_active_mpr_data()
dmb_data = get_active_dmb_data()
rca_actions = get_active_rca_actions()


# =========================================================
# DETAILED FUNCTION RCA DATA
# =========================================================

ACTION_STATUS_LABELS = {
    1: "Action not assigned",
    2: "Action assigned",
    3: "Action started",
    4: "Action in progress",
    5: "Action completed",
}


def clean_cell_text(value):
    if value is None or pd.isna(value):
        return ""

    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def function_key(value):
    return re.sub(r"[^a-z0-9]+", " ", clean_cell_text(value).lower()).strip()


FUNCTION_SHEET_TO_DASHBOARD = {
    "quality": "Quality",
    "regulatory": "Regulatory",
    "isc": "ISC & Procurement",
    "procurement": "ISC & Procurement",
    "isc procurement": "ISC & Procurement",
    "r d": "R&D",
    "marketing": "Marketing",
    "customer service": "Customer Service",
    "commercial excellence": "Commercial Excellence",
    "nar": "NAR",
    "europe": "Europe",
    "growth": "Growth",
    "finance": "Finance",
}

FUNCTION_ORDER = [
    "Quality",
    "Customer Service",
    "Marketing",
    "ISC & Procurement",
    "Regulatory",
    "R&D",
]

EXCLUDED_DMB_FUNCTIONS = {
    "nar",
    "europe",
    "growth",
    "finance",
    "commercial excellence",
    "isc",
    "procurement",
}



def kpi_key(value):
    normalized = function_key(value)

    if "profitability" in normalized or re.search(r"\bigm\b", normalized):
        return "igm"
    if "first vist fix" in normalized or "first visit fix" in normalized:
        return "fvf"
    if re.search(r"\bfvf\b", normalized):
        return "fvf"
    if "contract penetration" in normalized or normalized in {
        "cp",
        "cp percent",
    }:
        return "cp"
    if "production" in normalized:
        return "production"
    if "cdd" in normalized:
        return "cdd"
    if "ctb" in normalized:
        return "ctb"
    if "fsd" in normalized:
        return "fsd"
    if "market inventory" in normalized:
        return "market inventory"
    elif "inventory" in normalized:
        return "inventory"
    # Only plain sales KPIs share the key; "Sales tool delivery" and
    # "Sales Training" are different KPIs.
    if normalized in {"sales", "cs sales"}:
        return "sales"

    return normalized


def sheet_function_name(sheet_name):
    return re.sub(
        r"^\s*\d+\s*[.)_-]?\s*",
        "",
        clean_cell_text(sheet_name),
    )


def dashboard_function_name(sheet_name):
    source_function = sheet_function_name(sheet_name)
    return FUNCTION_SHEET_TO_DASHBOARD.get(
        function_key(source_function),
        source_function,
    )


def read_workbook_bytes(file_path):
    """
    Read a workbook even while it is open in Excel.

    Excel keeps an open workbook's handle with delete access, so Windows
    refuses a normal open() (which does not share delete access) with
    "Permission denied". Opening with full sharing reads the last saved copy.
    """
    if os.name != "nt":
        return Path(file_path).read_bytes()

    import ctypes
    import msvcrt
    from ctypes import wintypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    create_file.restype = wintypes.HANDLE

    generic_read = 0x80000000
    share_read_write_delete = 0x7
    open_existing = 3
    normal_attributes = 0x80

    handle = create_file(
        str(file_path),
        generic_read,
        share_read_write_delete,
        None,
        open_existing,
        normal_attributes,
        None,
    )
    if handle is None or handle == wintypes.HANDLE(-1).value:
        error_code = ctypes.get_last_error()
        raise ctypes.WinError(error_code)

    file_descriptor = msvcrt.open_osfhandle(handle, os.O_RDONLY)
    with open(file_descriptor, "rb") as workbook_file:
        return workbook_file.read()


def get_visible_rca_sheets(workbook_bytes):
    try:
        workbook = load_workbook(
            BytesIO(workbook_bytes),
            read_only=True,
            data_only=True,
        )
    except (OSError, ValueError, BadZipFile, InvalidFileException):
        return []

    matching_sheets = []

    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue

            has_rca_section = any(
                clean_cell_text(row[0]).lower().startswith(
                    "root cause analysis"
                )
                for row in worksheet.iter_rows(
                    min_col=1,
                    max_col=1,
                    values_only=True,
                )
            )

            if has_rca_section:
                matching_sheets.append(worksheet.title)
    finally:
        workbook.close()

    return matching_sheets


def find_detail_workbook():
    """
    Return (workbook path, its visible RCA sheets, its contents, workbooks
    that could not be read because they were locked).
    """
    data_directory = BASE_DIR / "data"

    preferred_candidates = [
        data_directory / "Functional DMB Review Sheets-17th_sept.xlsx",
    ]

    candidates = sorted(
        file_path
        for file_path in data_directory.glob("*.xlsx")
        if not file_path.name.startswith("~$")
        and file_path != MASTERFILE_PATH
    )

    for pref in reversed(preferred_candidates):
        if pref.exists() and pref in candidates:
            candidates.remove(pref)
            candidates.insert(0, pref)

    # Select the workbook with the greatest number of visible RCA sheets.
    # This makes the multi-function review workbook take precedence over an
    # older single-sheet Book1.xlsx if both files remain in the data folder.
    ranked_candidates = []
    locked_files = []
    for file_path in candidates:
        try:
            workbook_bytes = read_workbook_bytes(file_path)
        except PermissionError:
            locked_files.append(file_path)
            continue
        except OSError:
            continue

        rca_sheets = get_visible_rca_sheets(workbook_bytes)
        if rca_sheets:
            # Prefer the current review workbook over older copies in data/.
            is_pref = 1 if file_path in preferred_candidates else 0
            ranked_candidates.append(
                (
                    is_pref,
                    len(rca_sheets),
                    file_path.stat().st_mtime,
                    file_path,
                    rca_sheets,
                    workbook_bytes,
                )
            )

    if ranked_candidates:
        _, _, _, detail_file, rca_sheets, workbook_bytes = max(
            ranked_candidates,
            key=lambda candidate: candidate[:3],
        )
        return detail_file, rca_sheets, workbook_bytes, locked_files

    return None, [], None, locked_files


def locked_workbook_message(file_paths):
    names = ", ".join(file_path.name for file_path in file_paths)
    return (
        f"{names} could not be read because another program is locking it "
        "(for example OneDrive sync). Wait a moment, then click the card "
        "again."
    )


def format_due_date(value):
    if value is None or pd.isna(value):
        return "Not entered"

    text_value = clean_cell_text(value)
    if not text_value:
        return "Not entered"

    # The Excel action trackers intentionally contain month-only deadlines
    # such as "Aug", "Sept" and "NOV".  Parsing those as complete dates can
    # create a fictitious year 0001, so keep them as month labels.
    month_match = re.fullmatch(
        r"(?i)(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|"
        r"jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|"
        r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)",
        text_value,
    )
    if month_match:
        return text_value

    # Preserve workflow labels and multi-date notes exactly as entered.
    if (
        "|" in text_value
        or text_value.lower()
        in {"ongoing", "continuous", "tbd", "na", "n/a"}
    ):
        return text_value

    entered_year = re.search(r"\b(\d{4})\b", text_value)
    if entered_year and int(entered_year.group(1)) <= 1901:
        return text_value

    if isinstance(value, (pd.Timestamp, datetime, date)):
        timestamp = pd.Timestamp(value)
        if timestamp.year <= 1901:
            return timestamp.strftime("%b")
        return timestamp.strftime("%d %b %Y")

    try:
        parsed_date = pd.to_datetime(
            value,
            errors="coerce",
            format="mixed",
            dayfirst=True,
        )
    except TypeError:
        # Compatibility fallback for older pandas releases.
        parsed_date = pd.to_datetime(
            text_value,
            errors="coerce",
            dayfirst=True,
        )
    if not pd.isna(parsed_date):
        if parsed_date.year <= 1901:
            return text_value
        return parsed_date.strftime("%d %b %Y")

    return text_value


UNICODE_STATUS_MAP = {
    "⚪": "Action not assigned",
    "○": "Action not assigned",
    "◔": "Action assigned",
    "◑": "Action started",
    "◐": "Action started",
    "◕": "Action in progress",
    "●": "Action completed",
    "✔": "Action completed",
    "✓": "Action completed",
}


def format_action_status(value):
    if value is None or pd.isna(value):
        return "Action not assigned"

    numeric_value = pd.to_numeric(value, errors="coerce")

    if not pd.isna(numeric_value):
        return ACTION_STATUS_LABELS.get(
            int(numeric_value),
            f"Status {int(numeric_value)}",
        )

    text = clean_cell_text(value)
    if text in UNICODE_STATUS_MAP:
        return UNICODE_STATUS_MAP[text]

    t_lower = text.lower()
    if "not assigned" in t_lower or "unassigned" in t_lower or "⚪" in text or "○" in text:
        return "Action not assigned"
    if "assigned" in t_lower or "quarter" in t_lower or "◔" in text:
        return "Action assigned"
    if "in progress" in t_lower or "inprogress" in t_lower or "3/4" in t_lower or "three quarter" in t_lower or "◕" in text:
        return "Action in progress"
    if "started" in t_lower or "half" in t_lower or "1/2" in t_lower or "ongoing" in t_lower or "◑" in text or "◐" in text:
        return "Action started"
    if "completed" in t_lower or "closed" in t_lower or "confirmed" in t_lower or "resolution" in t_lower or "full" in t_lower or "●" in text:
        return "Action completed"

    return text or "Action not assigned"


def action_status_class(status):
    s = str(status).strip().lower()
    if "not assigned" in s or "unassigned" in s or s in {"1", "1.0", "⚪", "○"}:
        return "action-status-not-assigned"
    if "assigned" in s or s in {"2", "2.0", "◔"}:
        return "action-status-assigned"
    if "in progress" in s or "inprogress" in s or s in {"4", "4.0", "◕"} or "three quarter" in s:
        return "action-status-in-progress"
    if "started" in s or "ongoing" in s or s in {"3", "3.0", "◑", "◐"} or "half" in s:
        return "action-status-started"
    if "completed" in s or "closed" in s or "confirmed" in s or "resolution" in s or s in {"5", "5.0", "●"}:
        return "action-status-completed"
    return "action-status-unknown"


KPI_NAME_STOPWORDS = {
    "a",
    "an",
    "and",
    "at",
    "by",
    "for",
    "in",
    "is",
    "of",
    "on",
    "per",
    "the",
    "to",
    "vs",
    "with",
}


def kpi_name_tokens(value):
    return {
        token
        for token in function_key(value).split()
        if token not in KPI_NAME_STOPWORDS
    }


def kpi_names_match(first_name, second_name):
    # The KPI data, root-cause tables and action trackers name the same KPI
    # differently (for example "OOH %", "OOH%" and "OOH% to revunue in Qtr.").
    # Names match when their keys agree or one name's words contain the other's.
    first_key = kpi_key(first_name)
    if first_key and first_key == kpi_key(second_name):
        return True

    first_tokens = kpi_name_tokens(first_name)
    second_tokens = kpi_name_tokens(second_name)
    if not first_tokens or not second_tokens:
        return False

    return first_tokens <= second_tokens or second_tokens <= first_tokens


def kpi_match_score(first_name, second_name):
    """2 = same KPI key, 1 = one name's words contain the other's, 0 = no match."""
    first_key = kpi_key(first_name)
    if first_key and first_key == kpi_key(second_name):
        return 2
    if kpi_names_match(first_name, second_name):
        return 1
    return 0


def assign_to_best_match(item_names, target_aliases):
    """
    Map each item index to the target indexes it matches best, so "Inventory"
    goes to the "Inventory" KPI rather than also to "Market Inventory".
    """
    assignments = {}
    for item_index, item_name in enumerate(item_names):
        scores = [
            max(
                (kpi_match_score(item_name, alias) for alias in aliases),
                default=0,
            )
            for aliases in target_aliases
        ]
        best_score = max(scores, default=0)
        if best_score:
            assignments[item_index] = [
                target_index
                for target_index, score in enumerate(scores)
                if score == best_score
            ]
    return assignments


MONTH_TAG_PATTERN = re.compile(
    r"\b(jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|"
    r"jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|"
    r"dec(?:ember)?)[\s'’-]*(2\d|20\d{2})\b",
    re.IGNORECASE,
)


def name_applies_to_month(kpi_name, month):
    """
    RCA entries such as "Product Change assessment (Apr-26 & May-26)" belong
    only to the months they name; entries without a month apply to any month.
    """
    tags = MONTH_TAG_PATTERN.findall(clean_cell_text(kpi_name))
    if not tags:
        return True

    month = pd.Timestamp(month)
    for month_text, year_text in tags:
        year = int(year_text) if len(year_text) == 4 else 2000 + int(year_text)
        tagged_month = pd.to_datetime(
            f"{month_text[:3]} {year}", format="%b %Y", errors="coerce"
        )
        if (
            pd.notna(tagged_month)
            and tagged_month.year == month.year
            and tagged_month.month == month.month
        ):
            return True
    return False


def kpi_sheet_qualifier(kpi_name, source_sheets):
    """Return the sheet named in a KPI such as "CTB (12 weeks)(Procurement)"."""
    for qualifier in re.findall(r"\(([^()]*)\)", clean_cell_text(kpi_name)):
        qualifier_key = function_key(qualifier)
        if not qualifier_key:
            continue
        for source_sheet in source_sheets:
            if qualifier_key == function_key(sheet_function_name(source_sheet)):
                return source_sheet
    return None


def root_cause_matches(root_cause, causes):
    """True when an action's root cause repeats one of the given RCA causes."""
    root_key = function_key(root_cause)
    if len(root_key) < 12:
        return False
    for cause in causes:
        cause_key = function_key(cause)
        if len(cause_key) >= 12 and (cause_key in root_key or root_key in cause_key):
            return True
    return False


def parse_review_month(value):
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(value).to_period("M").to_timestamp()

    parsed_month = pd.to_datetime(
        clean_cell_text(value),
        format="%b-%Y",
        errors="coerce",
    )
    if pd.isna(parsed_month):
        return None

    return parsed_month


def parse_review_kpi_rows(raw_sheet, sheet_name):
    """Read the monthly Target/Actual KPI block at the top of a review sheet."""
    records = []
    month_columns = {}
    current_kpi = None

    if raw_sheet.shape[1] < 9:
        return records

    for row_index in range(len(raw_sheet)):
        first_cell = clean_cell_text(raw_sheet.iat[row_index, 0]).lower()
        if first_cell.startswith(("trends", "root cause analysis")):
            break

        row_type = clean_cell_text(raw_sheet.iat[row_index, 7]).lower()

        if row_type.startswith("target/"):
            month_columns = {}
            for column_index in range(8, min(20, raw_sheet.shape[1])):
                month = parse_review_month(
                    raw_sheet.iat[row_index, column_index]
                )
                if month is not None:
                    month_columns[column_index] = month
            continue

        if row_type == "target":
            # Sub-KPIs such as "Complaint rate for Z30" leave the KPI name
            # blank and carry their name in the definition column.
            kpi_name = clean_cell_text(
                raw_sheet.iat[row_index, 0]
            ) or clean_cell_text(raw_sheet.iat[row_index, 1])

            current_kpi = {
                "kpi_row": row_index,
                "kpi_name": kpi_name,
                "metric_nature": clean_cell_text(
                    raw_sheet.iat[row_index, 5]
                ),
                "targets": {
                    column_index: raw_sheet.iat[row_index, column_index]
                    for column_index in month_columns
                },
            }
            continue

        if row_type != "actual" or current_kpi is None:
            continue

        if current_kpi["kpi_name"]:
            for column_index, month in month_columns.items():
                records.append(
                    {
                        "source_sheet": sheet_name,
                        "kpi_row": current_kpi["kpi_row"],
                        "kpi_name": current_kpi["kpi_name"],
                        "metric_nature": current_kpi["metric_nature"],
                        "month": month,
                        "Target": current_kpi["targets"].get(column_index),
                        "Actual": raw_sheet.iat[row_index, column_index],
                    }
                )

        current_kpi = None

    return records


def parse_impact_percent_value(val):
    """
    Parses raw impact value from Excel which may be:
    - float/decimal: 0.554 -> 55.4, 0.4 -> 40.0
    - string percentage: '55.4%' -> 55.4
    - integer: 1 -> 100.0, 40 -> 40.0
    """
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (int, float)):
        num = float(val)
        if abs(num) <= 1.0 and num != 0:
            num *= 100
        return round(num, 2)
    s = str(val).strip().replace("\xa0", "")
    if not s or s in {"—", "-", "none", "nan", "null"}:
        return None
    has_percent = "%" in s
    s_clean = s.replace("%", "").strip()
    try:
        num = float(s_clean)
        if not has_percent and abs(num) <= 1.0 and num != 0:
            num *= 100
        return round(num, 2)
    except (ValueError, TypeError):
        return None


def load_function_rca_details():
    cause_columns = [
        "function",
        "function_key",
        "source_function",
        "source_sheet",
        "kpi_name",
        "kpi_key",
        "cause_rank",
        "cause",
        "impact_percent",
    ]
    action_columns = [
        "function",
        "function_key",
        "source_function",
        "source_sheet",
        "kpi_name",
        "kpi_key",
        "root_cause",
        "corrective_action",
        "owner",
        "due_date",
        "status",
        "status_class",
    ]
    review_kpi_columns = [
        "function",
        "function_key",
        "source_function",
        "source_sheet",
        "kpi_row",
        "kpi_name",
        "metric_nature",
        "month",
        "Target",
        "Actual",
        "status",
    ]

    (
        detail_file,
        visible_rca_sheets,
        workbook_bytes,
        locked_files,
    ) = find_detail_workbook()

    def failed_result(error):
        print(f"[RCA Loader Notice] {error}")
        return {
            "causes": pd.DataFrame(columns=cause_columns),
            "actions": pd.DataFrame(columns=action_columns),
            "review_kpis": pd.DataFrame(columns=review_kpi_columns),
            "file": detail_file,
            "error": error,
        }

    if detail_file is None:
        if locked_files:
            return failed_result(locked_workbook_message(locked_files))
        return failed_result(
            "No workbook with Root Cause Analysis sheets was found in "
            "the data folder."
        )

    cause_records = []
    action_records = []
    review_kpi_records = []

    try:
        worksheets = pd.read_excel(
            BytesIO(workbook_bytes),
            sheet_name=visible_rca_sheets,
            header=None,
            engine="openpyxl",
        )
    except (OSError, ValueError) as error:
        return failed_result(f"{detail_file.name} could not be read: {error}")

    try:
        wb_openpyxl = load_workbook(BytesIO(workbook_bytes), data_only=False, rich_text=True)
    except Exception:
        wb_openpyxl = None

    def get_clean_cell_text_without_strike(s_name, r_idx, c_idx, fallback=""):
        if wb_openpyxl and s_name in wb_openpyxl.sheetnames:
            try:
                ws_ox = wb_openpyxl[s_name]
                c_ox = ws_ox.cell(row=r_idx + 1, column=c_idx + 1)
                val = c_ox.value
                if val is None:
                    return ""
                if hasattr(val, "__iter__") and not isinstance(val, (str, bytes)):
                    valid_parts = []
                    for run in val:
                        f = getattr(run, "font", None)
                        if f and f.strike:
                            continue
                        t = run.text if hasattr(run, "text") else str(run)
                        valid_parts.append(t)
                    return clean_cell_text("".join(valid_parts))
                else:
                    if c_ox.font and c_ox.font.strike:
                        return ""
                    return clean_cell_text(val)
            except Exception:
                pass
        return clean_cell_text(fallback)

    def is_cell_strike(s_name, r_idx, c_idx):
        if wb_openpyxl and s_name in wb_openpyxl.sheetnames:
            try:
                ws_ox = wb_openpyxl[s_name]
                c_ox = ws_ox.cell(row=r_idx + 1, column=c_idx + 1)
                val = c_ox.value
                if hasattr(val, "__iter__") and not isinstance(val, (str, bytes)):
                    has_non_strike = False
                    for run in val:
                        t = (run.text if hasattr(run, "text") else str(run)).strip()
                        if not t:
                            continue
                        f = getattr(run, "font", None)
                        if not (f and f.strike):
                            has_non_strike = True
                            break
                    return not has_non_strike
                return bool(c_ox.font and c_ox.font.strike)
            except Exception:
                pass
        return False

    for sheet_name, raw_sheet in worksheets.items():
        source_function = sheet_function_name(sheet_name)
        function_name = dashboard_function_name(sheet_name)
        current_function_key = function_key(function_name)

        if raw_sheet.empty or raw_sheet.shape[1] < 6:
            continue

        for review_kpi in parse_review_kpi_rows(raw_sheet, sheet_name):
            review_kpi_records.append(
                {
                    "function": function_name,
                    "function_key": current_function_key,
                    "source_function": source_function,
                    **review_kpi,
                }
            )

        first_column = raw_sheet.iloc[:, 0].map(clean_cell_text)

        root_headers = [
            row_index
            for row_index, cell_value in first_column.items()
            if cell_value.lower() == "red kpi"
        ]

        action_header_index = None
        for row_index in root_headers:
            second_cell = (
                clean_cell_text(raw_sheet.iat[row_index, 1])
                if raw_sheet.shape[1] > 1
                else ""
            )
            third_cell = (
                clean_cell_text(raw_sheet.iat[row_index, 2])
                if raw_sheet.shape[1] > 2
                else ""
            )

            if (
                second_cell.lower() == "root cause description"
                and third_cell.lower() == "corrective action"
            ):
                action_header_index = row_index
                continue

            data_row_index = None
            for possible_row in range(
                row_index + 1,
                min(row_index + 4, len(raw_sheet)),
            ):
                first_cell = clean_cell_text(raw_sheet.iat[possible_row, 0])
                if (
                    first_cell
                    and first_cell.lower() != "red kpi"
                    and "% impact" not in first_cell.lower()
                ):
                    data_row_index = possible_row
                    break

            if data_row_index is None:
                continue

            source_kpi_name = clean_cell_text(
                raw_sheet.iat[data_row_index, 0]
            )
            if not source_kpi_name:
                continue

            impact_row_index = None
            for possible_row in range(
                data_row_index + 1,
                min(data_row_index + 3, len(raw_sheet)),
            ):
                first_cell = clean_cell_text(raw_sheet.iat[possible_row, 0])
                if "% impact" in first_cell.lower():
                    impact_row_index = possible_row
                    break

            cause_rank = 0
            for column_index in range(1, min(6, raw_sheet.shape[1])):
                if is_cell_strike(sheet_name, data_row_index, column_index):
                    continue  # Ignore cut / struck-through causes in Excel

                cause_text = get_clean_cell_text_without_strike(
                    sheet_name, data_row_index, column_index, raw_sheet.iat[data_row_index, column_index]
                )
                if not cause_text:
                    continue

                cause_rank += 1
                impact_percent = None
                if impact_row_index is not None:
                    raw_impact = raw_sheet.iat[impact_row_index, column_index]
                    impact_percent = parse_impact_percent_value(raw_impact)

                cause_records.append(
                    {
                        "function": function_name,
                        "function_key": current_function_key,
                        "source_function": source_function,
                        "source_sheet": sheet_name,
                        "kpi_name": source_kpi_name,
                        "kpi_key": kpi_key(source_kpi_name),
                        "cause_rank": cause_rank,
                        "cause": cause_text,
                        "impact_percent": impact_percent,
                    }
                )

        if action_header_index is None:
            continue

        for row_index in range(action_header_index + 1, len(raw_sheet)):
            if is_cell_strike(sheet_name, row_index, 1) or is_cell_strike(sheet_name, row_index, 2):
                continue  # Ignore cut / struck-through action rows in Excel
            source_kpi_name = clean_cell_text(raw_sheet.iat[row_index, 0])

            if not source_kpi_name:
                continue
            if source_kpi_name.lower() == "action not assigned":
                break

            root_cause = get_clean_cell_text_without_strike(
                sheet_name, row_index, 1, raw_sheet.iat[row_index, 1]
            )
            corrective_action = get_clean_cell_text_without_strike(
                sheet_name, row_index, 2, raw_sheet.iat[row_index, 2]
            )

            if not root_cause and not corrective_action:
                continue

            status = format_action_status(raw_sheet.iat[row_index, 5])
            owner_text = get_clean_cell_text_without_strike(
                sheet_name, row_index, 3, raw_sheet.iat[row_index, 3]
            )
            due_date_raw = get_clean_cell_text_without_strike(
                sheet_name, row_index, 4, raw_sheet.iat[row_index, 4]
            )

            action_records.append(
                {
                    "function": function_name,
                    "function_key": current_function_key,
                    "source_function": source_function,
                    "source_sheet": sheet_name,
                    "kpi_name": source_kpi_name,
                    "kpi_key": kpi_key(source_kpi_name),
                    "root_cause": root_cause or "Not entered",
                    "corrective_action": (
                        corrective_action or "Not entered"
                    ),
                    "owner": owner_text or "Not assigned",
                    "due_date": format_due_date(due_date_raw or raw_sheet.iat[row_index, 4]),
                    "status": status,
                    "status_class": action_status_class(status),
                }
            )

    causes = pd.DataFrame(cause_records, columns=cause_columns)
    actions = pd.DataFrame(action_records, columns=action_columns)

    if not causes.empty:
        causes = causes.drop_duplicates(
            subset=[
                "function_key",
                "source_sheet",
                "kpi_key",
                "cause_rank",
                "cause",
                "impact_percent",
            ]
        ).reset_index(drop=True)

    if not actions.empty:
        actions = actions.drop_duplicates().reset_index(drop=True)

    review_kpis = pd.DataFrame(review_kpi_records)
    if review_kpis.empty:
        review_kpis = pd.DataFrame(columns=review_kpi_columns)
    else:
        review_kpis = prepare_kpi_data(
            review_kpis,
            ["source_sheet", "kpi_row"],
        )[review_kpi_columns]

    return {
        "causes": causes,
        "actions": actions,
        "review_kpis": review_kpis,
        "file": detail_file,
        "error": None,
    }


_function_rca_lock = threading.Lock()
_function_rca_cache = {"signature": None, "data": None}


def data_folder_signature():
    signature = []
    for file_path in sorted((BASE_DIR / "data").glob("*.xlsx")):
        if file_path.name.startswith("~$"):
            continue
        try:
            file_stat = file_path.stat()
        except OSError:
            continue
        signature.append(
            (file_path.name, file_stat.st_mtime_ns, file_stat.st_size)
        )
    return tuple(signature)


def get_function_rca_data():
    """
    Return the Functional DMB Review data, re-reading the workbook when a data
    file changes or when the previous read failed (for example because the
    workbook was open in Excel), so no app restart is needed.
    """
    signature = data_folder_signature()

    with _function_rca_lock:
        cached_data = _function_rca_cache["data"]
        if (
            cached_data is None
            or cached_data["error"]
            or _function_rca_cache["signature"] != signature
        ):
            _function_rca_cache["data"] = load_function_rca_details()
            _function_rca_cache["signature"] = signature

        return _function_rca_cache["data"]


get_function_rca_data()


def get_dynamic_reporting_months(data=None, rca_actions_df=None, reference_date=None):
    """
    Returns (available_months, default_month) dynamically:
    - Automatically includes all completed months from January of the year up to
      the target reporting month (current month - 1).
      (e.g. in October 2026 -> Jan-Sep 2026;
       when November 2026 starts -> Jan-Oct 2026 automatically).
    - Capped at the target reporting month (current month - 1).
    - Plus any historical months with actual data present in the dataset (up to current month - 1).
    - Plus any historical months present in strategic/functional RCA actions (up to current month - 1).
    - Default month is current month - 1.
    """
    if reference_date is None:
        ref_dt = pd.Timestamp.now()
    else:
        ref_dt = pd.Timestamp(reference_date)

    current_month_start = pd.Timestamp(year=ref_dt.year, month=ref_dt.month, day=1)
    latest_completed_month = (current_month_start - pd.DateOffset(months=1)).floor("D")

    start_year_month = pd.Timestamp(year=latest_completed_month.year, month=1, day=1)
    start_month = start_year_month
    if data is not None and "month" in data.columns and not data["month"].dropna().empty:
        min_m = data["month"].dropna().min()
        if pd.notna(min_m) and pd.Timestamp(min_m) < start_month:
            start_month = pd.Timestamp(min_m).replace(day=1)

    calendar_months = set()
    if start_month <= latest_completed_month:
        calendar_months = set(pd.date_range(start=start_month, end=latest_completed_month, freq="MS"))
    else:
        calendar_months = {latest_completed_month}

    data_months = set()
    if data is not None and "month" in data.columns and "Actual" in data.columns:
        active_data = data
        if "function" in active_data.columns:
            active_data = active_data[
                ~active_data["function"].fillna("").astype(str).str.strip().str.lower().isin(EXCLUDED_DMB_FUNCTIONS)
            ]
        if "Target" in active_data.columns:
            actual_months = active_data.loc[
                active_data["Actual"].notna() & active_data["Target"].notna(), "month"
            ].dropna().unique()
        else:
            actual_months = active_data.loc[active_data["Actual"].notna(), "month"].dropna().unique()
        data_months.update(
            pd.Timestamp(m).replace(day=1) for m in actual_months
            if pd.Timestamp(m).replace(day=1) <= latest_completed_month
        )

    if rca_actions_df is None:
        try:
            rca_actions_df = get_active_rca_actions()
        except Exception:
            rca_actions_df = pd.DataFrame()

    if rca_actions_df is not None and not rca_actions_df.empty:
        if "actual" in rca_actions_df.columns:
            if "target" in rca_actions_df.columns:
                rca_actual_rows = rca_actions_df.loc[
                    rca_actions_df["actual"].notna() & rca_actions_df["target"].notna()
                ]
            else:
                rca_actual_rows = rca_actions_df.loc[rca_actions_df["actual"].notna()]
            if not rca_actual_rows.empty and "reporting_month" in rca_actual_rows.columns:
                rca_months = rca_actual_rows["reporting_month"].dropna().unique()
                data_months.update(
                    pd.Timestamp(m).replace(day=1) for m in rca_months
                    if pd.Timestamp(m).replace(day=1) <= latest_completed_month
                )

    all_months = sorted(m for m in (calendar_months | data_months) if m <= latest_completed_month)
    default_m = latest_completed_month if (latest_completed_month in all_months or not all_months) else all_months[-1]
    return all_months, default_m


def get_available_months(data):
    months, _ = get_dynamic_reporting_months(data)
    return months


def create_month_options(months):
    return [
        {
            "label": html.Span(
                pd.Timestamp(month).strftime("%B %Y"),
                className="month-option-label",
                style={
                    "color": "#142638",
                    "fontWeight": 600,
                },
            ),
            "value": pd.Timestamp(month).strftime("%Y-%m-%d"),
        }
        for month in months
    ]


mpr_months, default_mpr_month = get_dynamic_reporting_months(mpr_data, rca_actions)
dmb_months, default_dmb_month = get_dynamic_reporting_months(dmb_data, rca_actions)
default_month = max(default_mpr_month, default_dmb_month)


# =========================================================
# CREATE DASH APPLICATION
# =========================================================

app = Dash(
    __name__,
    assets_folder=str(ASSETS_DIR),
    assets_url_path="assets",
    serve_locally=True,
    suppress_callback_exceptions=True,
    title="DMB Performance Dashboard",
    update_title=None,
)

server = app.server


@server.route("/api/sync", methods=["GET", "POST"])
def api_sync_endpoint():
    res = sharepoint_sync.sync_now()
    data_loader.reload_all_data(force=True)
    return json.dumps(res), 200, {"Content-Type": "application/json"}


# =========================================================
# LOAD CSS
# =========================================================

if not CSS_FILE.exists():
    raise FileNotFoundError(f"CSS file was not found: {CSS_FILE}")

css_files = sorted(ASSETS_DIR.glob("*.css"))
custom_css = "\n\n".join(
    css_file.read_text(encoding="utf-8")
    for css_file in css_files
)

app.index_string = """
<!DOCTYPE html>
<html>
    <head>
        {%metas%}
        <title>{%title%}</title>
        {%favicon%}
        {%css%}
        <script src="https://html2canvas.hertzen.com/dist/html2canvas.min.js"></script>
        <style>
            __CUSTOM_CSS__
        </style>
    </head>
    <body>
        {%app_entry%}
        <footer>
            {%config%}
            {%scripts%}
            {%renderer%}
        </footer>
    </body>
</html>
""".replace("__CUSTOM_CSS__", custom_css)


# =========================================================
# REUSABLE COMPONENTS
# =========================================================

def insight_card(title, content_id, count_id, card_class, initial_content=None, initial_count=0):
    return html.Div(
        [
            html.Div(
                [
                    html.H3(title),
                    html.Span(str(initial_count), id=count_id, className="insight-count"),
                ],
                className="insight-title-row",
            ),
            html.Div(initial_content, id=content_id),
        ],
        className=f"insight-card {card_class}",
    )


def insight_list(items, empty_text):
    if not items:
        return html.P(empty_text, className="empty-message")

    return html.Ul(
        [html.Li(item) for item in items],
        className="insight-list",
    )


def kpi_card(
    label,
    value_id,
    value_color,
    initial_value="",
    secondary_id=None,
    initial_secondary=None,
    scroll_target=None,
    title=None,
    is_clickable=False,
):
    value_children = [
        html.Span(
            str(initial_value),
            id=value_id,
            className="kpi-card-value",
            style={"color": value_color},
        )
    ]

    if secondary_id:
        value_children.append(
            html.Span(
                str(initial_secondary) if initial_secondary is not None else "",
                id=secondary_id,
                className="kpi-card-secondary",
            )
        )

    card_kwargs = {
        "id": f"card-{value_id}",
    }

    if scroll_target:
        card_kwargs["className"] = "kpi-summary-card kpi-summary-card-clickable"
        card_kwargs["data-scroll-target"] = scroll_target
        card_kwargs["title"] = title or "Click to jump to Cause and Actions of Red KPIs at MoS Level"
        card_kwargs["n_clicks"] = 0
    elif is_clickable:
        card_kwargs["className"] = "kpi-summary-card kpi-summary-card-clickable"
        if title:
            card_kwargs["title"] = title
        card_kwargs["n_clicks"] = 0
    else:
        card_kwargs["className"] = "kpi-summary-card"
        if title:
            card_kwargs["title"] = title

    return html.Div(
        [
            html.P(label, className="kpi-card-label"),
            html.Div(value_children, className="kpi-value-row"),
        ],
        **card_kwargs,
    )


def section_header(
    section_id,
    title,
    subtitle,
    filter_label,
    filter_id,
    filter_options,
    default_value=None,
):
    if default_value is None and filter_options:
        default_value = filter_options[-1]["value"]
    elif default_value is None:
        default_value = default_month.strftime("%Y-%m-%d")

    title_children = [html.H2(title)]
    if subtitle:
        title_children.append(html.P(subtitle))

    control_children = []
    if filter_label:
        control_children.append(html.Label(filter_label))
    control_children.append(
        dcc.Dropdown(
            id=filter_id,
            options=filter_options,
            value=default_value,
            clearable=False,
            searchable=False,
            optionHeight=40,
            maxHeight=280,
            className="section-month-dropdown",
        )
    )

    return html.Div(
        [
            html.Div(title_children),
            html.Div(
                control_children,
                className="section-month-control",
            ),
        ],
        id=section_id,
        className="section-header",
    )


# =========================================================
# GAUGE CHART & FORMATTING
# =========================================================

def format_percentage(val):
    if val is None or pd.isna(val):
        return "—"
    try:
        val_float = float(val)
        rounded = round(val_float, 1)
        if rounded == 100 or rounded == int(rounded):
            return f"{int(rounded)}%"
        return f"{rounded:.1f}%"
    except (ValueError, TypeError):
        return str(val)


def create_gauge(value):
    gauge_color = "#168b69" if value >= GAUGE_TARGET else "#dc3d56"

    figure = go.Figure(
        go.Indicator(
            mode="gauge",
            value=value,
            domain={"x": [0.12, 0.88], "y": [0.05, 1.00]},
            gauge={
                "shape": "angular",
                "axis": {
                    "range": [0, 100],
                    "tickmode": "array",
                    "tickvals": [0, GAUGE_TARGET, 100],
                    "ticktext": ["0%", f"{GAUGE_TARGET}%", "100%"],
                    "tickfont": {
                        "family": "Segoe UI",
                        "size": 12,
                        "color": "#000000",
                    },
                    "tickcolor": "#000000",
                    "tickwidth": 1,
                    "ticklen": 4,
                },
                "bar": {
                    "color": gauge_color,
                    "thickness": 0.55,
                },
                "bgcolor": "#dce7ef",
                "borderwidth": 0,
                "steps": [
                    {
                        "range": [0, 100],
                        "color": "#dce7ef",
                    }
                ],
                "threshold": {
                    "line": {
                        "color": "#0877b9",
                        "width": 4,
                    },
                    "thickness": 0.88,
                    "value": GAUGE_TARGET,
                },
            },
        )
    )

    figure.add_annotation(
        x=0.5,
        y=0.38,
        text=f"<b>{format_percentage(value)}</b>",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 36,
            "color": "#000000",
        },
    )

    figure.add_annotation(
        x=0.5,
        y=0.18,
        text="met & improved",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 13,
            "color": "#000000",
        },
    )

    figure.update_layout(
        height=220,
        margin={"l": 35, "r": 35, "t": 12, "b": 5},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
    )

    return figure


# =========================================================
# DMB FUNCTION CARDS
# =========================================================


def create_mini_gauge(value):
    is_no_data = value is None or pd.isna(value)
    numeric_value = 0 if is_no_data else value
    gauge_color = "#dce7ef" if is_no_data else ("#168b69" if numeric_value >= GAUGE_TARGET else "#dc3d56")

    figure = go.Figure(
        go.Indicator(
            mode="gauge",
            value=numeric_value,
            domain={"x": [0.08, 0.92], "y": [0.04, 1.0]},
            gauge={
                "shape": "angular",
                "axis": {
                    "range": [0, 100],
                    "tickmode": "array",
                    "tickvals": [0, GAUGE_TARGET, 100],
                    "ticktext": [
                        "0%",
                        f"{GAUGE_TARGET}%",
                        "100%",
                    ],
                    "tickfont": {
                        "family": "Segoe UI",
                        "size": 9,
                        "color": "#000000",
                    },
                    "tickcolor": "#000000",
                    "tickwidth": 1,
                    "ticklen": 3,
                },
                "bar": {"color": gauge_color, "thickness": 0.52},
                "bgcolor": "#dce7ef",
                "borderwidth": 0,
                "steps": [
                    {
                        "range": [0, 100],
                        "color": "#dce7ef",
                    }
                ],
                "threshold": {
                    "line": {"color": "#0877b9", "width": 3},
                    "thickness": 0.85,
                    "value": GAUGE_TARGET,
                },
            },
        )
    )

    display_text = "<b>—</b>" if is_no_data else f"<b>{format_percentage(value)}</b>"
    figure.add_annotation(
        x=0.5,
        y=0.34,
        text=display_text,
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 27,
            "color": "#000000",
        },
    )

    figure.add_annotation(
        x=0.5,
        y=0.10,
        text="met & improved",
        showarrow=False,
        font={
            "family": "Segoe UI",
            "size": 11,
            "color": "#000000",
        },
    )

    figure.update_layout(
        height=140,
        margin={"l": 8, "r": 8, "t": 5, "b": 0},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
    )

    return figure


# RCA Table rendering is defined after create_cause_table below.


def create_function_card(function_name, current_data):
    active_dmb = get_active_dmb_data()
    if not current_data.empty and "function" in current_data.columns:
        if function_name == "ISC & Procurement":
            function_rows = current_data[
                current_data["function"].isin(["ISC & Procurement", "ISC", "Procurement"])
            ].copy()
            known_kpis = set(
                k for k in active_dmb[
                    active_dmb["function"].isin(["ISC & Procurement", "ISC", "Procurement"])
                ]["kpi_name"].dropna().unique()
                if clean_text(k)
            ) if not active_dmb.empty and "function" in active_dmb.columns else set()
        else:
            function_rows = current_data[current_data["function"].eq(function_name)].copy()
            known_kpis = set(
                k for k in active_dmb[
                    active_dmb["function"].eq(function_name)
                ]["kpi_name"].dropna().unique()
                if clean_text(k)
            ) if not active_dmb.empty and "function" in active_dmb.columns else set()
    else:
        function_rows = pd.DataFrame()
        known_kpis = set()

    month_kpis = set(
        k for k in function_rows["kpi_name"].dropna().unique()
        if clean_text(k)
    ) if not function_rows.empty and "kpi_name" in function_rows.columns else set()
    all_kpis = known_kpis.union(month_kpis)
    total_expected = max(len(all_kpis), len(function_rows))

    valid = (
        function_rows[
            function_rows["Actual"].notna()
            & function_rows["Target"].notna()
        ].copy()
        if not function_rows.empty
        else pd.DataFrame()
    )

    met = int(valid["is_met"].sum()) if not valid.empty else 0
    not_met = int(valid["status"].eq("Not Met").sum()) if not valid.empty else 0
    filled = met + not_met
    not_filled = max(0, total_expected - filled)
    total = filled + not_filled

    if total > 0:
        met_share = met / total * 100
        not_met_share = not_met / total * 100
        not_filled_share = not_filled / total * 100
    else:
        met_share = 0
        not_met_share = 0
        not_filled_share = 100

    bar_segments = []
    if met > 0:
        bar_segments.append(
            html.Div(
                str(met),
                className="function-bar-met",
                style={"width": f"{met_share}%"},
                title=f"{met} KPIs met",
            )
        )
    if not_met > 0:
        bar_segments.append(
            html.Div(
                str(not_met),
                className="function-bar-not-met",
                style={"width": f"{not_met_share}%"},
                title=f"{not_met} KPIs not met",
            )
        )
    if not_filled > 0:
        bar_segments.append(
            html.Div(
                str(not_filled),
                className="function-bar-not-filled",
                style={"width": f"{not_filled_share}%"},
                title=f"{not_filled} KPIs not filled",
            )
        )
    if not bar_segments:
        bar_segments.append(
            html.Div(
                "0",
                className="function-bar-not-filled",
                style={"width": "100%"},
                title="0 KPIs",
            )
        )

    if valid.empty:
        return html.Div(
            [
                html.Div(
                    [
                        html.H3(function_name),
                        html.Span("—", className="function-movement movement-neutral"),
                    ],
                    className="function-card-header",
                ),
                html.P(
                    "Core & enabling KPI performance",
                    className="function-card-subtitle",
                ),
                html.Div(
                    bar_segments,
                    className="function-stacked-bar",
                ),
                html.Div(
                    [
                        html.Button(
                            [
                                html.Div(
                                    [
                                        html.Span("Met & improved"),
                                        html.Span(
                                            f"Target {GAUGE_TARGET}%",
                                            className="mini-target-label",
                                        ),
                                    ],
                                    className="mini-gauge-title",
                                ),
                                dcc.Graph(
                                    figure=create_mini_gauge(None),
                                    config={"displayModeBar": False},
                                    className="mini-gauge-graph",
                                ),
                            ],
                            id={
                                "type": "function-gauge-card",
                                "function": function_name,
                            },
                            n_clicks=0,
                            className="mini-gauge-container mini-gauge-container-clickable",
                            title=f"Click to view {function_name} KPI definitions and trends",
                        ),
                        html.Button(
                            [
                                html.Span(
                                    "0",
                                    className="continuous-red-value no-red-value",
                                ),
                                html.P("Red KPIs"),
                                html.Small("No red KPI"),
                            ],
                            id={
                                "type": "continuous-red-card",
                                "function": function_name,
                            },
                            n_clicks=0,
                            disabled=True,
                            title="No data available for this function in this month",
                            className="continuous-red-panel",
                        ),
                    ],
                    className="function-card-bottom",
                ),
            ],
            className="function-card",
        )

    positive = int(valid["is_met_or_improved"].sum())
    performance = positive / len(valid) * 100 if len(valid) else 0

    comparison = valid[
        valid["previous_status"].isin(["Met", "Not Met"])
    ].copy()

    if comparison.empty:
        movement_text = "—"
        movement_class = "movement-neutral"
    else:
        current_rate = comparison["status"].eq("Met").mean() * 100
        previous_rate = comparison["previous_status"].eq("Met").mean() * 100
        movement = current_rate - previous_rate

        if movement > 0.05:
            movement_text = "▲"
            movement_class = "movement-up"
        elif movement < -0.05:
            movement_text = "▼"
            movement_class = "movement-down"
        else:
            movement_text = "—"
            movement_class = "movement-neutral"

    return html.Div(
        [
            html.Div(
                [
                    html.H3(function_name),
                    html.Span(
                        movement_text,
                        className=f"function-movement {movement_class}",
                    ),
                ],
                className="function-card-header",
            ),
            html.P(
                "Core & enabling KPI performance",
                className="function-card-subtitle",
            ),
            html.Div(
                bar_segments,
                className="function-stacked-bar",
            ),
            html.Div(
                [
                    html.Button(
                        [
                            html.Div(
                                [
                                    html.Span("Met & improved"),
                                    html.Span(
                                        f"Target {GAUGE_TARGET}%",
                                        className="mini-target-label",
                                    ),
                                ],
                                className="mini-gauge-title",
                            ),
                            dcc.Graph(
                                figure=create_mini_gauge(performance),
                                config={"displayModeBar": False},
                                className="mini-gauge-graph",
                            ),
                        ],
                        id={
                            "type": "function-gauge-card",
                            "function": function_name,
                        },
                        n_clicks=0,
                        className="mini-gauge-container mini-gauge-container-clickable",
                        title=f"Click to view {function_name} KPI definitions and 6-month trends",
                    ),
                    html.Button(
                        [
                            html.Span(
                                str(not_met),
                                className=(
                                    "continuous-red-value"
                                    if not_met
                                    else "continuous-red-value no-red-value"
                                ),
                            ),
                            html.P("Red KPIs"),
                            html.Small(
                                (
                                    "Click to view root causes and actions"
                                    if not_met
                                    else "No red KPI"
                                )
                            ),
                        ],
                        id={
                            "type": "continuous-red-card",
                            "function": function_name,
                        },
                        n_clicks=0,
                        disabled=not_met == 0,
                        title=(
                            "View root causes and corrective actions"
                            if not_met
                            else "No red KPI for this function"
                        ),
                        className=(
                            "continuous-red-panel continuous-red-panel-active"
                            if not_met
                            else "continuous-red-panel"
                        ),
                    ),
                ],
                className="function-card-bottom",
            ),
        ],
        className="function-card",
    )






def is_closed_action_status(status_val):
    if not status_val or pd.isna(status_val):
        return False
    s = str(status_val).strip().lower()
    return (
        s in {"5", "5.0", "●", "✔", "✓", "action completed", "completed", "closed", "resolution confirmed", "confirmed"}
        or "completed" in s
        or "confirmed" in s
        or "closed" in s
    )


def format_impact_badge(impact_val):
    parsed = parse_impact_percent_value(impact_val)
    if parsed is not None:
        if round(parsed, 1).is_integer():
            txt = f"{parsed:.0f}%"
        else:
            txt = f"{parsed:.1f}%"
        return html.Span(txt, className="rca-impact-pill")
    if impact_val is not None and pd.notna(impact_val):
        s = str(impact_val).strip()
        if s and s not in {"—", "-", "none", "nan", "null"}:
            disp = s if s.endswith("%") else f"{s}%"
            return html.Span(disp, className="rca-impact-pill")
    return html.Span("—", className="rca-dash-text")


def calculate_row_spans(total_rows, count):
    if count <= 0:
        return []
    base = total_rows // count
    extra = total_rows % count
    spans = []
    curr = 0
    for i in range(count):
        span = base + (1 if i < extra else 0)
        spans.append((curr, span))
        curr += span
    return spans


def format_kpi_metric_display(val, unit=""):
    if val is None or pd.isna(val):
        return None
    try:
        val_f = float(val)
    except (ValueError, TypeError):
        u_str = f" {unit}".strip() if unit else ""
        return f"{val}{u_str}".strip()
    unit_clean = str(unit).strip() if unit and pd.notna(unit) else ""
    if unit_clean == "%":
        if 0 < val_f <= 1.0:
            val_f *= 100
        return f"{val_f:.1f}%"
    elif unit_clean == "Mn":
        return f"{val_f:.2f} Mn"
    elif unit_clean.lower() in ["no.", "no", "nos"]:
        return f"{int(val_f)}" if val_f == int(val_f) else f"{val_f:.1f}"
    else:
        suffix = f" {unit_clean}" if unit_clean else ""
        return f"{val_f:.2f}{suffix}" if val_f % 1 != 0 else f"{int(val_f)}{suffix}"


def parse_num_and_unit(val):
    if val is None or pd.isna(val):
        return None, ""
    if isinstance(val, (int, float)):
        return float(val), ""
    s = str(val).strip()
    if not s or s in {"—", "-", "–", "nan", "None", "null"}:
        return None, ""
    unit = ""
    if "%" in s:
        unit = "%"
    elif "Mn" in s:
        unit = "Mn"
    elif "k" in s.lower():
        unit = "k"
    elif "day" in s.lower():
        unit = "days"

    clean_s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(clean_s), unit
    except (ValueError, TypeError):
        return None, ""


def get_kpi_metric_nature(kpi_name, default="Higher the better"):
    if not kpi_name:
        return default
    k_norm = function_key(str(kpi_name))
    k_clean = str(kpi_name).strip().lower()

    # 1. Look in DMB data
    try:
        active_dmb = get_active_dmb_data()
        if not active_dmb.empty and "kpi_name" in active_dmb.columns and "metric_nature" in active_dmb.columns:
            m = active_dmb[active_dmb["kpi_name"].apply(lambda x: function_key(str(x)) == k_norm or str(x).strip().lower() == k_clean)]
            if not m.empty:
                val = m["metric_nature"].dropna()
                if not val.empty and str(val.iloc[0]).strip():
                    return str(val.iloc[0]).strip()
    except Exception:
        pass

    # 2. Look in MPR data
    try:
        active_mpr = get_active_mpr_data()
        if not active_mpr.empty and "kpi_name" in active_mpr.columns and "metric_nature" in active_mpr.columns:
            m = active_mpr[active_mpr["kpi_name"].apply(lambda x: function_key(str(x)) == k_norm or str(x).strip().lower() == k_clean)]
            if not m.empty:
                val = m["metric_nature"].dropna()
                if not val.empty and str(val.iloc[0]).strip():
                    return str(val.iloc[0]).strip()
    except Exception:
        pass

    # 3. Look in review_kpis
    try:
        rca_data = get_function_rca_data()
        review_kpis = rca_data.get("review_kpis", pd.DataFrame())
        if not review_kpis.empty and "kpi_name" in review_kpis.columns and "metric_nature" in review_kpis.columns:
            m = review_kpis[review_kpis["kpi_name"].apply(lambda x: function_key(str(x)) == k_norm or str(x).strip().lower() == k_clean)]
            if not m.empty:
                val = m["metric_nature"].dropna()
                if not val.empty and str(val.iloc[0]).strip():
                    return str(val.iloc[0]).strip()
    except Exception:
        pass

    return default


def compute_kpi_delta_badge(
    target_val,
    actual_val,
    target_raw=None,
    actual_raw=None,
    unit_hint="",
    metric_nature=None,
    kpi_name=None,
):
    t_num, t_u = parse_num_and_unit(target_raw if target_raw is not None else target_val)
    a_num, a_u = parse_num_and_unit(actual_raw if actual_raw is not None else actual_val)
    if t_num is None or a_num is None:
        return None
    u = unit_hint or a_u or t_u
    if (u == "%" or unit_hint == "%") and (0 < t_num <= 1.0 and 0 < a_num <= 1.0):
        t_num *= 100
        a_num *= 100
        u = "%"

    diff = a_num - t_num
    if abs(diff) < 1e-6:
        val_str = "0.0%" if u == "%" else ("0" if diff == int(diff) else f"{diff:.2f}")
        return {
            "val": val_str,
            "class": "rca-metric-pill-delta rca-delta-neutral",
            "label": "DELTA: ",
            "title": "On target (no variance)",
        }

    # Determine metric nature: Higher the better vs Lower the better
    nature_str = metric_nature or (get_kpi_metric_nature(kpi_name) if kpi_name else "Higher the better")
    lower_is_better = "lower" in str(nature_str).lower()

    # Determine favorable vs unfavorable (Green vs Red):
    # - For Higher the better: Actual > Target (diff > 0) is Favorable (Green), Actual < Target (diff < 0) is Unfavorable (Red)
    # - For Lower the better: Actual < Target (diff < 0) is Favorable (Green), Actual > Target (diff > 0) is Unfavorable (Red)
    is_favorable = (diff < 0) if lower_is_better else (diff > 0)

    dir_class = "rca-delta-increase" if is_favorable else "rca-delta-dip"
    arrow = "▲" if diff > 0 else "▼"
    sign = "+" if diff > 0 else ""

    if u == "%":
        fmt_val = f"{arrow} {sign}{diff:.1f}%"
    elif u == "Mn":
        fmt_val = f"{arrow} {sign}{diff:.2f} Mn"
    elif u.lower() in ["no.", "no", "nos"]:
        fmt_val = f"{arrow} {sign}{int(diff)}" if diff == int(diff) else f"{arrow} {sign}{diff:.1f}"
    else:
        suffix = f" {u}".rstrip() if u else ""
        fmt_val = f"{arrow} {sign}{diff:.2f}{suffix}" if diff % 1 != 0 else f"{arrow} {sign}{int(diff)}{suffix}"

    if lower_is_better:
        if is_favorable:
            desc = f"Actual is {abs(diff):.1f}{u} below target (Favorable for Lower-the-better metric)"
        else:
            desc = f"Actual is {diff:.1f}{u} above target (Unfavorable for Lower-the-better metric)"
    else:
        if is_favorable:
            desc = f"Actual is {diff:.1f}{u} above target (Met / Exceeded Target)"
        else:
            desc = f"Actual is {abs(diff):.1f}{u} below target (Gap / Below Target)"

    return {
        "val": fmt_val,
        "class": f"rca-metric-pill-delta {dir_class}",
        "label": "DELTA: ",
        "title": desc,
    }


def compute_cause_action_match_score(cause_text, action_root_cause, action_desc=""):
    c_clean = str(cause_text).strip().lower()
    rc_clean = str(action_root_cause).strip().lower()
    desc_clean = str(action_desc).strip().lower()

    if not c_clean or c_clean in {"none", "nan", "—", ""}:
        return 0.0

    # 1. Exact match
    if c_clean == rc_clean:
        return 1.0

    # 2. Substring match
    if len(c_clean) >= 4 and (c_clean in rc_clean or rc_clean in c_clean):
        return 0.95

    # 3. Token overlap with root_cause (primary) and action_desc (secondary)
    stopwords = {'in', 'for', 'the', 'to', 'due', 'of', 'and', 'a', 'on', 'by', 'is', 'at', 'with', 'from', 'as', 'q1', 'q2', 'q3', 'q4', 'oit', 'npi', 'npis', 'kpi', 'gap', 'rate'}
    words = re.findall(r'[a-zA-Z0-9]+', c_clean)
    c_toks = set()
    for w in words:
        if len(w) >= 3 and w not in stopwords:
            c_toks.add(w)
            if w.endswith('s') and len(w) > 3:
                c_toks.add(w[:-1])
            if w.endswith('ing') and len(w) > 4:
                c_toks.add(w[:-3])

    if not c_toks:
        return 0.0

    rc_words = re.findall(r'[a-zA-Z0-9]+', rc_clean)
    rc_toks = set()
    for w in rc_words:
        if len(w) >= 3 and w not in stopwords:
            rc_toks.add(w)
            if w.endswith('s') and len(w) > 3:
                rc_toks.add(w[:-1])
            if w.endswith('ing') and len(w) > 4:
                rc_toks.add(w[:-3])

    desc_words = re.findall(r'[a-zA-Z0-9]+', desc_clean)
    desc_toks = set()
    for w in desc_words:
        if len(w) >= 3 and w not in stopwords:
            desc_toks.add(w)
            if w.endswith('s') and len(w) > 3:
                desc_toks.add(w[:-1])
            if w.endswith('ing') and len(w) > 4:
                desc_toks.add(w[:-3])

    common_rc = c_toks & rc_toks
    common_desc = c_toks & desc_toks

    score = 0.0
    if common_rc:
        score += len(common_rc) * 0.4
    if common_desc:
        score += len(common_desc) * 0.2

    ratio = score / len(c_toks)
    return min(1.0, ratio)


def create_kpi_rca_card(
    source_kpi_name,
    causes,
    actions_to_show,
    red_kpi_name=None,
    function_badge=None,
    card_id=None,
    is_red=True,
    missing_action_text=None,
    target_val=None,
    actual_val=None,
    reporting_month=None,
    metric_nature=None,
):
    disp_title = red_kpi_name or source_kpi_name
    if missing_action_text is None:
        missing_action_text = (
            "All corrective actions completed"
            if actions_to_show is not None and not actions_to_show.empty
            else "No Corrective Action provided"
        )

    # Prepare Causes DataFrame
    if causes is not None and not causes.empty:
        causes_df = (
            causes.sort_values("cause_rank").copy()
            if "cause_rank" in causes.columns
            else causes.copy()
        ).reset_index(drop=True)
    else:
        causes_df = pd.DataFrame()

    # Prepare Actions DataFrame
    if actions_to_show is not None and not actions_to_show.empty:
        open_actions = (
            actions_to_show[
                ~actions_to_show["status"].apply(is_closed_action_status)
            ].copy()
            if "status" in actions_to_show.columns
            else actions_to_show.copy()
        ).reset_index(drop=True)
    else:
        open_actions = pd.DataFrame()

    # Header Row 1: Top Categories (Top Causes = 2 cols, Corrective Action = 4 cols)
    thead_row_1 = html.Tr(
        [
            html.Th(
                "Top Causes",
                colSpan=2,
                className="rca-group-head-causes",
            ),
            html.Th(
                "Corrective Action",
                colSpan=4,
                className="rca-group-head-actions",
            ),
        ]
    )

    # Header Row 2: Sub-headers (6 columns total: Cause, Impact, Root Cause, Action Description, Owner, Status)
    thead_row_2 = html.Tr(
        [
            html.Th("Cause", className="rca-subhead-cell rca-col-cause"),
            html.Th("Impact on KPI Gap", className="rca-subhead-cell rca-col-impact"),
            html.Th("Root Cause", className="rca-subhead-cell rca-col-root-cause"),
            html.Th("Action Description", className="rca-subhead-cell rca-col-action-desc"),
            html.Th("Owner", className="rca-subhead-cell rca-col-owner"),
            html.Th("Status", className="rca-subhead-cell rca-col-status"),
        ]
    )

    # Match each Action to its most relevant Cause
    matched_actions_by_cause = {c_i: [] for c_i in range(len(causes_df))}
    assigned_action_indices = set()

    if not causes_df.empty and not open_actions.empty:
        for a_idx, a_row in open_actions.iterrows():
            rc_text = a_row.get("root_cause", "")
            a_desc = a_row.get("corrective_action", "")
            best_c_idx = None
            best_sc = 0.19  # Threshold for meaningful match
            for c_idx, c_row in causes_df.iterrows():
                c_text = c_row.get("cause", "")
                sc = compute_cause_action_match_score(c_text, rc_text, a_desc)
                if sc > best_sc:
                    best_sc = sc
                    best_c_idx = c_idx
            if best_c_idx is not None:
                matched_actions_by_cause[best_c_idx].append(a_row)
                assigned_action_indices.add(a_idx)

    unmatched_actions = [
        open_actions.iloc[a_i] for a_i in range(len(open_actions))
        if a_i not in assigned_action_indices
    ]

    tbody_rows = []

    # Case 1: Causes are present
    if not causes_df.empty:
        for c_idx, c_row in causes_df.iterrows():
            c_text = clean_cell_text(c_row.get("cause", "")) or "—"
            c_impact = format_impact_badge(c_row.get("impact_percent"))
            c_actions = matched_actions_by_cause.get(c_idx, [])

            if c_actions:
                for sub_i, a_row in enumerate(c_actions):
                    row_cells = []
                    if sub_i == 0:
                        row_cells.append(
                            html.Td(
                                c_text,
                                rowSpan=len(c_actions) if len(c_actions) > 1 else None,
                                className="rca-recovery-cell rca-cause-text",
                            )
                        )
                        row_cells.append(
                            html.Td(
                                c_impact,
                                rowSpan=len(c_actions) if len(c_actions) > 1 else None,
                                className="rca-recovery-cell text-center rca-col-impact-cell",
                            )
                        )
                    a_desc = clean_cell_text(a_row.get("corrective_action", ""))
                    while a_desc and a_desc[0] in {";", "-", "–", "—", ":", " ", "\t"}:
                        a_desc = a_desc[1:].strip()

                    root_cause_str = clean_cell_text(a_row.get("root_cause", ""))
                    while root_cause_str and root_cause_str[0] in {";", "-", "–", "—", ":", " ", "\t"}:
                        root_cause_str = root_cause_str[1:].strip()

                    root_cause_display = (
                        root_cause_str
                        if root_cause_str and root_cause_str.lower() not in {"not entered", "none", "—", ""}
                        else "—"
                    )
                    a_owner = clean_cell_text(a_row.get("owner", "")) or "—"
                    a_status_text = clean_cell_text(a_row.get("status", "")) or "Status not entered"
                    a_status_cls = a_row.get("status_class", "action-status-unknown")
                    a_status_pill = html.Span(a_status_text, className=f"action-status-pill {a_status_cls}")

                    row_cells.append(html.Td(root_cause_display, className="rca-recovery-cell rca-root-cause-text"))
                    row_cells.append(html.Td(a_desc or "—", className="rca-recovery-cell rca-action-text"))
                    row_cells.append(html.Td(a_owner, className="rca-recovery-cell text-center rca-owner-text"))
                    row_cells.append(html.Td(a_status_pill, className="rca-recovery-cell text-center"))
                    tbody_rows.append(html.Tr(row_cells, className="rca-recovery-row"))
            else:
                # No action matched for this cause -> write No Corrective Action provided in RED
                row_cells = [
                    html.Td(c_text, className="rca-recovery-cell rca-cause-text"),
                    html.Td(c_impact, className="rca-recovery-cell text-center rca-col-impact-cell"),
                    html.Td(
                        "No Corrective Action provided",
                        colSpan=2,
                        className="rca-recovery-cell rca-action-text rca-missing-text",
                    ),
                    html.Td("—", className="rca-recovery-cell text-center rca-dash-text"),
                    html.Td("—", className="rca-recovery-cell text-center rca-dash-text"),
                ]
                tbody_rows.append(html.Tr(row_cells, className="rca-recovery-row"))

        # Extra actions not matched to any cause
        for a_row in unmatched_actions:
            a_desc = clean_cell_text(a_row.get("corrective_action", ""))
            while a_desc and a_desc[0] in {";", "-", "–", "—", ":", " ", "\t"}:
                a_desc = a_desc[1:].strip()

            root_cause_str = clean_cell_text(a_row.get("root_cause", ""))
            while root_cause_str and root_cause_str[0] in {";", "-", "–", "—", ":", " ", "\t"}:
                root_cause_str = root_cause_str[1:].strip()

            root_cause_display = (
                root_cause_str
                if root_cause_str and root_cause_str.lower() not in {"not entered", "none", "—", ""}
                else "—"
            )
            a_owner = clean_cell_text(a_row.get("owner", "")) or "—"
            a_status_text = clean_cell_text(a_row.get("status", "")) or "Status not entered"
            a_status_cls = a_row.get("status_class", "action-status-unknown")
            a_status_pill = html.Span(a_status_text, className=f"action-status-pill {a_status_cls}")

            row_cells = [
                html.Td("", className="rca-recovery-cell rca-cause-text"),
                html.Td("", className="rca-recovery-cell text-center rca-col-impact-cell"),
                html.Td(root_cause_display, className="rca-recovery-cell rca-root-cause-text"),
                html.Td(a_desc or "—", className="rca-recovery-cell rca-action-text"),
                html.Td(a_owner, className="rca-recovery-cell text-center rca-owner-text"),
                html.Td(a_status_pill, className="rca-recovery-cell text-center"),
            ]
            tbody_rows.append(html.Tr(row_cells, className="rca-recovery-row"))

    # Case 2: Causes are empty, actions exist
    elif not open_actions.empty:
        for _, a_row in open_actions.iterrows():
            a_desc = clean_cell_text(a_row.get("corrective_action", ""))
            while a_desc and a_desc[0] in {";", "-", "–", "—", ":", " ", "\t"}:
                a_desc = a_desc[1:].strip()

            root_cause_str = clean_cell_text(a_row.get("root_cause", ""))
            while root_cause_str and root_cause_str[0] in {";", "-", "–", "—", ":", " ", "\t"}:
                root_cause_str = root_cause_str[1:].strip()

            root_cause_display = (
                root_cause_str
                if root_cause_str and root_cause_str.lower() not in {"not entered", "none", "—", ""}
                else "—"
            )
            a_owner = clean_cell_text(a_row.get("owner", "")) or "—"
            a_status_text = clean_cell_text(a_row.get("status", "")) or "Status not entered"
            a_status_cls = a_row.get("status_class", "action-status-unknown")
            a_status_pill = html.Span(a_status_text, className=f"action-status-pill {a_status_cls}")

            row_cells = [
                html.Td("", className="rca-recovery-cell rca-cause-text"),
                html.Td("", className="rca-recovery-cell text-center rca-col-impact-cell"),
                html.Td(root_cause_display, className="rca-recovery-cell rca-root-cause-text"),
                html.Td(a_desc or "—", className="rca-recovery-cell rca-action-text"),
                html.Td(a_owner, className="rca-recovery-cell text-center rca-owner-text"),
                html.Td(a_status_pill, className="rca-recovery-cell text-center"),
            ]
            tbody_rows.append(html.Tr(row_cells, className="rca-recovery-row"))

    # Case 3: Both Causes and Actions are empty
    else:
        row_cells = [
            html.Td("No RCA provided", className="rca-recovery-cell rca-cause-text rca-missing-text"),
            html.Td("—", className="rca-recovery-cell text-center rca-col-impact-cell rca-dash-text"),
            html.Td("No Corrective Action provided", colSpan=2, className="rca-recovery-cell rca-action-text rca-missing-text"),
            html.Td("—", className="rca-recovery-cell text-center rca-dash-text"),
            html.Td("—", className="rca-recovery-cell text-center rca-dash-text"),
        ]
        tbody_rows.append(html.Tr(row_cells, className="rca-recovery-row"))

    card_kwargs = {"className": "kpi-rca-action-card"}
    if card_id:
        card_kwargs["id"] = card_id

    meta_badges = []
    if target_val is not None and str(target_val).strip() and str(target_val).strip() not in {"—", "nan", "None"}:
        meta_badges.append(
            html.Span(
                [
                    html.Span("TARGET: ", className="rca-metric-pill-lbl"),
                    html.Strong(str(target_val), className="rca-metric-pill-val"),
                ],
                className="rca-metric-pill rca-metric-pill-target",
            )
        )
    if actual_val is not None and str(actual_val).strip() and str(actual_val).strip() not in {"—", "nan", "None"}:
        meta_badges.append(
            html.Span(
                [
                    html.Span("ACTUAL: ", className="rca-metric-pill-lbl"),
                    html.Strong(str(actual_val), className="rca-metric-pill-val rca-metric-val-actual"),
                ],
                className="rca-metric-pill rca-metric-pill-actual",
            )
        )

    delta_badge = compute_kpi_delta_badge(
        target_val,
        actual_val,
        metric_nature=metric_nature,
        kpi_name=red_kpi_name or source_kpi_name,
    )
    if delta_badge:
        meta_badges.append(
            html.Span(
                [
                    html.Span(delta_badge["label"], className="rca-metric-pill-lbl"),
                    html.Strong(delta_badge["val"], className="rca-metric-pill-val rca-metric-val-delta"),
                ],
                className=f"rca-metric-pill {delta_badge['class']}",
                title=delta_badge.get("title", ""),
            )
        )

    header_right_children = []
    if meta_badges:
        header_right_children.append(
            html.Div(meta_badges, className="rca-header-metrics-group")
        )
    if function_badge:
        header_right_children.append(
            html.Span(function_badge, className="rca-function-pill")
        )

    header_children = [
        html.Div(
            [
                html.Span(
                    [
                        html.Span(className="rca-red-badge-dot"),
                        html.Span(
                            "RED KPI" if is_red else "OTHER RCA",
                            className="rca-red-kpi-badge",
                        ),
                    ],
                    className=(
                        "rca-header-badge-group"
                        if is_red
                        else "rca-header-badge-group rca-badge-not-red"
                    ),
                ),
                html.H3(
                    disp_title,
                    className="rca-recovery-card-title",
                ),
            ],
            className="rca-card-title-group",
        ),
    ]
    if header_right_children:
        header_children.append(
            html.Div(header_right_children, className="rca-header-right-group")
        )

    return html.Div(
        [
            html.Div(
                header_children,
                className="rca-recovery-card-header",
            ),
            html.Div(
                html.Table(
                    [
                        html.Thead([thead_row_1, thead_row_2]),
                        html.Tbody(tbody_rows),
                    ],
                    className="rca-recovery-table",
                ),
                className="rca-recovery-table-wrap",
            ),
        ],
        **card_kwargs,
    )


def slug_kpi_id(s):
    return "rca-kpi-card-" + re.sub(r"[^a-zA-Z0-9_-]", "_", str(s)).lower().strip("_")


def create_rca_table(selected_month):
    selected_month = pd.Timestamp(selected_month)
    rca_actions = get_active_rca_actions()
    monthly_actions = rca_actions[
        rca_actions["reporting_month"].eq(selected_month)
    ].copy()

    if monthly_actions.empty:
        return html.Div(
            f"No Red KPI cause and action entries recorded for {selected_month.strftime('%B %Y')}.",
            className="rca-empty-message",
        )

    if "kpi_name" in monthly_actions.columns:
        monthly_actions = monthly_actions[~monthly_actions["kpi_name"].map(is_value_lever)].copy()

    for column in [
        "strategic_imperative",
        "kpi_name",
        "cause",
        "action",
    ]:
        if column in monthly_actions.columns:
            monthly_actions[column] = monthly_actions[column].fillna("").map(
                lambda value: str(value).strip()
            )
        else:
            monthly_actions[column] = ""

    # Filter for all Red KPIs in the selected month with both Target and Actual
    if "is_red" in monthly_actions.columns:
        monthly_actions = monthly_actions[
            monthly_actions["is_red"]
            & monthly_actions["kpi_name"].ne("")
            & monthly_actions["target"].notna()
            & monthly_actions["actual"].notna()
        ].drop_duplicates(["strategic_imperative", "kpi_name"], keep="last")
    else:
        monthly_actions = monthly_actions[
            monthly_actions["kpi_name"].ne("")
            & monthly_actions["target"].notna()
            & monthly_actions["actual"].notna()
        ].drop_duplicates(["strategic_imperative", "kpi_name"], keep="last")

    if monthly_actions.empty:
        return html.Div(
            f"No Red KPIs recorded for {selected_month.strftime('%B %Y')}.",
            className="rca-empty-message",
        )

    # Build exact lookup maps across strategic RCA records (exact KPI name only)
    strategic_cause_map = {}
    strategic_action_map = {}

    for _, r in rca_actions.iterrows():
        k_name = str(r.get("kpi_name", "")).strip()
        c_val = str(r.get("cause", "")).strip()
        a_val = str(r.get("action", "")).strip()

        c_clean = "" if (not c_val or c_val.lower() in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a", "no rca", "no rca provided"}) else data_loader.get_crisp_rca_text(k_name, c_val, "cause")
        a_clean = "" if (not a_val or a_val.lower() in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a", "no corrective actions provided", "no action", "no actions", "no actions provided"}) else data_loader.get_crisp_rca_text(k_name, a_val, "action")

        if c_clean and k_name not in strategic_cause_map:
            strategic_cause_map[k_name] = c_clean
        if a_clean and k_name not in strategic_action_map:
            strategic_action_map[k_name] = a_clean

    table_rows = []

    for _, row in monthly_actions.iterrows():
        kpi_name_val = row["kpi_name"]
        curr_c = str(row.get("cause", "")).strip()
        curr_a = str(row.get("action", "")).strip()

        raw_c = curr_c if (curr_c and curr_c.lower() not in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a", "no rca", "no rca provided"}) else strategic_cause_map.get(kpi_name_val, "")
        raw_a = curr_a if (curr_a and curr_a.lower() not in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a", "no corrective actions provided", "no action", "no actions", "no actions provided"}) else strategic_action_map.get(kpi_name_val, "")

        cause_val = data_loader.get_crisp_rca_text(kpi_name_val, raw_c, "cause")
        action_val = data_loader.get_crisp_rca_text(kpi_name_val, raw_a, "action")

        is_cause_missing = (
            not cause_val
            or cause_val in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a"}
            or cause_val.lower() in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a", "no rca", "no rca provided"}
        )
        if is_cause_missing:
            cause_cell = html.Span("No RCA provided", className="rca-missing-text")
        else:
            cause_cell = html.Span(cause_val)

        is_action_missing = (
            not action_val
            or action_val in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a"}
            or action_val.lower() in {"—", "-", "–", "none", "nan", "null", "not entered", "na", "n/a", "no corrective actions provided", "no action", "no actions", "no actions provided"}
        )
        if is_action_missing:
            action_cell = html.Span("No Actions Provided", className="rca-missing-text")
        else:
            action_cell = html.Span(action_val)

        # Get units for Target and Actual formatting
        u = str(row.get("units", "")).strip()
        if not u:
            active_mpr = get_active_mpr_data()
            if not active_mpr.empty and "kpi_name" in active_mpr.columns and "units" in active_mpr.columns:
                m = active_mpr[active_mpr["kpi_name"].str.strip().str.lower().eq(str(kpi_name_val).strip().lower())]
                if not m.empty and pd.notna(m["units"].iloc[0]):
                    u = str(m["units"].iloc[0]).strip()

        target_disp = format_kpi_metric_display(row.get("target"), unit=u)
        actual_disp = format_kpi_metric_display(row.get("actual"), unit=u)

        table_rows.append(
            html.Tr(
                [
                    html.Td(row["strategic_imperative"], className="rca-imperative-name"),
                    html.Td(row["kpi_name"], className="rca-kpi-name"),
                    html.Td(html.Span(target_disp or "—"), className="rca-target-cell"),
                    html.Td(html.Span(actual_disp or "—"), className="rca-actual-cell"),
                    html.Td(cause_cell, className="rca-cause-text"),
                    html.Td(action_cell, className="rca-action-text"),
                ]
            )
        )

    return html.Div(
        html.Table(
            [
                html.Thead(
                    html.Tr(
                        [
                            html.Th("Strategic Imperatives"),
                            html.Th("KPI Name"),
                            html.Th("Target"),
                            html.Th("Actual"),
                            html.Th("Cause"),
                            html.Th("Actions"),
                        ]
                    )
                ),
                html.Tbody(table_rows),
            ],
            className="rca-table",
        ),
        className="rca-table-wrap",
    )




def wrap_chart_label(value, maximum_line_length=20):
    words = clean_cell_text(value).split()
    lines = []
    current_line = []

    for word in words:
        candidate = " ".join(current_line + [word])
        if current_line and len(candidate) > maximum_line_length:
            lines.append(" ".join(current_line))
            current_line = [word]
        else:
            current_line.append(word)

    if current_line:
        lines.append(" ".join(current_line))

    return "<br>".join(lines)


def create_pareto_chart(causes):
    chart_data = causes.dropna(subset=["impact_percent"]).copy()
    chart_data["impact_percent"] = pd.to_numeric(
        chart_data["impact_percent"],
        errors="coerce",
    )
    chart_data = chart_data.dropna(subset=["impact_percent"])
    chart_data = chart_data.sort_values(
        ["impact_percent", "cause_rank"],
        ascending=[False, True],
    )

    figure = go.Figure()

    if chart_data.empty:
        figure.add_annotation(
            x=0.5,
            y=0.5,
            text="No quantified cause contribution is available.",
            showarrow=False,
            font={"family": "Segoe UI", "size": 12, "color": "#6e879b"},
        )
    else:
        total_impact = chart_data["impact_percent"].sum()
        if total_impact:
            cumulative_percentage = (
                chart_data["impact_percent"].cumsum()
                / total_impact
                * 100
            )
        else:
            cumulative_percentage = chart_data["impact_percent"] * 0

        cause_labels = [
            f"Cause {int(cause_rank)}"
            for cause_rank in chart_data["cause_rank"].tolist()
        ]

        figure.add_trace(
            go.Bar(
                x=cause_labels,
                y=chart_data["impact_percent"],
                name="Impact on KPI gap",
                marker={
                    "color": "#48a9c2",
                    "line": {"color": "#19718d", "width": 1},
                },
                text=[
                    f"{value:.1f}%" if round(value, 1) % 1 != 0 else f"{value:.0f}%"
                    for value in chart_data["impact_percent"]
                ],
                textposition="outside",
                cliponaxis=False,
                customdata=chart_data["cause"],
                hovertemplate=(
                    "<b>%{x}</b><br>%{customdata}"
                    "<br>Impact on KPI gap: %{y:.1f}%<extra></extra>"
                ),
            )
        )
        figure.add_trace(
            go.Scatter(
                x=cause_labels,
                y=cumulative_percentage,
                name="Cumulative impact",
                mode="lines+markers",
                yaxis="y2",
                line={"color": "#dc3d56", "width": 3},
                marker={
                    "size": 8,
                    "color": "#ffffff",
                    "line": {"color": "#dc3d56", "width": 2},
                },
                text=[
                    f"{value:.0f}%"
                    for value in cumulative_percentage
                ],
                textposition="top center",
                customdata=chart_data["cause"],
                hovertemplate=(
                    "<b>%{x}</b><br>%{customdata}"
                    "<br>Cumulative impact: %{y:.0f}%<extra></extra>"
                ),
            )
        )

    figure.update_layout(
        autosize=True,
        height=330,
        margin={"l": 58, "r": 62, "t": 48, "b": 52},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="#ffffff",
        bargap=0.32,
        hovermode="closest",
        legend={
            "orientation": "h",
            "x": 0,
            "xanchor": "left",
            "y": 1.02,
            "yanchor": "bottom",
            "font": {"size": 9},
        },
        xaxis={
            "title": "Causes ranked by impact",
            "tickfont": {"size": 10, "color": "#000000"},
            "showgrid": False,
            "automargin": True,
        },
        yaxis={
            "title": "Impact on KPI gap",
            "range": [0, 110],
            "ticksuffix": "%",
            "dtick": 20,
            "gridcolor": "#dfe8ee",
            "zeroline": False,
            "tickfont": {"size": 9, "color": "#000000"},
        },
        yaxis2={
            "title": {
                "text": "Cumulative impact",
                "font": {"size": 10, "color": "#dc3d56"},
            },
            "overlaying": "y",
            "side": "right",
            "range": [0, 110],
            "ticksuffix": "%",
            "dtick": 20,
            "showgrid": False,
            "zeroline": False,
            "tickfont": {"size": 9, "color": "#dc3d56"},
        },
        shapes=[
            {
                "type": "line",
                "xref": "paper",
                "x0": 0,
                "x1": 1,
                "yref": "y2",
                "y0": 80,
                "y1": 80,
                "line": {
                    "color": "#c18100",
                    "width": 1.5,
                    "dash": "dot",
                },
            }
        ],
        annotations=[
            {
                "xref": "paper",
                "x": 1,
                "xanchor": "right",
                "yref": "y2",
                "y": 80,
                "yanchor": "bottom",
                "text": "80% reference",
                "showarrow": False,
                "font": {"size": 8, "color": "#a66b00"},
            }
        ],
        font={"family": "Segoe UI", "color": "#000000"},
    )

    return figure


def create_pareto_card(source_kpi_name, causes):
    return html.Div(
        [
            html.Div(
                [
                    html.Div(
                        [
                            html.P("PARETO ANALYSIS"),
                            html.H4(source_kpi_name),
                        ]
                    ),
                    html.Span(
                        f"{len(causes)} causes",
                        className="pareto-cause-count",
                    ),
                ],
                className="pareto-card-heading",
            ),
            dcc.Graph(
                figure=create_pareto_chart(causes),
                config={
                    "displayModeBar": False,
                    "responsive": True,
                    "scrollZoom": False,
                },
                responsive=True,
                className="pareto-chart",
                style={
                    "width": "100%",
                    "height": "330px",
                    "minHeight": "330px",
                },
            ),
        ],
        className="pareto-chart-card",
    )


def get_strategic_executive_insights(selected_month):
    strat_df = get_active_strat_rca()
    selected_month = pd.Timestamp(selected_month)
    if strat_df.empty:
        mpr_data = get_active_mpr_data()
        current = mpr_data[mpr_data["month"].eq(selected_month)].copy()
        highlights = current.loc[
            (current["status"] == "Met") & (current["previous_status"] == "Not Met"),
            "kpi_name",
        ].drop_duplicates().tolist()
        if not highlights:
            highlights = current.loc[current["is_improved"], "kpi_name"].drop_duplicates().tolist()
        lowlights = current.loc[
            (current["status"] == "Not Met") & (current["previous_status"] == "Met"),
            "kpi_name",
        ].drop_duplicates().tolist()
        concerns = current.loc[current["is_continuous_red"], "kpi_name"].drop_duplicates().tolist()
        return (
            [f"{kpi} turned green in {selected_month.strftime('%b')}." for kpi in highlights],
            [f"{kpi} moved from green to red in {selected_month.strftime('%b')}." for kpi in lowlights],
            [f"{kpi} is continuously red." for kpi in concerns],
        )

    df = strat_df.sort_values(["kpi_name", "reporting_month"]).copy()
    df["has_data"] = df["actual"].notna()
    df["status"] = df.apply(
        lambda r: "Not Met" if r["has_data"] and r["is_red"] else ("Met" if r["has_data"] else "No Data"),
        axis=1,
    )

    kpis = df["kpi_name"].dropna().unique()

    highlight_text = []
    lowlight_text = []
    concern_text = []

    for kpi in kpis:
        kpi_df = df[df["kpi_name"] == kpi].sort_values("reporting_month").reset_index(drop=True)
        curr_idx_list = kpi_df[kpi_df["reporting_month"] == selected_month].index.tolist()
        if not curr_idx_list:
            continue
        curr_idx = curr_idx_list[0]
        curr_row = kpi_df.iloc[curr_idx]

        if curr_row["status"] == "No Data":
            continue

        curr_status = curr_row["status"]

        # History before selected month
        prev_rows = kpi_df.iloc[:curr_idx]
        prev_data_rows = prev_rows[prev_rows["status"] != "No Data"]

        # Count consecutive reds ending at selected month
        consecutive_reds = 0
        for i in range(curr_idx, -1, -1):
            if kpi_df.iloc[i]["status"] == "Not Met":
                consecutive_reds += 1
            elif kpi_df.iloc[i]["status"] == "Met":
                break

        # Count consecutive reds immediately prior to current month
        prev_consecutive_reds = 0
        if curr_status == "Met":
            for i in range(len(prev_data_rows) - 1, -1, -1):
                if prev_data_rows.iloc[i]["status"] == "Not Met":
                    prev_consecutive_reds += 1
                else:
                    break

        # Count consecutive greens immediately prior to current month
        prev_consecutive_greens = 0
        if curr_status == "Not Met":
            for i in range(len(prev_data_rows) - 1, -1, -1):
                if prev_data_rows.iloc[i]["status"] == "Met":
                    prev_consecutive_greens += 1
                else:
                    break

        prev_status = prev_data_rows.iloc[-1]["status"] if not prev_data_rows.empty else None

        # 1. Highlights:
        # - Turned green after 2+ consecutive red months
        # - Or turned green from previous red
        if curr_status == "Met" and prev_status == "Not Met":
            if prev_consecutive_reds >= 2:
                highlight_text.append(f"{kpi} turned green after {prev_consecutive_reds} consecutive red months.")
            else:
                highlight_text.append(f"{kpi} turned green in {selected_month.strftime('%b')}.")

        # 2. Lowlights:
        # - Turned red after 2+ consecutive green months
        # - Or turned red from previous green
        # - Or newly red
        if curr_status == "Not Met" and prev_status == "Met":
            if prev_consecutive_greens >= 2:
                lowlight_text.append(f"{kpi} moved from green to red after {prev_consecutive_greens} green months.")
            else:
                lowlight_text.append(f"{kpi} moved from green to red in {selected_month.strftime('%b')}.")
        elif curr_status == "Not Met" and prev_status is None:
            lowlight_text.append(f"{kpi} is below target in {selected_month.strftime('%b')}.")

        # 3. Concerns:
        # - Red for 2 or more consecutive months
        if curr_status == "Not Met" and consecutive_reds >= 2:
            concern_text.append(f"{kpi} is continuously red ({consecutive_reds} consecutive months).")

    return highlight_text, lowlight_text, concern_text


def create_modal_section_header(title, subtitle):
    return html.Div(
        [
            html.Div(
                [
                    html.H3(title, className="rca-section-heading"),
                    html.P(subtitle, className="rca-section-subheading"),
                ],
                className="rca-header-title-box",
            ),
        ],
        className="rca-modal-section-heading",
    )


def create_continuous_red_detail(function_name, selected_month):
    selected_month = pd.Timestamp(selected_month)
    selected_function_key = function_key(function_name)
    active_dmb = get_active_dmb_data()

    if function_name == "ISC & Procurement":
        current_rows = active_dmb[
            active_dmb["month"].eq(selected_month)
            & active_dmb["function"].isin(["ISC & Procurement", "ISC", "Procurement"])
        ].copy()
    else:
        current_rows = active_dmb[
            active_dmb["month"].eq(selected_month)
            & active_dmb["function"].eq(function_name)
        ].copy()
    red_rows = current_rows[
        current_rows["status"].eq("Not Met")
    ].copy()
    continuous_rows = current_rows[
        current_rows["is_continuous_red"]
    ].copy()

    rca_data = get_function_rca_data()
    if rca_data["error"]:
        return html.Div(
            [
                html.Strong("The RCA workbook could not be read."),
                html.Span(rca_data["error"]),
            ],
            className="modal-data-warning",
        )

    review_kpis = rca_data["review_kpis"]
    review_red_rows = (
        review_kpis[
            review_kpis["function_key"].eq(selected_function_key)
            & review_kpis["month"].eq(selected_month)
            & review_kpis["status"].eq("Not Met")
        ].drop_duplicates(["source_sheet", "kpi_row"])
        if not review_kpis.empty
        else pd.DataFrame()
    )

    dmb_red_rows = red_rows.drop_duplicates("kpi_name")

    if not dmb_red_rows.empty:
        red_kpis = dmb_red_rows
    elif not review_red_rows.empty:
        red_kpis = review_red_rows
    else:
        red_kpis = pd.DataFrame(columns=["kpi_name", "source_sheet"])

    if red_kpis.empty:
        return html.Div(
            [
                html.Div("0", className="modal-empty-number"),
                html.H3("No Red KPIs"),
                html.P(
                    f"This function has no KPI below target for {selected_month.strftime('%B %Y')}."
                ),
            ],
            className="modal-empty-state",
        )

    selected_kpi_names = continuous_rows["kpi_name"].drop_duplicates().tolist()

    function_causes = rca_data["causes"][
        rca_data["causes"]["function_key"].eq(selected_function_key)
    ].copy()

    function_actions = rca_data["actions"][
        rca_data["actions"]["function_key"].eq(selected_function_key)
    ].copy()

    source_sheets = list(
        dict.fromkeys(
            (red_kpis["source_sheet"].dropna().tolist() if "source_sheet" in red_kpis.columns else [])
            + function_causes["source_sheet"].tolist()
            + function_actions["source_sheet"].tolist()
        )
    )
    show_source_function = len(source_sheets) > 1

    def display_name(source_sheet, kpi_name):
        if show_source_function and source_sheet:
            return f"{sheet_function_name(source_sheet)} — {kpi_name}"
        return kpi_name

    red_kpi_names = red_kpis["kpi_name"].drop_duplicates().tolist()
    rca_kpis = (
        function_causes[["source_sheet", "kpi_name"]]
        .drop_duplicates()
        .reset_index(drop=True)
    )
    function_actions = function_actions.reset_index(drop=True)

    def rca_causes(rca_index):
        rca_kpi = rca_kpis.loc[rca_index]
        return function_causes[
            function_causes["source_sheet"].eq(rca_kpi["source_sheet"])
            & function_causes["kpi_name"].eq(rca_kpi["kpi_name"])
        ]

    def applies_to_red_kpi(source_sheet, item_name, red_kpi_name):
        # A sheet-qualified KPI such as "CTB (12 weeks)(Procurement)" only uses
        # its own sheet, and month-tagged RCA only applies to the months it names.
        qualified_sheet = kpi_sheet_qualifier(red_kpi_name, source_sheets)
        if qualified_sheet and qualified_sheet != source_sheet:
            return False
        return name_applies_to_month(item_name, selected_month)

    def best_indexes(scores):
        best_score = max(scores, default=0)
        if not best_score:
            return []
        return [index for index, score in enumerate(scores) if score == best_score]

    # Match each RCA block to the red KPI it explains. Every red KPI is matched
    # once across all of the function's sheets.
    rca_by_red_kpi = {red_index: [] for red_index in range(len(red_kpi_names))}
    for rca_index, rca_kpi in rca_kpis.iterrows():
        scores = [
            kpi_match_score(rca_kpi["kpi_name"], red_kpi_name)
            if applies_to_red_kpi(
                rca_kpi["source_sheet"], rca_kpi["kpi_name"], red_kpi_name
            )
            else 0
            for red_kpi_name in red_kpi_names
        ]
        for red_index in best_indexes(scores):
            rca_by_red_kpi[red_index].append(rca_index)

    # A breakdown RCA such as "Win rate for Z30" belongs to the red KPI whose
    # RCA "Win rate for NPIs" lists "Z30" as a cause.
    breakdown_rca_indexes = set()
    directly_matched = {
        rca_index for indexes in rca_by_red_kpi.values() for rca_index in indexes
    }
    for rca_index, rca_kpi in rca_kpis.iterrows():
        if rca_index in directly_matched:
            continue
        breakdown_tokens = kpi_name_tokens(rca_kpi["kpi_name"])
        for red_index, red_rca_indexes in rca_by_red_kpi.items():
            parent_indexes = [
                parent_index
                for parent_index in red_rca_indexes
                if parent_index in directly_matched
                and rca_kpis.at[parent_index, "source_sheet"] == rca_kpi["source_sheet"]
            ]
            for parent_index in parent_indexes:
                parent_tokens = kpi_name_tokens(rca_kpis.at[parent_index, "kpi_name"])
                extra_tokens = breakdown_tokens - parent_tokens
                cause_tokens = set().union(
                    *(
                        kpi_name_tokens(cause)
                        for cause in rca_causes(parent_index)["cause"]
                    )
                )
                if (
                    extra_tokens
                    and breakdown_tokens - extra_tokens
                    and extra_tokens <= cause_tokens
                ):
                    rca_by_red_kpi[red_index].append(rca_index)
                    breakdown_rca_indexes.add(rca_index)
                    break

    # The RCA names a red KPI was matched to also identify its actions.
    actions_by_red_kpi = {red_index: [] for red_index in range(len(red_kpi_names))}
    for action_index, action in function_actions.iterrows():
        allowed = [
            applies_to_red_kpi(
                action["source_sheet"], action["kpi_name"], red_kpi_name
            )
            for red_kpi_name in red_kpi_names
        ]
        scores = []
        for red_index, red_kpi_name in enumerate(red_kpi_names):
            aliases = [red_kpi_name] + [
                rca_kpis.at[rca_index, "kpi_name"]
                for rca_index in rca_by_red_kpi[red_index]
                if rca_kpis.at[rca_index, "source_sheet"] == action["source_sheet"]
            ]
            scores.append(
                max(kpi_match_score(action["kpi_name"], alias) for alias in aliases)
                if allowed[red_index]
                else 0
            )

        if not max(scores, default=0):
            scores = [
                1
                if allowed[red_index]
                and any(
                    rca_kpis.at[rca_index, "source_sheet"] == action["source_sheet"]
                    and root_cause_matches(
                        action["root_cause"], rca_causes(rca_index)["cause"]
                    )
                    for rca_index in rca_by_red_kpi[red_index]
                )
                else 0
                for red_index in range(len(red_kpi_names))
            ]

        for red_index in best_indexes(scores):
            actions_by_red_kpi[red_index].append(action_index)

    matched_rca_indexes = {
        rca_index for indexes in rca_by_red_kpi.values() for rca_index in indexes
    }

    total_red_causes = sum(len(rca_causes(rca_idx)) for rca_idx in matched_rca_indexes)

    red_kpi_label = f"Red KPIs in {selected_month.strftime('%B %Y')}"

    summary = html.Div(
        [
            html.Div(
                [
                    html.Strong(str(len(red_kpis))),
                    html.Span(red_kpi_label),
                ],
                className="modal-summary-card summary-card-red",
            ),
            html.Div(
                [
                    html.Strong(str(len(selected_kpi_names))),
                    html.Span("Continuous-red KPIs"),
                ],
                className="modal-summary-card summary-card-red",
            ),
            html.Div(
                [
                    html.Strong(str(total_red_causes)),
                    html.Span("Top causes"),
                ],
                className="modal-summary-card summary-card-amber",
            ),
        ],
        className="modal-summary-grid",
    )

    function_review_kpis = (
        review_kpis[review_kpis["function_key"].eq(selected_function_key)][
            ["source_sheet", "kpi_name"]
        ].drop_duplicates()
        if not review_kpis.empty
        else pd.DataFrame(columns=["source_sheet", "kpi_name"])
    )

    def home_sheet(red_kpi_name, kpi_actions):
        qualified_sheet = kpi_sheet_qualifier(red_kpi_name, source_sheets)
        if qualified_sheet:
            return qualified_sheet
        if not kpi_actions.empty:
            return kpi_actions["source_sheet"].iloc[0]
        review_indexes = best_indexes(
            [
                kpi_match_score(red_kpi_name, review_name)
                for review_name in function_review_kpis["kpi_name"]
            ]
        )
        if review_indexes:
            return function_review_kpis["source_sheet"].iloc[review_indexes[0]]
        return ""

    red_kpi_tables = []
    modal_kpi_options = []

    def resolve_kpi_target_actual(title_name, red_name=None):
        candidates = []
        primary = red_name or title_name
        if primary:
            candidates.append(primary.split("—")[-1].strip() if "—" in primary else primary.strip())
        secondary = title_name if primary == red_name else red_name
        if secondary and secondary != primary:
            candidates.append(secondary.split("—")[-1].strip() if "—" in secondary else secondary.strip())

        # Pass 1: Exact match on candidate names
        for cand in candidates:
            norm_name = function_key(cand)
            if not current_rows.empty:
                for _, r in current_rows.iterrows():
                    rk_name = str(r.get("kpi_name", "")).strip()
                    if function_key(rk_name) == norm_name:
                        t_val = r.get("Target")
                        a_val = r.get("Actual")
                        u_val = r.get("units", "")
                        m_val = r.get("metric_nature", "")
                        return format_kpi_metric_display(t_val, u_val), format_kpi_metric_display(a_val, u_val), m_val
            if not red_kpis.empty:
                for _, r in red_kpis.iterrows():
                    rk_name = str(r.get("kpi_name", "")).strip()
                    if function_key(rk_name) == norm_name:
                        t_val = r.get("Target") if "Target" in r else r.get("target")
                        a_val = r.get("Actual") if "Actual" in r else r.get("actual")
                        u_val = r.get("units", "") if "units" in r else ""
                        m_val = r.get("metric_nature", "") if "metric_nature" in r else ""
                        return format_kpi_metric_display(t_val, u_val), format_kpi_metric_display(a_val, u_val), m_val
            if not review_kpis.empty:
                rev_match = review_kpis[
                    review_kpis["function_key"].eq(selected_function_key)
                    & review_kpis["month"].eq(selected_month)
                ]
                for _, r in rev_match.iterrows():
                    rk_name = str(r.get("kpi_name", "")).strip()
                    if function_key(rk_name) == norm_name:
                        t_val = r.get("target")
                        a_val = r.get("actual")
                        u_val = r.get("units", "") if "units" in r else ""
                        m_val = r.get("metric_nature", "") if "metric_nature" in r else ""
                        return format_kpi_metric_display(t_val, u_val), format_kpi_metric_display(a_val, u_val), m_val

        # Pass 2: Fuzzy match
        for cand in candidates:
            if not current_rows.empty:
                for _, r in current_rows.iterrows():
                    rk_name = str(r.get("kpi_name", "")).strip()
                    if kpi_match_score(rk_name, cand) >= 0.8:
                        t_val = r.get("Target")
                        a_val = r.get("Actual")
                        u_val = r.get("units", "")
                        m_val = r.get("metric_nature", "")
                        return format_kpi_metric_display(t_val, u_val), format_kpi_metric_display(a_val, u_val), m_val
        return None, None, None

    def add_card(
        title,
        causes,
        actions,
        red_kpi_name=None,
        is_red=True,
        missing_action_text=None,
    ):
        card_id = slug_kpi_id(f"modal-card-{red_kpi_name or title}")
        t_disp, a_disp, m_nature = resolve_kpi_target_actual(title, red_kpi_name)
        card = create_kpi_rca_card(
            title,
            causes,
            actions,
            red_kpi_name=red_kpi_name,
            card_id=card_id,
            is_red=is_red,
            missing_action_text=missing_action_text,
            target_val=t_disp,
            actual_val=a_disp,
            reporting_month=selected_month.strftime("%B %Y"),
            metric_nature=m_nature,
        )
        red_kpi_tables.append(card)
        modal_kpi_options.append({"label": red_kpi_name or title, "value": card_id})

    for red_index, red_kpi_name in enumerate(red_kpi_names):
        kpi_actions = function_actions.loc[actions_by_red_kpi[red_index]]
        red_rca_indexes = rca_by_red_kpi[red_index]

        if not red_rca_indexes:
            add_card(
                display_name(home_sheet(red_kpi_name, kpi_actions), red_kpi_name),
                function_causes.iloc[0:0],
                kpi_actions,
            )
            continue

        rca_sheets = {rca_kpis.at[index, "source_sheet"] for index in red_rca_indexes}
        for position, rca_index in enumerate(red_rca_indexes):
            rca_kpi = rca_kpis.loc[rca_index]
            title = display_name(rca_kpi["source_sheet"], rca_kpi["kpi_name"])
            if rca_index in breakdown_rca_indexes:
                # The parent card already lists this red KPI's actions.
                add_card(
                    title,
                    rca_causes(rca_index),
                    kpi_actions.iloc[0:0],
                    missing_action_text=f"Actions are listed under {display_name(rca_kpi['source_sheet'], red_kpi_name)}",
                )
                continue
            # Actions from a sheet without an RCA card go on the first card.
            card_actions = kpi_actions[
                kpi_actions["source_sheet"].eq(rca_kpi["source_sheet"])
                | (
                    (position == 0)
                    & ~kpi_actions["source_sheet"].isin(rca_sheets)
                )
            ]
            names_differ = function_key(red_kpi_name) != function_key(
                rca_kpi["kpi_name"]
            )
            add_card(
                title,
                rca_causes(rca_index),
                card_actions,
                red_kpi_name=(
                    display_name(rca_kpi["source_sheet"], red_kpi_name)
                    if names_differ
                    else None
                ),
            )

    unique_modal_options = []
    seen_vals = set()
    for opt in modal_kpi_options:
        if opt["value"] not in seen_vals:
            seen_vals.add(opt["value"])
            unique_modal_options.append(opt)

    pill_box = None
    if unique_modal_options:
        pill_box = html.Div(
            [
                html.Span("Quick Jump to KPI:", className="rca-pill-label"),
                html.Div(
                    [
                        html.Button(
                            opt["label"],
                            id={"type": "modal-rca-kpi-pill", "target": opt["value"]},
                            className="modal-rca-kpi-pill",
                            n_clicks=0,
                        )
                        for opt in unique_modal_options
                    ],
                    className="modal-rca-pills-row",
                ),
            ],
            className="modal-rca-pill-box",
        )

    section_children = [
        create_modal_section_header(
            "Root Cause Analysis & Actions on Red KPIs",
            (
                "Red KPIs with their top causes, percentage "
                "contribution to the KPI gap and open corrective "
                "actions"
            ),
        ),
    ]
    if pill_box:
        section_children.append(pill_box)
    section_children.append(
        html.Div(
            red_kpi_tables,
            className="cause-table-stack",
        )
    )

    return html.Div(
        [
            summary,
            html.Section(
                section_children,
                id="modal-root-causes",
                className="rca-workspace-section",
            ),
        ],
        className="rca-fullscreen-content",
    )


# =========================================================
# STRATEGIC-IMPERATIVE PERFORMANCE CHART
# =========================================================

def create_imperative_chart(current_data):
    if current_data.empty:
        empty_figure = go.Figure()
        empty_figure.add_annotation(
            text="No strategic-imperative data available.",
            x=0.5,
            y=0.5,
            xref="paper",
            yref="paper",
            showarrow=False,
            font={"size": 13, "color": "#8298a9"},
        )
        empty_figure.update_layout(
            height=245,
            margin={"l": 20, "r": 20, "t": 20, "b": 20},
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="rgba(0,0,0,0)",
            xaxis={"visible": False},
            yaxis={"visible": False},
        )
        return empty_figure

    data = current_data.copy()
    if "kpi_name" in data.columns and "strategic_imperative" in data.columns:
        data.loc[
            data["kpi_name"].str.strip().str.lower().str.startswith("win rate"),
            "strategic_imperative",
        ] = "Drive growth through Commercial Excellence"

    total_df = (
        data.groupby(
            "strategic_imperative",
            sort=False,
            as_index=False,
        )
        .agg(total=("kpi_name", "count"))
    )

    met_df = (
        data[data["is_met"].eq(True)]
        .groupby(
            "strategic_imperative",
            sort=False,
            as_index=False,
        )
        .agg(met=("kpi_name", "count"))
    )

    summary = total_df.merge(met_df, on="strategic_imperative", how="left")
    summary["met"] = summary["met"].fillna(0).astype(int)
    summary["percentage"] = summary.apply(
        lambda r: (r["met"] / r["total"] * 100) if r["total"] > 0 else 0.0,
        axis=1,
    )

    display_order = [
        "Customer Focus",
        "Drive growth through Commercial Excellence",
        "Environmental, Social & Governance",
        "Improve Deliverability & Profitability",
        "Patient Safety & Quality",
        "Roadmap competitiveness & Innovation agility",
    ]
    summary["sort_order"] = summary["strategic_imperative"].map(
        lambda x: display_order.index(x) if x in display_order else 99
    )
    summary = summary.sort_values("sort_order").reset_index(drop=True)

    compact_labels = {
        "Customer Focus": "Customer<br>Focus",
        "Drive growth through Commercial Excellence": (
            "Commercial<br>Excellence"
        ),
        "Improve Deliverability & Profitability": (
            "Deliverability &<br>Profitability"
        ),
        "Patient Safety & Quality": "Patient Safety<br>& Quality",
        "Roadmap competitiveness & Innovation agility": (
            "Roadmap &<br>Innovation"
        ),
        "Environmental, Social & Governance": "Environmental,<br>Social & Governance",
        "Environmental, Social & Governance (ESG)": "Environmental,<br>Social & Governance",
    }

    summary["chart_label"] = summary["strategic_imperative"].map(
        lambda value: compact_labels.get(value, value)
    )

    custom_data = summary[
        ["strategic_imperative", "met", "total"]
    ].to_numpy()

    figure = go.Figure(
        go.Bar(
            x=summary["chart_label"],
            y=summary["percentage"],
            customdata=custom_data,
            marker={
                "color": [
                    "#168b69" if percentage >= GAUGE_TARGET else "#dc3d56"
                    for percentage in summary["percentage"]
                ],
                "line": {"color": "#ffffff", "width": 1},
            },
            text=summary["percentage"].map(format_percentage),
            textposition="outside",
            textfont={"size": 12, "color": "#000000"},
            cliponaxis=False,
            hovertemplate=(
                "<b>%{customdata[0]}</b><br>"
                "KPIs met: %{customdata[1]} of %{customdata[2]}"
                "<extra></extra>"
            ),
        )
    )

    figure.update_layout(
        height=245,
        margin={"l": 42, "r": 14, "t": 25, "b": 66},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        bargap=0.34,
        hovermode="closest",
        hoverlabel={
            "bgcolor": "#ffffff",
            "bordercolor": "#b8c9d6",
            "font": {
                "family": "Segoe UI",
                "size": 13,
                "color": "#000000",
            },
            "align": "left",
        },
        dragmode=False,
        showlegend=False,
        uniformtext={"minsize": 9, "mode": "show"},
        xaxis={
            "showgrid": False,
            "zeroline": False,
            "fixedrange": True,
            "tickfont": {"size": 10, "color": "#000000"},
            "automargin": True,
        },
        yaxis={
            "range": [0, 110],
            "tickmode": "array",
            "tickvals": [0, 25, 50, 75, 100],
            "ticktext": ["0%", "25%", "50%", "75%", "100%"],
            "gridcolor": "#dce6ed",
            "gridwidth": 1,
            "zeroline": False,
            "fixedrange": True,
            "tickfont": {"size": 10, "color": "#000000"},
        },
        font={"family": "Segoe UI", "color": "#000000"},
    )

    return figure


# =========================================================
# MONTHLY TREND CHART
# =========================================================

def create_trend_chart(selected_month):
    active_mpr = get_active_mpr_data()
    selected_month = pd.Timestamp(selected_month)
    month_data = active_mpr[active_mpr["month"] <= selected_month].copy()

    if month_data.empty:
        return go.Figure()

    months_with_data = (
        month_data[month_data["Actual"].notna()]["month"].unique()
    )
    if len(months_with_data) == 0:
        return go.Figure()

    month_data = month_data[month_data["month"].isin(months_with_data)].copy()

    # Re-calculate using user's formula:
    # KPIs met -> total kpi met / total critical * 100
    # KPIs improved -> total kpis improved / total critical * 100
    trend = (
        month_data.groupby("month", as_index=False)
        .agg(
            total_critical=("kpi_name", "count"),
            met=("is_met", "sum"),
            improved=("is_improved", "sum"),
        )
    )

    trend["met_percentage"] = trend.apply(
        lambda r: (r["met"] / r["total_critical"] * 100) if r["total_critical"] > 0 else 0.0,
        axis=1,
    )
    trend["improved_percentage"] = trend.apply(
        lambda r: (r["improved"] / r["total_critical"] * 100) if r["total_critical"] > 0 else 0.0,
        axis=1,
    )

    figure = go.Figure()

    custom_met = trend[["met", "total_critical"]].to_numpy()
    custom_imp = trend[["improved", "total_critical"]].to_numpy()

    figure.add_trace(
        go.Scatter(
            x=trend["month"],
            y=trend["met_percentage"],
            customdata=custom_met,
            mode="lines+markers",
            name="KPIs met",
            line={"color": "#168b69", "width": 3},
            marker={
                "size": 7,
                "color": "#ffffff",
                "line": {"color": "#168b69", "width": 2},
            },
            hovertemplate=(
                "%{x|%B %Y}<br>KPIs met: %{customdata[0]} of %{customdata[1]} (%{y:.1f}%)<extra></extra>"
            ),
        )
    )

    figure.add_trace(
        go.Scatter(
            x=trend["month"],
            y=trend["improved_percentage"],
            customdata=custom_imp,
            mode="lines+markers",
            name="KPIs improved",
            line={"color": "#c18100", "width": 3},
            marker={
                "size": 7,
                "color": "#ffffff",
                "line": {"color": "#c18100", "width": 2},
            },
            hovertemplate=(
                "%{x|%B %Y}<br>KPIs improved: %{customdata[0]} of %{customdata[1]} (%{y:.1f}%)<extra></extra>"
            ),
        )
    )

    selected_rows = trend[trend["month"].eq(selected_month)]

    if not selected_rows.empty:
        selected_result = selected_rows.iloc[0]

        figure.add_vline(
            x=selected_month,
            line_width=2,
            line_dash="dot",
            line_color="#0877b9",
        )

        figure.add_annotation(
            x=selected_month,
            y=selected_result["met_percentage"],
            text=f"<b>{format_percentage(selected_result['met_percentage'])}</b>",
            showarrow=False,
            xshift=30,
            yshift=10,
            font={"size": 13, "color": "#168b69"},
        )

        figure.add_annotation(
            x=selected_month,
            y=selected_result["improved_percentage"],
            text=f"<b>{format_percentage(selected_result['improved_percentage'])}</b>",
            showarrow=False,
            xshift=30,
            yshift=-12,
            font={"size": 13, "color": "#c18100"},
        )

    figure.update_layout(
        height=245,
        margin={"l": 45, "r": 30, "t": 30, "b": 35},
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        hovermode="x unified",
        legend={
            "orientation": "h",
            "x": 0.62,
            "y": 1.18,
            "font": {"size": 12, "color": "#000000"},
        },
        xaxis={
            "tickformat": "%b",
            "gridcolor": "#dce6ed",
            "zeroline": False,
            "tickfont": {"color": "#000000"},
        },
        yaxis={
            "range": [0, 105],
            "ticksuffix": "%",
            "gridcolor": "#dce6ed",
            "zeroline": False,
            "dtick": 25,
            "tickfont": {"color": "#000000"},
        },
        font={
            "family": "Segoe UI",
            "size": 12,
            "color": "#000000",
        },
    )

    return figure


# =========================================================
# ULTRA-FAST IN-MEMORY CACHING & PRECOMPUTATION
# =========================================================

_memo_cache = {}
_cached_signature = None


def clear_dmb_cache():
    global _memo_cache, _cached_signature
    _memo_cache.clear()
    _cached_signature = None


def get_current_data_sig():
    try:
        return data_folder_signature()
    except Exception:
        return None


def memoize_by_data_signature(func):
    def wrapper(*args):
        global _cached_signature, _memo_cache
        current_sig = get_current_data_sig()
        if _cached_signature != current_sig:
            _memo_cache.clear()
            _cached_signature = current_sig

        cache_key = (func.__name__, args)
        if cache_key in _memo_cache:
            return _memo_cache[cache_key]

        result = func(*args)
        _memo_cache[cache_key] = result
        return result

    return wrapper


@memoize_by_data_signature
def get_mpr_dashboard_content(month_value):
    mpr_data = get_active_mpr_data()
    selected_month = pd.Timestamp(month_value)
    current = mpr_data[mpr_data["month"].eq(selected_month)].copy()
    summary = get_month_summary(mpr_data, selected_month)

    highlight_text, lowlight_text, concern_text = get_strategic_executive_insights(selected_month)

    # All 3 frequencies (monthly, quarterly, half-yearly) for the gauge
    gauge_percentage = summary["performance_percentage"]

    return (
        summary["total_kpis"],
        summary["met"],
        summary["not_met"],
        summary["improved"],
        summary["neither"],
        create_gauge(gauge_percentage),
        create_imperative_chart(current),
        selected_month.strftime("%B %Y"),
        create_trend_chart(selected_month),
        selected_month.strftime("%B %Y"),
        insight_list(
            highlight_text,
            "No positive movement identified.",
        ),
        len(highlight_text),
        insight_list(
            lowlight_text,
            "No negative movement identified.",
        ),
        len(lowlight_text),
        insight_list(
            concern_text,
            "No continuous-red KPI identified.",
        ),
        len(concern_text),
    )


@memoize_by_data_signature
def get_dmb_function_cards_content(month_value):
    dmb_data = get_active_dmb_data()
    selected_month = pd.Timestamp(month_value)
    current = dmb_data[dmb_data["month"].eq(selected_month)].copy()

    ordered_functions = [
        f for f in FUNCTION_ORDER
        if str(f).strip().lower() not in EXCLUDED_DMB_FUNCTIONS
    ]

    if not ordered_functions:
        return html.P(
            "No function-wise KPI data is available for this month.",
            className="function-empty-message",
        )

    return [
        create_function_card(function_name, current)
        for function_name in ordered_functions
    ]


@memoize_by_data_signature
def get_rca_table_content(month_value):
    selected_month = pd.Timestamp(month_value)
    return (
        create_rca_table(selected_month),
        selected_month.strftime("%B %Y"),
    )


def create_kpi_trend_chart(kpi_df, kpi_name, unit_str="%", metric_nature="Higher the better"):
    if kpi_df.empty:
        fig = go.Figure()
        fig.add_annotation(
            text="No trend data available for this KPI.",
            xref="paper",
            yref="paper",
            x=0.5,
            y=0.5,
            showarrow=False,
            font=dict(size=14, color="#64748b", family="Segoe UI"),
        )
        fig.update_layout(
            height=370,
            paper_bgcolor="rgba(0,0,0,0)",
            plot_bgcolor="#ffffff",
        )
        return fig

    months_labels = [r["month"].strftime("%b %Y") for _, r in kpi_df.iterrows()]
    targets = []
    actuals = []
    marker_colors = []
    target_texts = []
    actual_texts = []
    is_met_list = []

    valid_targets = kpi_df["Target"].dropna()
    valid_actuals = kpi_df["Actual"].dropna()

    is_pct = (
        unit_str == "%"
        or ("%" in str(kpi_df.get("units", "").iloc[0] if not kpi_df.empty and "units" in kpi_df.columns else ""))
        or (not valid_targets.empty and 0 < float(valid_targets.max()) <= 1.0)
        or (not valid_actuals.empty and 0 < float(valid_actuals.max()) <= 1.0)
    )

    lower_is_better = "lower" in str(metric_nature).lower()

    for _, r in kpi_df.iterrows():
        t_raw = r["Target"]
        a_raw = r["Actual"]
        st = r.get("status", "")

        t_val = float(t_raw) if pd.notna(t_raw) else None
        a_val = float(a_raw) if pd.notna(a_raw) else None

        if is_pct and t_val is not None and t_val <= 1.0:
            t_disp_num = round(t_val * 100, 1)
        else:
            t_disp_num = round(t_val, 2) if t_val is not None else None

        if is_pct and a_val is not None and a_val <= 1.0:
            a_disp_num = round(a_val * 100, 1)
        else:
            a_disp_num = round(a_val, 2) if a_val is not None else None

        targets.append(t_disp_num)
        actuals.append(a_disp_num)

        if t_disp_num is not None:
            t_str = f"{t_disp_num:.1f}%" if is_pct else (f"{int(t_disp_num)}" if t_disp_num == int(t_disp_num) else f"{t_disp_num:.1f}")
            target_texts.append(t_str)
        else:
            target_texts.append(None)

        if a_disp_num is not None:
            a_str = f"{a_disp_num:.1f}%" if is_pct else (f"{int(a_disp_num)}" if a_disp_num == int(a_disp_num) else f"{a_disp_num:.1f}")
            actual_texts.append(a_str)

            # Polarity coloring
            if st == "Met":
                is_met_point = True
            elif st == "Not Met":
                is_met_point = False
            elif t_disp_num is not None:
                is_met_point = (a_disp_num <= t_disp_num) if lower_is_better else (a_disp_num >= t_disp_num)
            else:
                is_met_point = True

            is_met_list.append(is_met_point)
            marker_colors.append("#168b69" if is_met_point else "#dc3d56")
        else:
            actual_texts.append(None)
            is_met_list.append(None)
            marker_colors.append("#94a3b8")

    fig = go.Figure()

    # Target Line - Exact Philips Blue (#0877b9)
    fig.add_trace(
        go.Scatter(
            x=months_labels,
            y=targets,
            mode="lines+markers",
            name="Target",
            line=dict(color="#0877b9", width=3, dash="solid"),
            marker=dict(size=9, color="#0877b9", line=dict(color="#ffffff", width=2)),
            hovertemplate="<b>%{x}</b><br>🎯 Target: <b>%{y}" + ("%" if is_pct else "") + "</b><extra></extra>",
        )
    )

    # Add colored segments for Actual Line (Red for dips / Not Met, Green for Met)
    for i in range(len(months_labels) - 1):
        a1 = actuals[i]
        a2 = actuals[i + 1]
        if a1 is not None and a2 is not None:
            m1_met = is_met_list[i]
            m2_met = is_met_list[i + 1]
            seg_is_red = (m2_met is False) or (m1_met is False and m2_met is not True)
            seg_color = "#dc3d56" if seg_is_red else "#168b69"
            fig.add_trace(
                go.Scatter(
                    x=[months_labels[i], months_labels[i + 1]],
                    y=[a1, a2],
                    mode="lines",
                    showlegend=False,
                    line=dict(color=seg_color, width=3.2),
                    hoverinfo="skip",
                )
            )

    # Split actual points by Met / Not Met so both green and red dots display in legend
    met_x = [months_labels[i] for i, m in enumerate(is_met_list) if m is True and actuals[i] is not None]
    met_y = [actuals[i] for i, m in enumerate(is_met_list) if m is True and actuals[i] is not None]

    not_met_x = [months_labels[i] for i, m in enumerate(is_met_list) if m is False and actuals[i] is not None]
    not_met_y = [actuals[i] for i, m in enumerate(is_met_list) if m is False and actuals[i] is not None]

    # Actual (Met) - Green dot
    fig.add_trace(
        go.Scatter(
            x=met_x if met_x else [None],
            y=met_y if met_y else [None],
            mode="markers",
            name="Actual (Met)",
            marker=dict(
                size=12,
                color="#168b69",
                line=dict(color="#ffffff", width=2.5),
            ),
            hovertemplate="<b>%{x}</b><br>🟢 Actual (Met): <b>%{y}" + ("%" if is_pct else "") + "</b><extra></extra>",
        )
    )

    # Actual (Not Met) - Red dot
    if not_met_x or not met_x:
        fig.add_trace(
            go.Scatter(
                x=not_met_x if not_met_x else [None],
                y=not_met_y if not_met_y else [None],
                mode="markers",
                name="Actual (Not Met)",
                marker=dict(
                    size=12,
                    color="#dc3d56",
                    line=dict(color="#ffffff", width=2.5),
                ),
                hovertemplate="<b>%{x}</b><br>🔴 Actual (Not Met): <b>%{y}" + ("%" if is_pct else "") + "</b><extra></extra>",
            )
        )

    # Calculate smart non-overlapping callouts
    valid_nums = [v for v in targets + actuals if v is not None]
    if valid_nums:
        y_min = min(valid_nums)
        y_max = max(valid_nums)
        span = max(y_max - y_min, 8 if is_pct else 2)
        if is_pct:
            y_range = [min(-8.0, y_min - span * 0.18), max(112.0, y_max + span * 0.18)]
        else:
            y_range = [y_min - span * 0.22, y_max + span * 0.22]
    else:
        y_range = [-10, 115]

    # Clean Data Numbers (NO BOX BORDERS, crisp contrast typography)
    for i, m_label in enumerate(months_labels):
        t_val = targets[i]
        a_val = actuals[i]
        t_txt = target_texts[i]
        a_txt = actual_texts[i]
        pt_met = is_met_list[i] if i < len(is_met_list) else None

        if t_val is not None and a_val is not None:
            # Both Target and Actual exist
            if abs(t_val - a_val) < (span * 0.08):
                if a_val <= (y_min + span * 0.1):  # near bottom
                    a_yshift = 18
                    t_yshift = 40
                else:
                    a_yshift = 18
                    t_yshift = -20
            elif a_val >= t_val:
                a_yshift = 18
                t_yshift = -18
            else:
                t_yshift = 18
                a_yshift = -18
                if a_val <= (y_min + span * 0.06):
                    a_yshift = 18
                    t_yshift = 40

            # Target Number (Exact Philips Blue #0877b9, clean floating text)
            fig.add_annotation(
                x=m_label,
                y=t_val,
                text=f"<b>{t_txt}</b>",
                showarrow=False,
                yshift=t_yshift,
                bgcolor="rgba(255, 255, 255, 0.85)",
                bordercolor="rgba(0,0,0,0)",
                borderwidth=0,
                font=dict(size=13, color="#0877b9", family="Segoe UI", weight=750),
            )

            # Actual Number (Green for Met, Red for Not Met, clean floating text)
            act_color = "#168b69" if pt_met else "#dc3d56"
            fig.add_annotation(
                x=m_label,
                y=a_val,
                text=f"<b>{a_txt}</b>",
                showarrow=False,
                yshift=a_yshift,
                bgcolor="rgba(255, 255, 255, 0.85)",
                bordercolor="rgba(0,0,0,0)",
                borderwidth=0,
                font=dict(size=13, color=act_color, family="Segoe UI", weight=800),
            )

        elif t_val is not None:
            fig.add_annotation(
                x=m_label,
                y=t_val,
                text=f"<b>{t_txt}</b>",
                showarrow=False,
                yshift=18,
                bgcolor="rgba(255, 255, 255, 0.85)",
                bordercolor="rgba(0,0,0,0)",
                borderwidth=0,
                font=dict(size=13, color="#0877b9", family="Segoe UI", weight=750),
            )

        elif a_val is not None:
            act_color = "#168b69" if pt_met else "#dc3d56"
            a_yshift = 18 if a_val <= (y_min + span * 0.1) else -18
            fig.add_annotation(
                x=m_label,
                y=a_val,
                text=f"<b>{a_txt}</b>",
                showarrow=False,
                yshift=a_yshift,
                bgcolor="rgba(255, 255, 255, 0.85)",
                bordercolor="rgba(0,0,0,0)",
                borderwidth=0,
                font=dict(size=13, color=act_color, family="Segoe UI", weight=800),
            )

    fig.update_layout(
        autosize=True,
        height=390,
        margin=dict(l=50, r=50, t=55, b=45),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="#ffffff",
        legend=dict(
            orientation="h",
            x=0.5,
            xanchor="center",
            y=1.14,
            font=dict(size=13.5, family="Segoe UI", color="#082d4c", weight=750),
            bgcolor="rgba(245, 249, 253, 0.96)",
            bordercolor="#c9ddec",
            borderwidth=1,
            itemclick=False,
            itemdoubleclick=False,
        ),
        xaxis=dict(
            showgrid=False,
            tickfont=dict(size=12, color="#0f172a", family="Segoe UI", weight=600),
            linecolor="#cbd5e1",
            linewidth=1.2,
        ),
        yaxis=dict(
            range=y_range,
            ticksuffix="%" if is_pct else "",
            gridcolor="#e9f1f6",
            gridwidth=1,
            zeroline=True,
            zerolinecolor="#cbd5e1",
            zerolinewidth=1,
            tickfont=dict(size=11.5, color="#475569", family="Segoe UI", weight=600),
            linecolor="#cbd5e1",
            linewidth=1.2,
        ),
        font=dict(family="Segoe UI"),
    )
    return fig


def create_function_kpis_table_detail(function_name, selected_month):
    selected_month = pd.Timestamp(selected_month)
    active_dmb = get_active_dmb_data()

    if function_name == "ISC & Procurement":
        fn_rows = active_dmb[
            active_dmb["function"].isin(["ISC & Procurement", "ISC", "Procurement"])
        ].copy()
    else:
        fn_rows = active_dmb[
            active_dmb["function"].eq(function_name)
        ].copy()

    if fn_rows.empty:
        return html.Div(
            [
                html.H3("No KPI Data Found"),
                html.P(f"No KPI records found for {function_name}."),
            ],
            className="modal-empty-state",
        )

    kpi_names = fn_rows["kpi_name"].dropna().unique().tolist()
    if not kpi_names:
        return html.Div(
            [
                html.H3("No KPIs Available"),
                html.P(f"No KPIs available for {function_name}."),
            ],
            className="modal-empty-state",
        )

    # Active completed months up to current year/data (e.g. Jan to Aug or full months)
    all_months = sorted(fn_rows["month"].dropna().unique())
    if not all_months:
        all_months = [selected_month]

    # Calculate summary metrics for selected month
    sel_rows = fn_rows[fn_rows["month"].eq(selected_month)].copy()
    total_kpis = len(kpi_names)
    met_count = int(sel_rows["is_met"].sum()) if not sel_rows.empty and "is_met" in sel_rows.columns else 0
    not_met_count = int(sel_rows["status"].eq("Not Met").sum()) if not sel_rows.empty and "status" in sel_rows.columns else 0
    filled_count = met_count + not_met_count
    met_rate = round(met_count / filled_count * 100, 1) if filled_count > 0 else 0.0

    # 1. Executive Metric Stat Cards (Large, prominent, elevated)
    summary_bar = html.Div(
        [
            html.Div(
                [
                    html.Div(
                        [
                            html.Span("TOTAL KPIS", className="fn-stat-label"),
                            html.Span(f"{function_name} Scope", className="fn-stat-sub"),
                        ],
                        className="fn-stat-info",
                    ),
                    html.Strong(str(total_kpis), className="fn-stat-value fn-stat-val-total"),
                ],
                className="fn-stat-card fn-stat-total",
            ),
            html.Div(
                [
                    html.Div(
                        [
                            html.Span(f"MET IN {selected_month.strftime('%b').upper()}", className="fn-stat-label"),
                            html.Span("On-Target Benchmark", className="fn-stat-sub"),
                        ],
                        className="fn-stat-info",
                    ),
                    html.Strong(str(met_count), className="fn-stat-value fn-stat-val-met"),
                ],
                className="fn-stat-card fn-stat-met",
            ),
            html.Div(
                [
                    html.Div(
                        [
                            html.Span(f"NOT MET IN {selected_month.strftime('%b').upper()}", className="fn-stat-label"),
                            html.Span("Performance Gap / Red", className="fn-stat-sub"),
                        ],
                        className="fn-stat-info",
                    ),
                    html.Strong(str(not_met_count), className="fn-stat-value fn-stat-val-not-met"),
                ],
                className="fn-stat-card fn-stat-not-met",
            ),
            html.Div(
                [
                    html.Div(
                        [
                            html.Span("MET RATE", className="fn-stat-label"),
                            html.Span(f"{selected_month.strftime('%B %Y')} Compliance", className="fn-stat-sub"),
                        ],
                        className="fn-stat-info",
                    ),
                    html.Strong(f"{met_rate}%", className="fn-stat-value fn-stat-val-rate"),
                ],
                className="fn-stat-card fn-stat-rate",
            ),
        ],
        className="fn-stat-cards-grid",
    )

    # Format helpers
    def format_val_str(val, unit, is_pct_hint=False):
        if val is None or pd.isna(val) or str(val).strip() in {"", "nan", "None", "—"}:
            return "—"
        try:
            num = float(val)
            if unit == "%" or is_pct_hint or (abs(num) <= 1.5 and abs(num) > 0 and unit != "Mn"):
                return f"{num * 100:.1f}%" if abs(num) <= 1.5 else f"{num:.1f}%"
            elif unit == "Mn":
                return f"€{num:,.2f}M" if num >= 0 else f"-€{abs(num):,.2f}M"
            elif float(num).is_integer():
                return f"{int(num)}"
            else:
                return f"{num:,.2f}".rstrip("0").rstrip(".")
        except (ValueError, TypeError):
            return str(val)

    # 2. Build Individual KPI Cards with 2-Row Monthly Trend Tables
    kpi_cards = []
    for idx, k_name in enumerate(kpi_names, 1):
        k_hist = fn_rows[fn_rows["kpi_name"] == k_name].sort_values("month")
        def_val = k_hist["definition"].dropna().iloc[0] if not k_hist["definition"].dropna().empty else ""
        unit_val = k_hist["units"].dropna().iloc[0] if not k_hist["units"].dropna().empty else ""
        op_val = k_hist["operator"].dropna().iloc[0] if not k_hist["operator"].dropna().empty else ""
        aop_val = k_hist["target_aop_2026"].dropna().iloc[0] if not k_hist["target_aop_2026"].dropna().empty else ""
        nature_val = k_hist["metric_nature"].dropna().iloc[0] if not k_hist["metric_nature"].dropna().empty else ""

        is_pct = unit_val == "%" or "percent" in k_name.lower() or "%" in k_name or (isinstance(aop_val, (int, float)) and abs(aop_val) <= 1.0 and unit_val != "Mn")

        aop_fmt = format_val_str(aop_val, unit_val, is_pct)
        target_aop_display = f"{op_val} {aop_fmt}".strip() if op_val else aop_fmt
        nature_display = f"{unit_val} · {nature_val}".strip(" ·") if unit_val or nature_val else ""

        curr_record = k_hist[k_hist["month"] == selected_month]
        st_curr = curr_record["status"].iloc[0] if not curr_record.empty else "No Data"

        if st_curr == "Met":
            status_pill = html.Span("🟢 Met", className="fn-badge-met")
        elif st_curr == "Not Met":
            status_pill = html.Span("🔴 Not Met", className="fn-badge-not-met")
        else:
            status_pill = html.Span("— No Data", className="fn-badge-nodata")

        # Header for this KPI card
        meta_items = []
        if target_aop_display:
            meta_items.append(
                html.Div(
                    [
                        html.Span("Target AOP:", className="fn-kpi-meta-lbl"),
                        html.Strong(target_aop_display, className="fn-kpi-meta-val"),
                        html.Span(f"({nature_display})" if nature_display else "", className="fn-kpi-meta-nature"),
                    ],
                    className="fn-kpi-meta-pill",
                )
            )

        meta_items.append(
            html.Div(
                [
                    html.Span(f"{selected_month.strftime('%b %Y')}:", className="fn-kpi-curr-lbl"),
                    status_pill,
                ],
                className="fn-kpi-curr-status-wrap",
            )
        )

        card_header = html.Div(
            [
                html.Div(
                    [
                        html.Span(f"#{idx}", className="fn-kpi-idx-badge"),
                        html.Div(
                            [
                                html.H3(k_name, className="fn-kpi-title"),
                                html.P(def_val or f"Performance metric for {function_name}", className="fn-kpi-def"),
                            ],
                            className="fn-kpi-title-wrap",
                        ),
                    ],
                    className="fn-kpi-header-left",
                ),
                html.Div(
                    meta_items,
                    className="fn-kpi-header-right",
                ),
            ],
            className="fn-kpi-header-bar",
        )

        # Build 2-Row Monthly Trend Table
        trend_th_cells = [
            html.Th("Monthly Trend", className="fn-trend-th fn-trend-th-metric"),
        ]
        for m in all_months:
            is_cur = (m == selected_month)
            th_cls = "fn-trend-th fn-trend-th-month" + (" fn-trend-th-selected" if is_cur else "")
            trend_th_cells.append(
                html.Th(m.strftime("%b"), className=th_cls)
            )

        # Row 1: Target (Excel Blue Format)
        target_td_cells = [
            html.Td(
                html.Span("Target", className="fn-row-type-badge fn-badge-target"),
                className="fn-trend-td-type fn-td-target-label",
            ),
        ]
        for m in all_months:
            m_rec = k_hist[k_hist["month"] == m]
            m_tgt = m_rec["Target"].iloc[0] if not m_rec.empty else None
            t_str = format_val_str(m_tgt, unit_val, is_pct)
            is_cur = (m == selected_month)
            td_cls = "fn-trend-td fn-trend-tgt-cell" + (" fn-cell-selected-month" if is_cur else "")
            target_td_cells.append(
                html.Td(t_str, className=td_cls, title=f"{m.strftime('%b %Y')} Target: {t_str}")
            )

        # Row 2: Actual (Red/Green Format)
        actual_td_cells = [
            html.Td(
                html.Span("Actual", className="fn-row-type-badge fn-badge-actual"),
                className="fn-trend-td-type fn-td-actual-label",
            ),
        ]
        for m in all_months:
            m_rec = k_hist[k_hist["month"] == m]
            m_act = m_rec["Actual"].iloc[0] if not m_rec.empty else None
            m_st = m_rec["status"].iloc[0] if not m_rec.empty else "No Data"
            a_str = format_val_str(m_act, unit_val, is_pct)
            is_cur = (m == selected_month)

            if m_st == "Met":
                st_cls = "fn-act-met"
            elif m_st == "Not Met":
                st_cls = "fn-act-not-met"
            else:
                st_cls = "fn-act-nodata"

            td_cls = f"fn-trend-td fn-trend-act-cell {st_cls}" + (" fn-cell-selected-month" if is_cur else "")
            actual_td_cells.append(
                html.Td(a_str, className=td_cls, title=f"{m.strftime('%b %Y')} Actual: {a_str} ({m_st})")
            )

        table_el = html.Div(
            html.Table(
                [
                    html.Thead(html.Tr(trend_th_cells)),
                    html.Tbody(
                        [
                            html.Tr(target_td_cells, className="fn-trend-row-target"),
                            html.Tr(actual_td_cells, className="fn-trend-row-actual"),
                        ]
                    ),
                ],
                className="fn-kpi-trend-table",
            ),
            className="fn-trend-table-container",
        )

        kpi_cards.append(
            html.Div(
                [
                    card_header,
                    table_el,
                ],
                className="fn-kpi-card",
            )
        )

    return html.Div(
        [
            summary_bar,
            html.Div(kpi_cards, className="fn-kpi-cards-stack"),
        ],
        className="fn-modal-content-container",
    )


@memoize_by_data_signature
def get_continuous_red_modal_content(function_name, month_value):
    selected_month = pd.Timestamp(month_value)
    return (
        "continuous-red-modal",
        f"{function_name} — Root Cause Analysis",
        selected_month.strftime("%B %Y"),
        create_continuous_red_detail(
            function_name,
            selected_month,
        ),
    )


@memoize_by_data_signature
def get_function_kpis_trend_modal_content(function_name, month_value, active_kpi=None):
    selected_month = pd.Timestamp(month_value)
    return create_function_kpis_table_detail(function_name, selected_month)


def create_mpr_breakdown_pills(summary, active_filter="all"):
    pills_config = [
        ("all", "All Critical KPIs", summary.get("total_kpis", 0), "pill-all"),
        ("met", "KPIs Met", summary.get("met", 0), "pill-met"),
        ("not_met", "KPIs Not Met", summary.get("not_met", 0), "pill-not-met"),
        ("improved", "KPIs Improved", summary.get("improved", 0), "pill-improved"),
        ("neither", "Neither Met nor Improved", summary.get("neither", 0), "pill-neither"),
        ("outside", "KPIs Without Data / Outside", summary.get("outside_kpis", 0), "pill-outside"),
    ]

    pills = []
    for f_key, label, count, color_cls in pills_config:
        is_active = f_key == active_filter
        cls_name = f"mpr-modal-pill {color_cls}"
        if is_active:
            cls_name += " mpr-pill-active"
        pills.append(
            html.Button(
                [
                    html.Span(label, className="mpr-pill-label"),
                    html.Span(f" ({count})", className="mpr-pill-count"),
                ],
                id={"type": "mpr-modal-pill", "filter": f_key},
                className=cls_name,
                n_clicks=0,
            )
        )
    return pills


def get_mpr_kpis_breakdown_modal_content(month_value, filter_type="all"):
    selected_month = pd.Timestamp(month_value)
    mpr_data = get_active_mpr_data()
    current = mpr_data[mpr_data["month"].eq(selected_month)].copy()
    summary = get_month_summary(mpr_data, selected_month)

    if current.empty:
        return (
            "Overall MoS KPI Performance Details",
            selected_month.strftime("%B %Y"),
            [],
            html.Div(f"No KPI records found for {selected_month.strftime('%B %Y')}.", className="mpr-empty-state"),
        )

    # Filter data based on filter_type
    filter_type = filter_type or "all"
    if filter_type == "met":
        sub_df = current[current["is_met"]].copy()
        title_suffix = "KPIs Met"
    elif filter_type == "not_met":
        sub_df = current[
            current["status"].eq("Not Met")
            | current["Target"].isna()
            | current["Actual"].isna()
        ].copy()
        title_suffix = "KPIs Not Met"
    elif filter_type == "improved":
        sub_df = current[current["is_improved"]].copy()
        title_suffix = "KPIs Improved"
    elif filter_type == "neither":
        sub_df = current[current["is_neither"]].copy()
        title_suffix = "Neither Met nor Improved KPIs"
    elif filter_type == "outside":
        sub_df = current[current["Target"].isna() | current["Actual"].isna()].copy()
        title_suffix = "KPIs Without Data / Outside Met & Not Met"
    else:
        sub_df = current.copy()
        title_suffix = "All Critical KPIs"

    pills = create_mpr_breakdown_pills(summary, filter_type)

    if sub_df.empty:
        table_content = html.Div(
            f"No KPIs matching '{title_suffix}' for {selected_month.strftime('%B %Y')}.",
            className="mpr-empty-state",
        )
    else:
        rows = []
        for _, r in sub_df.iterrows():
            k_name = str(r.get("kpi_name", "")).strip()
            imperative = str(r.get("strategic_imperative", "")).strip()
            freq = str(r.get("frequency", "")).strip() or "Monthly"
            t_raw = r.get("Target")
            a_raw = r.get("Actual")
            u = str(r.get("units", "")).strip()

            t_disp = format_kpi_metric_display(t_raw, unit=u)
            a_disp = format_kpi_metric_display(a_raw, unit=u)

            is_met = pd.notna(r.get("is_met")) and bool(r.get("is_met"))
            actual_class = "mpr-actual-text-met" if is_met else "mpr-actual-text-red"

            rows.append(
                html.Tr(
                    [
                        html.Td(imperative, className="mpr-imperative-td"),
                        html.Td(html.Strong(k_name), className="mpr-kpi-name-td"),
                        html.Td(freq, className="mpr-freq-td"),
                        html.Td(t_disp or "—", className="mpr-target-td mpr-target-text"),
                        html.Td(a_disp or "—", className=f"mpr-actual-td {actual_class}"),
                    ]
                )
            )

        table_content = html.Table(
            [
                html.Thead(
                    html.Tr(
                        [
                            html.Th("Strategic Imperative"),
                            html.Th("KPI Name"),
                            html.Th("Frequency"),
                            html.Th("Target"),
                            html.Th("Actual"),
                        ]
                    )
                ),
                html.Tbody(rows),
            ],
            className="mpr-breakdown-table",
        )

    title = f"{title_suffix} Breakdown"
    subtitle = f"{selected_month.strftime('%B %Y')} • {len(sub_df)} {title_suffix} listed"

    return title, subtitle, pills, table_content


def warmup_cache():
    try:
        active_mpr = get_active_mpr_data()
        active_dmb = get_active_dmb_data()
        active_strat = get_active_rca_actions()
        mpr_months, default_mpr_m = get_dynamic_reporting_months(active_mpr, active_strat)
        dmb_months, default_dmb_m = get_dynamic_reporting_months(active_dmb, active_strat)
        
        # 1. Warm up default reporting months first immediately (< 50ms)
        for def_m in set([default_dmb_m, default_mpr_m]):
            m_str = def_m.strftime("%Y-%m-%d")
            get_mpr_dashboard_content(m_str)
            get_dmb_function_cards_content(m_str)
            get_rca_table_content(m_str)
            for fn in ["Quality", "Regulatory", "ISC & Procurement", "R&D", "Customer Service", "Marketing"]:
                get_continuous_red_modal_content(fn, m_str)
                get_function_kpis_trend_modal_content(fn, m_str)
                
        # 2. Warm up remaining months in background gently with small yield
        all_months = sorted(set(mpr_months + dmb_months))
        for m in all_months:
            m_str = m.strftime("%Y-%m-%d")
            time.sleep(0.08)
            get_mpr_dashboard_content(m_str)
            get_dmb_function_cards_content(m_str)
            get_rca_table_content(m_str)
            for fn in ["Quality", "Regulatory", "ISC & Procurement", "R&D", "Customer Service", "Marketing"]:
                get_continuous_red_modal_content(fn, m_str)
                get_function_kpis_trend_modal_content(fn, m_str)
    except Exception as e:
        print(f"[Cache Warmup Notice] {e}")


threading.Thread(target=warmup_cache, daemon=True).start()


# =========================================================
# DASHBOARD LAYOUT
# =========================================================

def serve_layout():
    active_mpr = get_active_mpr_data()
    active_dmb = get_active_dmb_data()
    active_strat = get_active_rca_actions()

    curr_mpr_months, curr_default_mpr_month = get_dynamic_reporting_months(active_mpr, active_strat)
    curr_dmb_months, curr_default_dmb_month = get_dynamic_reporting_months(active_dmb, active_strat)

    mpr_str = curr_default_mpr_month.strftime("%Y-%m-%d")
    dmb_str = curr_default_dmb_month.strftime("%Y-%m-%d")

    mpr_init = get_mpr_dashboard_content(mpr_str)
    dmb_init_cards = get_dmb_function_cards_content(dmb_str)
    rca_init_table, rca_init_month = get_rca_table_content(mpr_str)

    return html.Div(
        [
            html.Header(
                [
                    html.Div(
                        [
                            html.H1(
                                f"{curr_default_mpr_month.strftime('%B')} MoS DMB Performance Review - AOP Critical",
                                id="dashboard-header-title",
                            ),
                        ],
                        className="brand-text",
                    ),
                    html.Div(
                        [
                            html.Button(
                                "Download 1 Pager",
                                id="download-one-pager-button",
                                className="download-button",
                                n_clicks=0,
                                title="Download the complete dashboard as one long PNG image",
                            ),
                        ],
                        className="top-nav-actions",
                    ),
                ],
                className="top-navigation",
            ),
            dcc.Store(id="one-pager-download-state"),
            dcc.Store(id="active-rca-modal-state", storage_type="memory", data={"is_open": False, "function": None, "month": None}),
            dcc.Store(id="active-gauge-modal-state", storage_type="memory", data={"is_open": False, "function": None, "selected_kpi": None, "month": None}),
            dcc.Store(id="active-mpr-breakdown-modal-state", storage_type="memory", data={"is_open": False, "filter": "outside", "month": None}),
            dcc.Interval(id="live-sync-interval", interval=10000, n_intervals=0),
            dcc.Store(id="live-sync-state-store"),
            html.Section(
                [
                    html.Div(
                        [
                            html.H2(
                                "Highlights & Lowlights of KPIs at MoS Level"
                            ),
                            html.Span(
                                mpr_init[9],
                                id="insights-reporting-month",
                                className="insights-month",
                            ),
                        ],
                        className="insights-title-bar",
                    ),
                    html.Div(
                        [
                            insight_card(
                                "Highlights",
                                "highlights-content",
                                "highlights-count",
                                "highlight-card",
                                initial_content=mpr_init[10],
                                initial_count=mpr_init[11],
                            ),
                            insight_card(
                                "Lowlights",
                                "lowlights-content",
                                "lowlights-count",
                                "lowlight-card",
                                initial_content=mpr_init[12],
                                initial_count=mpr_init[13],
                            ),
                            insight_card(
                                "Concerns",
                                "concerns-content",
                                "concerns-count",
                                "concern-card",
                                initial_content=mpr_init[14],
                                initial_count=mpr_init[15],
                            ),
                        ],
                        className="executive-insights",
                    ),
                ],
                id="executive-insights-section",
                className="executive-insights-section",
            ),
            html.Section(
                [
                    section_header(
                        "mpr-section",
                        "Overall MoS KPI Performance",
                        None,
                        None,
                        "mpr-month-filter",
                        create_month_options(curr_mpr_months),
                        default_value=mpr_str,
                    ),
                    html.Div(
                        [
                            html.Div(
                                [
                                    kpi_card(
                                        "Critical KPIs",
                                        "mpr-total-kpis",
                                        "#082d4c",
                                        initial_value=mpr_init[0],
                                        is_clickable=True,
                                        title="Click to view all Critical KPIs table",
                                    ),
                                    kpi_card(
                                        "KPIs met",
                                        "mpr-met-kpis",
                                        "#168b69",
                                        initial_value=mpr_init[1],
                                        is_clickable=True,
                                        title="Click to view Met KPIs table",
                                    ),
                                    kpi_card(
                                        "KPIs not met",
                                        "mpr-not-met-kpis",
                                        "#dc3d56",
                                        initial_value=mpr_init[2],
                                        scroll_target="rca-section",
                                        is_clickable=True,
                                        title="Click to jump to Cause and Actions of Red KPIs at MoS Level",
                                    ),
                                ],
                                className="kpi-summary-group kpi-summary-group-3",
                            ),
                            html.Div(
                                [
                                    kpi_card(
                                        "KPIs improved",
                                        "mpr-improved-kpis",
                                        "#c18100",
                                        initial_value=mpr_init[3],
                                        is_clickable=True,
                                        title="Click to view Improved KPIs table",
                                    ),
                                    kpi_card(
                                        "Neither improved nor met",
                                        "mpr-neither-kpis",
                                        "#dc3d56",
                                        initial_value=mpr_init[4],
                                        is_clickable=True,
                                        title="Click to view Neither improved nor met KPIs table",
                                    ),
                                ],
                                className="kpi-summary-group kpi-summary-group-2",
                            ),
                        ],
                        className="kpi-summary-grid",
                    ),
                    html.Div(
                        [
                            html.Div(
                                [
                                    html.Div(
                                        [
                                            html.H3("KPIs met & improved"),
                                            html.Span(
                                                f"Threshold {GAUGE_TARGET}%",
                                                className="chart-note",
                                            ),
                                        ],
                                        className="chart-title-row",
                                    ),
                                    dcc.Graph(
                                        id="mpr-gauge",
                                        figure=mpr_init[5],
                                        config={"displayModeBar": False},
                                    ),
                                ],
                                className="mpr-chart-card gauge-chart-card",
                            ),
                            html.Div(
                                [
                                    html.Div(
                                        [
                                            html.H3(
                                                "Strategic imperative performance"
                                            ),
                                            html.Span(
                                                mpr_init[7],
                                                id="imperative-month",
                                                className="chart-note",
                                            ),
                                        ],
                                        className="chart-title-row",
                                    ),
                                    dcc.Graph(
                                        id="mpr-imperative-chart",
                                        figure=mpr_init[6],
                                        config={
                                             "displayModeBar": False,
                                             "scrollZoom": False,
                                        },
                                    ),
                                ],
                                className=(
                                    "mpr-chart-card imperative-chart-card"
                                ),
                            ),
                            html.Div(
                                [
                                    html.Div(
                                        [html.H3("Monthly Performance Trend")],
                                        className="chart-title-row",
                                    ),
                                    dcc.Graph(
                                        id="mpr-trend-chart",
                                        figure=mpr_init[8],
                                        config={"displayModeBar": False},
                                    ),
                                ],
                                className="mpr-chart-card trend-chart-card",
                            ),
                        ],
                        className="mpr-visual-grid",
                    ),
                ],
                className="dashboard-section",
            ),
            html.Section(
                [
                    section_header(
                        "dmb-section",
                        "DMB — Function-wise KPI Performance",
                        None,
                        None,
                        "dmb-month-filter",
                        create_month_options(curr_dmb_months),
                        default_value=dmb_str,
                    ),
                    html.Div(
                        [
                            html.Div(
                                [
                                    html.Span(
                                        className="legend-dot legend-green"
                                    ),
                                    html.Span("KPIs met"),
                                ],
                                className="dmb-legend-item",
                            ),
                            html.Div(
                                [
                                    html.Span(
                                        className="legend-dot legend-red"
                                    ),
                                    html.Span("KPIs not met"),
                                ],
                                className="dmb-legend-item",
                            ),
                            html.Div(
                                [
                                    html.Span(
                                        className="legend-dot legend-not-filled"
                                    ),
                                    html.Span("KPIs not filled"),
                                ],
                                className="dmb-legend-item",
                            ),
                            html.Div(
                                [
                                    html.Span(
                                        className="legend-dot legend-blue"
                                    ),
                                    html.Span(f"Performance target {GAUGE_TARGET}%"),
                                ],
                                className="dmb-legend-item",
                            ),
                        ],
                        className="dmb-legend",
                    ),
                    html.Div(
                        dmb_init_cards,
                        id="function-cards-container",
                        className="function-grid",
                    ),
                ],
                className="dashboard-section",
            ),
            html.Section(
                [
                    html.Div(
                        [
                            html.H2(
                                "Cause and Actions of Red KPIs at MoS Level"
                            ),
                            html.Span(
                                rca_init_month,
                                id="rca-reporting-month",
                                className="rca-month",
                            ),
                        ],
                        className="rca-title-bar",
                    ),
                    html.Div(
                        rca_init_table,
                        id="rca-table-container",
                    ),
                ],
                id="rca-section",
                className="rca-section",
            ),
            html.Div(
                [
                    html.Button(
                        id="continuous-red-modal-backdrop",
                        className="continuous-red-modal-backdrop",
                        n_clicks=0,
                        title="Close details",
                    ),
                    html.Div(
                        [
                            html.Div(
                                [
                                    html.Div(
                                        [
                                            html.P(
                                                "DMB ROOT-CAUSE REVIEW",
                                                className="modal-eyebrow",
                                            ),
                                            html.H2(
                                                id="continuous-red-modal-title"
                                            ),
                                            html.P(
                                                id="continuous-red-modal-month",
                                                className="modal-reporting-month",
                                            ),
                                        ]
                                    ),
                                    html.Button(
                                        [
                                            html.Span("←", className="modal-back-arrow"),
                                            html.Span("Back", className="modal-back-text"),
                                        ],
                                        id="close-continuous-red-modal",
                                        n_clicks=0,
                                        className="continuous-red-modal-close",
                                        title="Back to Dashboard",
                                    ),
                                ],
                                className="continuous-red-modal-header",
                            ),
                            html.Div(
                                id="continuous-red-modal-body",
                                className="continuous-red-modal-body",
                            ),
                            html.Div(
                                id="continuous-red-modal-scroll-dummy",
                                style={"display": "none"},
                            ),
                        ],
                        className="continuous-red-modal-dialog",
                    ),
                ],
                id="continuous-red-modal",
                className=(
                    "continuous-red-modal continuous-red-modal-hidden"
                ),
            ),
            html.Div(
                [
                    html.Button(
                        id="function-gauge-modal-backdrop",
                        className="function-gauge-modal-backdrop",
                        n_clicks=0,
                        title="Close details",
                    ),
                    html.Div(
                        [
                            html.Div(
                                [
                                    html.Div(
                                        [
                                            html.Div(
                                                html.Span("📈", className="trend-header-icon"),
                                                className="trend-header-icon-wrap",
                                            ),
                                            html.Div(
                                                [
                                                    html.H2(
                                                        id="function-gauge-modal-title"
                                                    ),
                                                    html.P(
                                                        id="function-gauge-modal-subtitle",
                                                        className="trend-header-subtitle",
                                                    ),
                                                ],
                                                className="trend-header-title-box",
                                            ),
                                        ],
                                        className="trend-header-left",
                                    ),
                                    html.Div(
                                        [
                                            html.Button(
                                                [
                                                    html.Span("←", className="modal-back-arrow"),
                                                    html.Span("Back", className="modal-back-text"),
                                                ],
                                                id="close-function-gauge-modal",
                                                n_clicks=0,
                                                className="function-gauge-modal-close",
                                                title="Back to Dashboard",
                                            ),
                                        ],
                                        className="trend-header-right",
                                    ),
                                ],
                                className="function-gauge-modal-header",
                            ),
                            html.Div(
                                id="function-gauge-modal-body",
                                className="function-gauge-modal-body",
                            ),
                        ],
                        className="function-gauge-modal-dialog",
                    ),
                ],
                id="function-gauge-modal",
                className="function-gauge-modal function-gauge-modal-hidden",
            ),
            html.Div(
                [
                    html.Button(
                        id="mpr-breakdown-modal-backdrop",
                        className="mpr-breakdown-modal-backdrop",
                        n_clicks=0,
                        title="Close details",
                    ),
                    html.Div(
                        [
                            html.Div(
                                [
                                    html.Div(
                                        [
                                            html.P(
                                                "OVERALL MOS KPI PERFORMANCE",
                                                className="mpr-breakdown-eyebrow",
                                            ),
                                            html.H2(
                                                id="mpr-breakdown-modal-title",
                                                className="mpr-breakdown-modal-title",
                                            ),
                                            html.P(
                                                id="mpr-breakdown-modal-subtitle",
                                                className="mpr-breakdown-modal-subtitle",
                                            ),
                                        ],
                                        className="mpr-breakdown-header-left",
                                    ),
                                    html.Button(
                                        [
                                            html.Span("←", className="modal-back-arrow"),
                                            html.Span("Back", className="modal-back-text"),
                                        ],
                                        id="close-mpr-breakdown-modal",
                                        n_clicks=0,
                                        className="mpr-breakdown-modal-close",
                                        title="Back to Dashboard",
                                    ),
                                ],
                                className="mpr-breakdown-modal-header",
                            ),
                            html.Div(
                                id="mpr-breakdown-modal-pills-container",
                                className="mpr-modal-pills-row",
                            ),
                            html.Div(
                                id="mpr-breakdown-modal-body",
                                className="mpr-breakdown-modal-body",
                            ),
                        ],
                        className="mpr-breakdown-modal-dialog",
                    ),
                ],
                id="mpr-breakdown-modal",
                className="mpr-breakdown-modal mpr-breakdown-modal-hidden",
            ),
        ],
        className="page-shell",
    )


app.layout = serve_layout


# =========================================================
# BACKGROUND LIVE AUTO-SYNC (CONTINUOUS REAL-TIME CYCLE)
# =========================================================

@app.callback(
    Output("live-sync-state-store", "data"),
    Input("live-sync-interval", "n_intervals"),
    State("live-sync-state-store", "data"),
    prevent_initial_call=True,
)
def handle_live_sync_trigger(n_intervals, current_sync_state=None):
    # Lightweight check: background thread updates _last_loaded
    last_loaded = getattr(_data_store, "_last_loaded", 0.0)
    prev_loaded = (current_sync_state or {}).get("last_loaded", 0.0)

    if last_loaded > 0 and last_loaded != prev_loaded:
        clear_dmb_cache()
        return {"last_loaded": last_loaded, "ts": time.time(), "auto": True}

    return no_update


@app.callback(
    Output("mpr-month-filter", "options"),
    Output("dmb-month-filter", "options"),
    Input("live-sync-state-store", "data"),
    prevent_initial_call=True,
)
def refresh_month_dropdown_options(sync_data):
    active_mpr = get_active_mpr_data()
    active_dmb = get_active_dmb_data()
    active_strat = get_active_rca_actions()

    curr_mpr_months, _ = get_dynamic_reporting_months(active_mpr, active_strat)
    curr_dmb_months, _ = get_dynamic_reporting_months(active_dmb, active_strat)

    mpr_opts = create_month_options(curr_mpr_months)
    return mpr_opts, create_month_options(curr_dmb_months)


@app.callback(
    Output("mpr-month-filter", "value"),
    Output("dmb-month-filter", "value"),
    Input("mpr-month-filter", "value"),
    Input("dmb-month-filter", "value"),
    prevent_initial_call=True,
)
def sync_all_month_filters(mpr_val, dmb_val):
    try:
        ctx = dash.callback_context
        if ctx and ctx.triggered:
            triggered_id = ctx.triggered[0]["prop_id"].split(".")[0]
            if triggered_id == "mpr-month-filter" and mpr_val:
                return no_update, mpr_val
            elif triggered_id == "dmb-month-filter" and dmb_val:
                return dmb_val, no_update
    except Exception:
        pass
    raise PreventUpdate


# =========================================================
# MPR CALLBACK (Executive MoS & Imperative Performance)
# =========================================================

@app.callback(
    Output("dashboard-header-title", "children"),
    Output("mpr-total-kpis", "children"),
    Output("mpr-met-kpis", "children"),
    Output("mpr-not-met-kpis", "children"),
    Output("mpr-improved-kpis", "children"),
    Output("mpr-neither-kpis", "children"),
    Output("mpr-gauge", "figure"),
    Output("mpr-imperative-chart", "figure"),
    Output("imperative-month", "children"),
    Output("mpr-trend-chart", "figure"),
    Output("insights-reporting-month", "children"),
    Output("highlights-content", "children"),
    Output("highlights-count", "children"),
    Output("lowlights-content", "children"),
    Output("lowlights-count", "children"),
    Output("concerns-content", "children"),
    Output("concerns-count", "children"),
    Input("mpr-month-filter", "value"),
    Input("dmb-month-filter", "value"),
    Input("live-sync-state-store", "data"),
    prevent_initial_call=True,
)
def update_mpr_dashboard(mpr_val, dmb_val, sync_data):
    month_value = mpr_val or dmb_val
    try:
        ctx = dash.callback_context
        if ctx and ctx.triggered:
            trig = ctx.triggered[0]["prop_id"].split(".")[0]
            if trig == "mpr-month-filter" and mpr_val:
                month_value = mpr_val
            elif trig == "dmb-month-filter" and dmb_val:
                month_value = dmb_val
    except Exception:
        pass
    selected_month = pd.Timestamp(month_value)
    heading_title = f"{selected_month.strftime('%B')} MoS DMB Performance Review - AOP Critical"
    mpr_content = get_mpr_dashboard_content(month_value)
    return (heading_title, *mpr_content)


# =========================================================
# DMB FUNCTION-CARD CALLBACK (Function-wise Performance)
# =========================================================

@app.callback(
    Output("function-cards-container", "children"),
    Input("dmb-month-filter", "value"),
    Input("mpr-month-filter", "value"),
    Input("live-sync-state-store", "data"),
    prevent_initial_call=True,
)
def update_dmb_function_cards(dmb_val, mpr_val, sync_data):
    month_value = dmb_val or mpr_val
    try:
        ctx = dash.callback_context
        if ctx and ctx.triggered:
            trig = ctx.triggered[0]["prop_id"].split(".")[0]
            if trig == "dmb-month-filter" and dmb_val:
                month_value = dmb_val
            elif trig == "mpr-month-filter" and mpr_val:
                month_value = mpr_val
    except Exception:
        pass
    return get_dmb_function_cards_content(month_value)


# =========================================================
# RCA CALLBACK (MoS Level Cause & Actions)
# =========================================================

@app.callback(
    Output("rca-table-container", "children"),
    Output("rca-reporting-month", "children"),
    Input("mpr-month-filter", "value"),
    Input("dmb-month-filter", "value"),
    Input("live-sync-state-store", "data"),
    prevent_initial_call=True,
)
def update_rca_table(mpr_val, dmb_val, sync_data):
    month_value = mpr_val or dmb_val
    try:
        ctx = dash.callback_context
        if ctx and ctx.triggered:
            trig = ctx.triggered[0]["prop_id"].split(".")[0]
            if trig == "mpr-month-filter" and mpr_val:
                month_value = mpr_val
            elif trig == "dmb-month-filter" and dmb_val:
                month_value = dmb_val
    except Exception:
        pass
    return get_rca_table_content(month_value)


app.clientside_callback(
    """
    function(nClicksList) {
        if (!nClicksList || nClicksList.length === 0) {
            return window.dash_clientside.no_update;
        }
        var triggered = window.dash_clientside.callback_context.triggered;
        if (!triggered || triggered.length === 0) {
            return window.dash_clientside.no_update;
        }
        var trig = triggered[0];
        if (!trig || !trig.value) {
            return window.dash_clientside.no_update;
        }
        try {
            var propId = trig.prop_id;
            var jsonStr = propId.replace(/\\.n_clicks$/, '');
            var parsed = JSON.parse(jsonStr);
            var targetId = parsed.target;
            if (targetId) {
                var el = document.getElementById(targetId);
                if (el) {
                    el.scrollIntoView({ behavior: 'smooth', block: 'start' });
                    el.classList.add('kpi-card-highlighted');
                    setTimeout(function() {
                        el.classList.remove('kpi-card-highlighted');
                    }, 2500);
                }
            }
        } catch(err) {
            console.error("Scroll error:", err);
        }
        return window.dash_clientside.no_update;
    }
    """,
    Output("continuous-red-modal-scroll-dummy", "children"),
    Input({"type": "modal-rca-kpi-pill", "target": ALL}, "n_clicks"),
    prevent_initial_call=True,
)


# =========================================================
# CONTINUOUS-RED DETAIL MODAL CALLBACK
# =========================================================

@app.callback(
    Output("continuous-red-modal", "className"),
    Output("continuous-red-modal-title", "children"),
    Output("continuous-red-modal-month", "children"),
    Output("continuous-red-modal-body", "children"),
    Output("active-rca-modal-state", "data"),
    Input({"type": "continuous-red-card", "function": ALL}, "n_clicks"),
    Input("close-continuous-red-modal", "n_clicks"),
    Input("continuous-red-modal-backdrop", "n_clicks"),
    Input("dmb-month-filter", "value"),
    Input("live-sync-state-store", "data"),
    State("active-rca-modal-state", "data"),
    prevent_initial_call=True,
)
def manage_continuous_red_modal(
    card_clicks,
    close_clicks,
    backdrop_clicks,
    month_filter_val,
    sync_data,
    current_modal_state,
):
    ctx_inst = dash.callback_context
    current_modal_state = current_modal_state or {
        "is_open": False,
        "function": None,
        "month": None,
    }
    selected_month = (
        month_filter_val
        or current_modal_state.get("month")
        or default_month.strftime("%Y-%m-%d")
    )

    if not ctx_inst or not ctx_inst.triggered:
        raise PreventUpdate

    trig_id = getattr(dash.ctx, "triggered_id", None)

    # 1. Close Button or Backdrop Clicked -> Hide modal cleanly
    if trig_id in ("close-continuous-red-modal", "continuous-red-modal-backdrop"):
        close_has_click = any(
            item.get("value")
            for item in ctx_inst.triggered
            if item.get("prop_id", "").startswith(str(trig_id))
        )
        if close_has_click:
            return (
                "continuous-red-modal continuous-red-modal-hidden",
                no_update,
                no_update,
                no_update,
                {
                    "is_open": False,
                    "function": None,
                    "month": selected_month,
                },
            )
        raise PreventUpdate

    # 2. RCA Card Clicked from Dashboard (Must be an actual user click with n_clicks > 0)
    clicked_fn = None
    has_card_clicks = any(c for c in (card_clicks or []) if c and c > 0)

    if has_card_clicks:
        for item in (ctx_inst.triggered or []):
            prop_id = item.get("prop_id", "")
            val = item.get("value")
            if "continuous-red-card" in prop_id and val is not None and val > 0:
                comp_id_str = prop_id.rsplit(".", 1)[0]
                try:
                    parsed = json.loads(comp_id_str) if comp_id_str.startswith("{") else None
                    if isinstance(parsed, dict) and parsed.get("type") == "continuous-red-card":
                        clicked_fn = parsed.get("function")
                        break
                except Exception:
                    pass

        if not clicked_fn:
            parsed_trig = trig_id
            if isinstance(trig_id, str) and trig_id.startswith("{"):
                try:
                    parsed_trig = json.loads(trig_id)
                except Exception:
                    parsed_trig = None
            if isinstance(parsed_trig, dict) and parsed_trig.get("type") == "continuous-red-card":
                clicked_fn = parsed_trig.get("function")

    if clicked_fn:
        _, c_title, c_month, c_body = get_continuous_red_modal_content(
            clicked_fn, selected_month
        )
        return (
            "continuous-red-modal",
            c_title,
            c_month,
            c_body,
            {
                "is_open": True,
                "function": clicked_fn,
                "month": selected_month,
            },
        )

    # 3. Live Sync / Month Change (Update content ONLY if modal is currently open and trigger was filter/sync)
    if current_modal_state.get("is_open") and current_modal_state.get("function"):
        if trig_id in ("dmb-month-filter", "live-sync-state-store"):
            current_fn = current_modal_state["function"]
            _, c_title, c_month, c_body = get_continuous_red_modal_content(
                current_fn, selected_month
            )
            return (
                "continuous-red-modal",
                c_title,
                c_month,
                c_body,
                {
                    "is_open": True,
                    "function": current_fn,
                    "month": selected_month,
                },
            )

    # 4. Modal is closed or no valid action -> Do nothing
    raise PreventUpdate


# =========================================================
# FUNCTION GAUGE KPI DEFINITIONS & TREND MODAL CALLBACK
# =========================================================

@app.callback(
    Output("function-gauge-modal", "className"),
    Output("function-gauge-modal-title", "children"),
    Output("function-gauge-modal-subtitle", "children"),
    Output("function-gauge-modal-body", "children"),
    Output("active-gauge-modal-state", "data"),
    Input({"type": "function-gauge-card", "function": ALL}, "n_clicks"),
    Input("close-function-gauge-modal", "n_clicks"),
    Input("function-gauge-modal-backdrop", "n_clicks"),
    Input("dmb-month-filter", "value"),
    Input("live-sync-state-store", "data"),
    State("active-gauge-modal-state", "data"),
    prevent_initial_call=True,
)
def manage_function_gauge_modal(
    gauge_card_clicks,
    close_clicks,
    backdrop_clicks,
    month_filter_val,
    sync_data,
    current_modal_state,
):
    ctx_inst = dash.callback_context
    current_modal_state = current_modal_state or {
        "is_open": False,
        "function": None,
        "month": None,
    }
    selected_month = (
        month_filter_val
        or current_modal_state.get("month")
        or default_month.strftime("%Y-%m-%d")
    )

    if not ctx_inst or not ctx_inst.triggered:
        raise PreventUpdate

    trig_id = getattr(dash.ctx, "triggered_id", None)

    # 1. Close Button or Backdrop Clicked -> Hide modal cleanly
    if trig_id in ("close-function-gauge-modal", "function-gauge-modal-backdrop"):
        close_has_click = any(
            item.get("value")
            for item in ctx_inst.triggered
            if item.get("prop_id", "").startswith(str(trig_id))
        )
        if close_has_click:
            return (
                "function-gauge-modal function-gauge-modal-hidden",
                no_update,
                no_update,
                no_update,
                {
                    "is_open": False,
                    "function": None,
                    "month": selected_month,
                },
            )
        raise PreventUpdate

    # 2. Gauge Card Clicked from Dashboard (Must be an actual user click with n_clicks > 0)
    clicked_gauge_fn = None
    has_gauge_clicks = any(c for c in (gauge_card_clicks or []) if c and c > 0)

    if has_gauge_clicks:
        for item in (ctx_inst.triggered or []):
            prop_id = item.get("prop_id", "")
            val = item.get("value")
            if "function-gauge-card" in prop_id and val is not None and val > 0:
                comp_id_str = prop_id.rsplit(".", 1)[0]
                try:
                    parsed = json.loads(comp_id_str) if comp_id_str.startswith("{") else None
                    if isinstance(parsed, dict) and parsed.get("type") == "function-gauge-card":
                        clicked_gauge_fn = parsed.get("function")
                        break
                except Exception:
                    pass

        if not clicked_gauge_fn:
            parsed_trig = trig_id
            if isinstance(trig_id, str) and trig_id.startswith("{"):
                try:
                    parsed_trig = json.loads(trig_id)
                except Exception:
                    parsed_trig = None
            if isinstance(parsed_trig, dict) and parsed_trig.get("type") == "function-gauge-card":
                clicked_gauge_fn = parsed_trig.get("function")

    if clicked_gauge_fn:
        detail_content = get_function_kpis_trend_modal_content(
            clicked_gauge_fn, selected_month
        )
        sel_dt = pd.Timestamp(selected_month)
        return (
            "function-gauge-modal",
            f"{clicked_gauge_fn} — Function KPI Performance Review",
            f"All Key Performance Indicators · {sel_dt.strftime('%B %Y')}",
            detail_content,
            {
                "is_open": True,
                "function": clicked_gauge_fn,
                "month": selected_month,
            },
        )

    # 3. Live Sync or Month Filter Change (Update content ONLY if modal currently open and trigger was filter/sync)
    if current_modal_state.get("is_open") and current_modal_state.get("function"):
        if trig_id in ("dmb-month-filter", "live-sync-state-store"):
            func_name = current_modal_state["function"]
            detail_content = get_function_kpis_trend_modal_content(
                func_name, selected_month
            )
            sel_dt = pd.Timestamp(selected_month)
            return (
                "function-gauge-modal",
                f"{func_name} — Function KPI Performance Review",
                f"All Key Performance Indicators · {sel_dt.strftime('%B %Y')}",
                detail_content,
                {
                    "is_open": True,
                    "function": func_name,
                    "month": selected_month,
                },
            )

    # 4. Modal is closed or no valid action -> Do nothing
    raise PreventUpdate

# =========================================================
# MPR BREAKDOWN MODAL CALLBACK
# =========================================================

@app.callback(
    Output("mpr-breakdown-modal", "className"),
    Output("mpr-breakdown-modal-title", "children"),
    Output("mpr-breakdown-modal-subtitle", "children"),
    Output("mpr-breakdown-modal-pills-container", "children"),
    Output("mpr-breakdown-modal-body", "children"),
    Output("active-mpr-breakdown-modal-state", "data"),
    Input("card-mpr-total-kpis", "n_clicks"),
    Input("card-mpr-met-kpis", "n_clicks"),
    Input("card-mpr-improved-kpis", "n_clicks"),
    Input("card-mpr-neither-kpis", "n_clicks"),
    Input({"type": "mpr-modal-pill", "filter": ALL}, "n_clicks"),
    Input("close-mpr-breakdown-modal", "n_clicks"),
    Input("mpr-breakdown-modal-backdrop", "n_clicks"),
    Input("mpr-month-filter", "value"),
    State("active-mpr-breakdown-modal-state", "data"),
    prevent_initial_call=True,
)
def manage_mpr_breakdown_modal(
    total_clicks,
    met_clicks,
    improved_clicks,
    neither_clicks,
    pill_clicks,
    close_clicks,
    backdrop_clicks,
    month_val,
    current_state,
):
    ctx_inst = dash.callback_context
    if not ctx_inst or not ctx_inst.triggered:
        raise PreventUpdate

    current_state = current_state or {"is_open": False, "filter": "all", "month": None}
    selected_month = month_val or current_state.get("month") or default_month.strftime("%Y-%m-%d")
    trig_id = getattr(dash.ctx, "triggered_id", None)

    # 1. Close Button or Backdrop Clicked
    if trig_id in ("close-mpr-breakdown-modal", "mpr-breakdown-modal-backdrop"):
        return (
            "mpr-breakdown-modal mpr-breakdown-modal-hidden",
            no_update,
            no_update,
            no_update,
            no_update,
            {"is_open": False, "filter": current_state.get("filter", "all"), "month": selected_month},
        )

    # 2. Pill button clicked inside modal
    if isinstance(trig_id, dict) and trig_id.get("type") == "mpr-modal-pill":
        active_f = trig_id.get("filter", "all")
        title, subtitle, pills, body = get_mpr_kpis_breakdown_modal_content(selected_month, active_f)
        return (
            "mpr-breakdown-modal",
            title,
            subtitle,
            pills,
            body,
            {"is_open": True, "filter": active_f, "month": selected_month},
        )

    # 3. Card clicked on dashboard
    target_filter = None
    if trig_id == "card-mpr-neither-kpis" and neither_clicks:
        target_filter = "neither"
    elif trig_id == "card-mpr-improved-kpis" and improved_clicks:
        target_filter = "improved"
    elif trig_id == "card-mpr-met-kpis" and met_clicks:
        target_filter = "met"
    elif trig_id == "card-mpr-total-kpis" and total_clicks:
        target_filter = "all"

    # 4. Month filter changed while modal is open
    if trig_id == "mpr-month-filter":
        if current_state.get("is_open"):
            target_filter = current_state.get("filter", "all")
        else:
            raise PreventUpdate

    if target_filter is not None:
        title, subtitle, pills, body = get_mpr_kpis_breakdown_modal_content(selected_month, target_filter)
        return (
            "mpr-breakdown-modal",
            title,
            subtitle,
            pills,
            body,
            {"is_open": True, "filter": target_filter, "month": selected_month},
        )

    raise PreventUpdate


# =========================================================
# ONE-PAGER PNG EXPORT
# =========================================================

app.clientside_callback(
    """
    function (nClicks) {
        if (!nClicks) {
            return window.dash_clientside.no_update;
        }

        if (typeof window.html2canvas !== "function") {
            window.alert(
                "The one-pager export library did not load. " +
                "Please check your internet connection, refresh the page, and try again."
            );
            return nClicks;
        }

        const dashboard = document.querySelector(".page-shell");
        if (!dashboard) {
            window.alert("The dashboard could not be found for export.");
            return nClicks;
        }

        const root = document.documentElement;
        root.classList.add("one-pager-capturing");

        const fontsReady = document.fonts && document.fonts.ready
            ? document.fonts.ready
            : Promise.resolve();

        return fontsReady
            .then(function () {
                return new Promise(function (resolve) {
                    window.setTimeout(resolve, 250);
                });
            })
            .then(function () {
                const captureWidth = Math.max(
                    dashboard.scrollWidth,
                    dashboard.offsetWidth
                );
                const captureHeight = Math.max(
                    dashboard.scrollHeight,
                    dashboard.offsetHeight
                );
                const longestSide = Math.max(captureWidth, captureHeight);
                const exportScale = Math.max(
                    0.5,
                    Math.min(2, 16000 / longestSide)
                );

                return window.html2canvas(dashboard, {
                    backgroundColor: "#edf4f8",
                    scale: exportScale,
                    useCORS: true,
                    allowTaint: false,
                    logging: false,
                    scrollX: -window.scrollX,
                    scrollY: -window.scrollY,
                    windowWidth: Math.max(
                        document.documentElement.clientWidth,
                        captureWidth
                    ),
                    windowHeight: Math.max(
                        document.documentElement.clientHeight,
                        captureHeight
                    ),
                    onclone: function (clonedDocument) {
                        const clonedDashboard = clonedDocument.querySelector(
                            ".page-shell"
                        );

                        if (clonedDashboard) {
                            clonedDashboard.style.width = captureWidth + "px";
                            clonedDashboard.style.maxWidth = "none";
                            clonedDashboard.style.height = "auto";
                            clonedDashboard.style.maxHeight = "none";
                            clonedDashboard.style.overflow = "visible";
                        }

                        clonedDocument.documentElement.style.overflow = "visible";
                        clonedDocument.body.style.overflow = "visible";

                        clonedDocument.querySelectorAll(
                            ".download-button, " +
                            "._dash-debug-menu, " +
                            ".dash-debug-menu, " +
                            ".continuous-red-modal"
                        ).forEach(function (element) {
                            element.style.display = "none";
                        });
                    }
                });
            })
            .then(function (canvas) {
                return new Promise(function (resolve, reject) {
                    canvas.toBlob(function (blob) {
                        if (!blob) {
                            reject(new Error("The dashboard image could not be created."));
                            return;
                        }

                        const objectUrl = URL.createObjectURL(blob);
                        const downloadLink = document.createElement("a");
                        const exportDate = new Date().toISOString().slice(0, 10);

                        downloadLink.href = objectUrl;
                        downloadLink.download =
                            "DMB_One_Pager_" + exportDate + ".png";
                        document.body.appendChild(downloadLink);
                        downloadLink.click();
                        downloadLink.remove();

                        window.setTimeout(function () {
                            URL.revokeObjectURL(objectUrl);
                        }, 1000);

                        resolve(nClicks);
                    }, "image/png");
                });
            })
            .catch(function (error) {
                console.error("One-pager export failed:", error);
                window.alert(
                    "The one-pager image could not be downloaded. " +
                    "Please refresh the page and try again."
                );
                return nClicks;
            })
            .finally(function () {
                root.classList.remove("one-pager-capturing");
            });
    }
    """,
    Output("one-pager-download-state", "data"),
    Input("download-one-pager-button", "n_clicks"),
    prevent_initial_call=True,
)
# Starts the 24/7 background SharePoint Live Sync Engine
sharepoint_sync.start_live_sync(
    callback=lambda: data_loader.reload_all_data(force=True)
)

if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 10000)),
        debug=False,
    )

