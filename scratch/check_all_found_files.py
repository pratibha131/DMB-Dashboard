import openpyxl
from pathlib import Path

files = [
    Path(r"c:\Users\320320898\OneDrive - Philips\Functional DMB Review Sheets.xlsx"),
    Path(r"c:\Users\320320898\OneDrive - Philips\Desktop\DMB\Functional DMB Review Sheets.xlsx"),
    Path(r"c:\Users\320320898\OneDrive - Philips\Desktop\Trigger-workflow\Functional DMB Review Sheets.xlsx"),
    Path(r"c:\Users\320320898\OneDrive - Philips\Desktop\automation\Excel-automation\Functional DMB Review Sheets.xlsx"),
    Path(r"c:\Users\320320898\OneDrive - Philips\Desktop\DMB-Kavya\DMB-Dashboard\data\Functional DMB Review Sheets-16th Aug.xlsx"),
]

for p in files:
    if not p.exists():
        continue
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        print(f"\nChecking {p} (sheets: {wb.sheetnames})")
        if "7.Customer Service" in wb.sheetnames:
            ws = wb["7.Customer Service"]
            for r in range(45, min(80, ws.max_row + 1)):
                vals = [ws.cell(r, c).value for c in range(1, 8)]
                if any(v is not None for v in vals):
                    print(f"  Row {r:2d}: {[v for v in vals if v is not None]}")
        wb.close()
    except Exception as e:
        print(f"Error {p}: {e}")
