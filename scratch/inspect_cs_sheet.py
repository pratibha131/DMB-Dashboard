import openpyxl
from pathlib import Path

wb = openpyxl.load_workbook("data/Functional DMB Review Sheets-16th Aug.xlsx", data_only=True)
ws = wb["7.Customer Service"]

print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
for r in range(45, min(90, ws.max_row + 1)):
    row_vals = [ws.cell(r, c).value for c in range(1, 15)]
    if any(v is not None for v in row_vals):
        print(f"Row {r:2d}: {[v for v in row_vals[:8] if v is not None]}")
