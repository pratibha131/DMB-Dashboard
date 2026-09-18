import openpyxl

wb = openpyxl.load_workbook("data/Functional DMB Review Sheets-16th Aug.xlsx")
ws = wb["7.Customer Service"]

# Insert 3 rows before row 71
ws.insert_rows(71, 3)

# Row 71: Header
ws.cell(71, 1, "Red KPI")
ws.cell(71, 2, "Cause 1 ")
ws.cell(71, 3, "Cause 2")
ws.cell(71, 4, "Cause 3")
ws.cell(71, 5, "Cause 4")
ws.cell(71, 6, "Cause 5")

# Row 72: KPI Name and Causes
ws.cell(72, 1, "CP%")
ws.cell(72, 2, "META Contract Renewal GAP")
ws.cell(72, 3, "GRC CP drop")
ws.cell(72, 4, "EOL / EOS Impact")
ws.cell(72, 5, "Pending delayed renewal (Global)")

# Row 73: Impacts
ws.cell(73, 1, "% impact on KPI gap")
ws.cell(73, 2, 0.40)
ws.cell(73, 3, 0.30)
ws.cell(73, 4, 0.20)
ws.cell(73, 5, 0.10)

wb.save("data/Functional DMB Review Sheets-16th Aug.xlsx")
print("Saved CP% RCA block to sheet 7.Customer Service!")
