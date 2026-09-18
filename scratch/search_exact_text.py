import openpyxl

wb = openpyxl.load_workbook("data/Functional DMB Review Sheets-16th Aug.xlsx", data_only=True)
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or "").strip()
            if any(term in val.lower() for term in ["meta contract", "renewal gap", "grc cp", "contract renewals", "warranty conversion"]):
                print(f"Sheet: '{sname}' | Row {r}, Col {c} ({ws.cell(r, c).coordinate}): {val}")
