import openpyxl
from pathlib import Path

p1 = Path("data/Functional DMB Review Sheets-16th Aug.xlsx")
p2 = Path("data/Functional DMB Review Sheets-16thAug.xlsx")

print("File 1 exists:", p1.exists(), "size:", p1.stat().st_size if p1.exists() else 0)
print("File 2 exists:", p2.exists(), "size:", p2.stat().st_size if p2.exists() else 0)

if p1.exists():
    wb1 = openpyxl.load_workbook(p1, read_only=True)
    print("File 1 sheet names:", wb1.sheetnames)
    wb1.close()

if p2.exists():
    wb2 = openpyxl.load_workbook(p2, read_only=True)
    print("File 2 sheet names:", wb2.sheetnames)
    wb2.close()
