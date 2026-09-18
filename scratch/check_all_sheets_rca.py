import openpyxl

wb = openpyxl.load_workbook("data/Functional DMB Review Sheets-16th Aug.xlsx", data_only=True)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ---")
    for r in range(1, ws.max_row + 1):
        v = str(ws.cell(r, 1).value or "").strip()
        if "root cause" in v.lower() or "rca" in v.lower() or "action tracker" in v.lower() or "red kpi" in v.lower():
            row_sample = [ws.cell(r, c).value for c in range(1, 10)]
            print(f"  Row {r:3d}: {[x for x in row_sample if x is not None]}")
        # also check col 2
        v2 = str(ws.cell(r, 2).value or "").strip()
        if "cp" in v2.lower() or "cp%" in v2.lower() or "contract" in v2.lower():
            row_sample = [ws.cell(r, c).value for c in range(1, 10)]
            print(f"  Row {r:3d} (CP match): {[x for x in row_sample if x is not None]}")
