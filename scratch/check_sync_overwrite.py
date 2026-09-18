import openpyxl
from pathlib import Path

p_data = Path("data/Functional DMB Review Sheets-16th Aug.xlsx")
p_one = Path(r"c:\Users\320320898\OneDrive - Philips\Functional DMB Review Sheets.xlsx")

print("p_data exists:", p_data.exists(), "mtime:", p_data.stat().st_mtime if p_data.exists() else None)
print("p_one exists:", p_one.exists(), "mtime:", p_one.stat().st_mtime if p_one.exists() else None)

wb = openpyxl.load_workbook(p_data, data_only=True)
ws = wb["7.Customer Service"]
for r in range(48, min(80, ws.max_row + 1)):
    v = [ws.cell(r, c).value for c in range(1, 8)]
    if any(x is not None for x in v):
        print(f"p_data row {r}: {[x for x in v if x is not None]}")
wb.close()
