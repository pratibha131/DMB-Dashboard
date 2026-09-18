import openpyxl
from pathlib import Path

for p in Path(".").rglob("*.xlsx"):
    if p.name.startswith("~$"):
        continue
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for r in range(1, min(150, ws.max_row + 1)):
                for c in range(1, min(40, ws.max_column + 1)):
                    v = str(ws.cell(r, c).value or "")
                    if "META Contract Renewal GAP" in v or "GRC CP drop" in v or "META Impact (War)" in v:
                        print(f"FOUND in file '{p}', sheet '{sname}', row {r}, col {c}: {v}")
        wb.close()
    except Exception as e:
        print(f"Error reading {p}: {e}")
