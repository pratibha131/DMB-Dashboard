from pathlib import Path
import re
import datetime
import io
import pandas as pd
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "data" / "Masterfile_DMB_Dashboard.xlsx"
STRATEGIC_EXCEL_PATH = BASE_DIR / "data" / "Strategic Execution Dashboard-17Th_sept.xlsx"


import os
import ctypes


def read_file_safe_bytes(file_path: Path | str) -> bytes:
    """
    Read file safely on Windows even if open in Excel / OneDrive sync.
    Uses Win32 CreateFile with FILE_SHARE_READ | FILE_SHARE_WRITE | FILE_SHARE_DELETE.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if os.name != "nt":
        return file_path.read_bytes()

    import msvcrt
    from ctypes import wintypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    create_file.restype = wintypes.HANDLE

    generic_read = 0x80000000
    share_all = 0x7  # read, write, delete
    open_existing = 3
    normal_attributes = 0x80

    handle = create_file(
        str(file_path),
        generic_read,
        share_all,
        None,
        open_existing,
        normal_attributes,
        None,
    )
    if handle is None or handle == wintypes.HANDLE(-1).value or handle == -1:
        return file_path.read_bytes()

    try:
        fd = msvcrt.open_osfhandle(handle, os.O_RDONLY)
        with open(fd, "rb") as f:
            return f.read()
    except Exception:
        return file_path.read_bytes()



def clean_text(value):
    if pd.isna(value):
        return ""
    return str(value).replace("\xa0", " ").strip()


def first_non_blank(series):
    for value in series:
        if pd.notna(value) and clean_text(value):
            return value
    return ""


def combine_target_actual(records, keys, metadata_columns):
    long_df = pd.DataFrame(records)

    if long_df.empty:
        return pd.DataFrame()

    values = (
        long_df.pivot_table(
            index=keys,
            columns="series",
            values="value",
            aggfunc="first",
        )
        .reset_index()
        .rename_axis(None, axis=1)
    )

    metadata = (
        long_df.groupby(keys, as_index=False)[metadata_columns]
        .agg(first_non_blank)
    )

    result = metadata.merge(values, on=keys, how="outer")

    for column in ["Target", "Actual"]:
        if column not in result.columns:
            result[column] = pd.NA

    return result.sort_values(keys).reset_index(drop=True)


def load_mpr_data_from_strategic(excel_path=None):
    if excel_path is None:
        candidates = sorted(
            file_path
            for file_path in (BASE_DIR / "data").glob("*Strategic*.xlsx")
            if not file_path.name.startswith("~$")
        )
        if candidates:
            excel_path = candidates[0]
        else:
            excel_path = STRATEGIC_EXCEL_PATH

    if not Path(excel_path).exists():
        return pd.DataFrame()

    try:
        file_bytes = read_file_safe_bytes(excel_path)
        xl = pd.ExcelFile(io.BytesIO(file_bytes), engine="openpyxl")
        sheet_name = "AOP Critical" if "AOP Critical" in xl.sheet_names else xl.sheet_names[0]
        raw = xl.parse(sheet_name, header=None)
    except Exception as exc:
        print(f"[Strategic Loader Notice] {exc}")
        return pd.DataFrame()

    header_row_idx = 4
    for r in range(min(12, len(raw))):
        row_vals = [str(val).strip().lower() for val in raw.iloc[r] if pd.notna(val)]
        if "jan" in row_vals or "feb" in row_vals:
            header_row_idx = r
            break

    header_row = raw.iloc[header_row_idx]
    month_prefixes = {
        "jan": 1, "feb": 2, "mar": 3, "q1": 3, "apr": 4, "may": 5, "jun": 6, "q2": 6,
        "jul": 7, "aug": 8, "sep": 9, "q3": 9, "oct": 10, "nov": 11, "dec": 12, "q4": 12,
    }
    month_cols = {}
    for c in range(len(header_row)):
        val = str(header_row[c]).strip().lower()
        for prefix, m_idx in month_prefixes.items():
            if val.startswith(prefix):
                month_cols[c] = pd.Timestamp(year=2026, month=m_idx, day=1)
                break

    current_strat = ""
    records = []

    for r in range(header_row_idx + 1, len(raw)):
        strat_val = raw.iloc[r, 0] if raw.shape[1] > 0 else None
        if pd.notna(strat_val) and str(strat_val).strip() and not str(strat_val).startswith("MoS") and not str(strat_val).startswith("Strategic"):
            raw_s = str(strat_val).strip()
            if "cusotomer" in raw_s.lower() or "customer" in raw_s.lower():
                current_strat = "Customer Focus"
            elif "deliverability" in raw_s.lower() or "profitability" in raw_s.lower():
                current_strat = "Improve Deliverability & Profitability"
            elif "safety" in raw_s.lower() or "quality" in raw_s.lower():
                current_strat = "Patient Safety & Quality"
            elif "commercial" in raw_s.lower() or "growth" in raw_s.lower():
                current_strat = "Drive growth through Commercial Excellence"
            elif "roadmap" in raw_s.lower() or "innovation" in raw_s.lower():
                current_strat = "Roadmap competitiveness & Innovation agility"
            elif "esg" in raw_s.lower() or "environmental" in raw_s.lower() or "governance" in raw_s.lower():
                current_strat = "Environmental, Social & Governance (ESG)"
            else:
                current_strat = raw_s

        t_or_a = str(raw.iloc[r, 10] if raw.shape[1] > 10 else "").strip().upper()
        if t_or_a == "T":
            kpi_val = raw.iloc[r, 1] if raw.shape[1] > 1 else ""
            kpi_name = clean_text(kpi_val)
            if not kpi_name or kpi_name.lower().startswith("value lever") or kpi_name.lower().startswith("kpi "):
                continue

            owner = clean_text(raw.iloc[r, 2] if raw.shape[1] > 2 else "")
            kpi_def = clean_text(raw.iloc[r, 3] if raw.shape[1] > 3 else "")
            aop_2026 = pd.to_numeric(raw.iloc[r, 4] if raw.shape[1] > 4 else None, errors="coerce")
            nature = str(raw.iloc[r, 7] if raw.shape[1] > 7 else "").strip()
            units = str(raw.iloc[r, 8] if raw.shape[1] > 8 else "").strip()
            freq = str(raw.iloc[r, 9] if raw.shape[1] > 9 else "").strip()

            if pd.notna(aop_2026) and units == "%" and aop_2026 > 1:
                aop_2026 = aop_2026 / 100.0

            if "(EQ)" in kpi_def and "(EQ)" not in kpi_name:
                kpi_name = f"{kpi_name} (EQ)"
            elif "(CS)" in kpi_def and "(CS)" not in kpi_name:
                kpi_name = f"{kpi_name} (CS)"

            a_row = None
            if r + 1 < len(raw) and str(raw.iloc[r + 1, 10] if raw.shape[1] > 10 else "").strip().upper() == "A":
                a_row = raw.iloc[r + 1]

            for c_idx, m_ts in month_cols.items():
                t_val = raw.iloc[r, c_idx]
                a_val = a_row[c_idx] if a_row is not None else None

                t_num = pd.to_numeric(t_val, errors="coerce")
                if pd.isna(t_num) and pd.notna(aop_2026):
                    t_num = aop_2026
                a_num = pd.to_numeric(a_val, errors="coerce")

                records.append({
                    "strategic_imperative": current_strat,
                    "kpi_name": kpi_name,
                    "month": m_ts,
                    "metric_owner": owner,
                    "definition": kpi_def,
                    "aop_2026": aop_2026,
                    "metric_nature": nature,
                    "units": units,
                    "frequency": freq,
                    "data_type": "Monthly",
                    "series": "Target",
                    "value": t_num,
                })
                records.append({
                    "strategic_imperative": current_strat,
                    "kpi_name": kpi_name,
                    "month": m_ts,
                    "metric_owner": owner,
                    "definition": kpi_def,
                    "aop_2026": aop_2026,
                    "metric_nature": nature,
                    "units": units,
                    "frequency": freq,
                    "data_type": "Monthly",
                    "series": "Actual",
                    "value": a_num,
                })

    return combine_target_actual(
        records,
        keys=["strategic_imperative", "kpi_name", "month"],
        metadata_columns=[
            "metric_owner",
            "definition",
            "aop_2026",
            "metric_nature",
            "units",
            "frequency",
            "data_type",
        ],
    )


def load_mpr_data(excel_path=None):
    if excel_path is None:
        excel_path = find_masterfile()

    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    file_bytes = read_file_safe_bytes(excel_path)
    raw = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name="MPR Masterfile",
        header=None,
        engine="openpyxl",
    )

    records = []
    current_imperative = ""
    month_headers = {}
    reporting_year = 2026

    for row_number in range(len(raw)):
        first_cell = clean_text(raw.iat[row_number, 0])

        next_first_cell = (
            clean_text(raw.iat[row_number + 1, 0])
            if row_number + 1 < len(raw)
            else ""
        )

        normalized_next = next_first_cell.lower().replace(" ", "")

        if normalized_next in {"corekpi's", "corekpis"}:
            current_imperative = first_cell
            continue

        normalized_first = first_cell.lower().replace(" ", "")

        if normalized_first in {"corekpi's", "corekpis"}:
            year_match = re.search(
                r"20\d{2}",
                clean_text(raw.iat[row_number, 3]),
            )

            if year_match:
                reporting_year = int(year_match.group())

            month_headers = {
                column: clean_text(raw.iat[row_number, column])
                for column in range(10, min(22, raw.shape[1]))
            }
            continue

        row_type = clean_text(raw.iat[row_number, 9]).upper()

        if row_type not in {"T", "A", "TARGET", "ACTUAL"}:
            continue

        series_name = (
            "Target"
            if row_type in {"T", "TARGET"}
            else "Actual"
        )

        for column, month_name in month_headers.items():
            value = raw.iat[row_number, column]

            if not month_name or pd.isna(value):
                continue

            month_date = pd.to_datetime(
                f"{month_name}-{reporting_year}",
                format="%b-%Y",
                errors="coerce",
            )

            if pd.isna(month_date):
                continue

            records.append(
                {
                    "strategic_imperative": current_imperative,
                    "kpi_name": first_cell,
                    "month": month_date,
                    "metric_owner": clean_text(raw.iat[row_number, 1]),
                    "definition": clean_text(raw.iat[row_number, 2]),
                    "aop_2026": raw.iat[row_number, 3],
                    "metric_nature": clean_text(raw.iat[row_number, 6]),
                    "units": clean_text(raw.iat[row_number, 7]),
                    "frequency": clean_text(raw.iat[row_number, 8]),
                    "data_type": clean_text(raw.iat[row_number, 22]),
                    "series": series_name,
                    "value": value,
                }
            )

    return combine_target_actual(
        records,
        keys=["strategic_imperative", "kpi_name", "month"],
        metadata_columns=[
            "metric_owner",
            "definition",
            "aop_2026",
            "metric_nature",
            "units",
            "frequency",
            "data_type",
        ],
    )


def load_dmb_data(excel_path=EXCEL_PATH):
    file_bytes = read_file_safe_bytes(excel_path)
    raw = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name="DMB Masterfile",
        header=None,
        engine="openpyxl",
    )

    records = []
    current_function = ""
    month_headers = {}

    for row_number in range(len(raw)):
        first_cell = clean_text(raw.iat[row_number, 0])

        next_first_cell = (
            clean_text(raw.iat[row_number + 1, 0])
            if row_number + 1 < len(raw)
            else ""
        )

        if next_first_cell.lower() == "kpi name":
            current_function = first_cell.replace(" DMB", "").strip()
            continue

        if first_cell.lower() == "kpi name":
            month_headers = {
                column: clean_text(raw.iat[row_number, column])
                for column in range(8, min(20, raw.shape[1]))
            }
            continue

        row_type = clean_text(raw.iat[row_number, 7]).lower()

        if row_type not in {"target", "actual"}:
            continue

        series_name = row_type.title()

        for column, month_name in month_headers.items():
            value = raw.iat[row_number, column]

            if not month_name or pd.isna(value):
                continue

            month_date = pd.to_datetime(
                month_name,
                format="%b-%Y",
                errors="coerce",
            )

            if pd.isna(month_date):
                continue

            records.append(
                {
                    "function": current_function,
                    "kpi_name": first_cell,
                    "month": month_date,
                    "definition": clean_text(raw.iat[row_number, 1]),
                    "operator": clean_text(raw.iat[row_number, 2]),
                    "target_aop_2026": raw.iat[row_number, 3],
                    "units": clean_text(raw.iat[row_number, 4]),
                    "metric_nature": clean_text(raw.iat[row_number, 5]),
                    "frequency": clean_text(raw.iat[row_number, 6]),
                    "data_type": clean_text(raw.iat[row_number, 20]),
                    "kpi_category": clean_text(raw.iat[row_number, 21]),
                    "series": series_name,
                    "value": value,
                }
            )

    return combine_target_actual(
        records,
        keys=["function", "kpi_name", "month"],
        metadata_columns=[
            "definition",
            "operator",
            "target_aop_2026",
            "units",
            "metric_nature",
            "frequency",
            "data_type",
            "kpi_category",
        ],
    )


FUNCTIONAL_REVIEW_PATH = BASE_DIR / "data" / "Functional DMB Review Sheets-17th_sept.xlsx"


def load_red_kpis_rca(file_path=None):
    """
    Parses Functional DMB Review Sheets-17th_sept.xlsx (or configured file) for Red KPIs RCA
    and corrective action tracker entries.
    """
    if file_path is None:
        file_path = FUNCTIONAL_REVIEW_PATH
    file_path = Path(file_path)
    if not file_path.exists():
        data_dir = file_path.parent
        matching = (
            list(data_dir.glob("*17th*sept*.xlsx"))
            or list(data_dir.glob("*17Th*sept*.xlsx"))
            or list(data_dir.glob("*Functional*Review*.xlsx"))
            or list(data_dir.glob("*Review*.xlsx"))
            or list(data_dir.glob("*Strategic*Execution*.xlsx"))
        )
        if matching:
            file_path = matching[0]
        else:
            return []

    try:
        file_bytes = read_file_safe_bytes(file_path)
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        records = []

        status_map = {
            1: "Action not assigned",
            2: "Action assigned",
            3: "Action started",
            4: "Action completed",
            5: "Resolution confirmed",
        }

        for sheet_name in wb.sheetnames:
            if sheet_name in ["MPR Review", "MasterSheet", "Sheet1 (2)"]:
                continue
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            function_name = sheet_name.replace("1.", "").replace("2.", "").replace("3.", "").replace("4.", "").replace("5.", "").replace("6.", "").replace("7.", "").replace("8.", "").replace("9.", "").replace("10.", "").replace("11.", "").replace("12.", "").strip()

            in_rca = False
            in_act = False

            for r_idx, row in enumerate(rows):
                if not row:
                    continue
                c0 = clean_text(row[0]) if len(row) > 0 else ""
                c1 = clean_text(row[1]) if len(row) > 1 else ""

                if "root cause analysis" in c0.lower() or "root cause analysis" in c1.lower():
                    in_rca = True
                    in_act = False
                    continue

                if "action tracker" in c0.lower() or "action tracker" in c1.lower():
                    in_rca = False
                    in_act = True
                    continue

                if in_act:
                    if (
                        c0
                        and c0.lower() not in ["red kpi", "action tracker", "action not assigned"]
                        and "if the kpi is reported" not in c0.lower()
                    ):
                        rc_desc = clean_text(row[1]) if len(row) > 1 else ""
                        corr_act = clean_text(row[2]) if len(row) > 2 else ""
                        owner = clean_text(row[3]) if len(row) > 3 else ""
                        due_val = row[4] if len(row) > 4 else None
                        status_val = row[5] if len(row) > 5 else None

                        if rc_desc or corr_act:
                            status_text = ""
                            if isinstance(status_val, (int, float)) and int(status_val) in status_map:
                                status_text = status_map[int(status_val)]
                            elif status_val is not None:
                                status_text = clean_text(status_val)

                            due_text = ""
                            if isinstance(due_val, (pd.Timestamp, datetime.datetime, datetime.date)):
                                due_text = pd.Timestamp(due_val).strftime("%d %b %Y")
                            elif due_val is not None:
                                due_text = clean_text(due_val)

                            records.append({
                                "function": function_name,
                                "source_sheet": sheet_name,
                                "kpi_name": c0,
                                "root_cause": rc_desc or "Not entered",
                                "action": corr_act or "Not entered",
                                "owner": owner or "Not assigned",
                                "due_date": due_text or "Not entered",
                                "status": status_text or "Status not entered",
                            })

        return records
    except Exception as exc:
        print(f"[RCA Loader Notice] {exc}")
        return []


def find_strategic_workbook():
    if STRATEGIC_EXCEL_PATH.exists():
        return STRATEGIC_EXCEL_PATH

    # Otherwise use the most recently saved Strategic workbook, not the first
    # name alphabetically ("... (1).xlsx" sorts before "...-17Th_sept.xlsx").
    candidates = [
        file_path
        for file_path in (BASE_DIR / "data").glob("*Strategic*.xlsx")
        if not file_path.name.startswith("~$")
    ]
    if candidates:
        return max(candidates, key=lambda file_path: file_path.stat().st_mtime)
    return STRATEGIC_EXCEL_PATH


def find_performance_month(sheet):
    """Read the "Performance of the Month" date from the top of the sheet."""
    for row_index in range(min(4, len(sheet))):
        for column_index in range(sheet.shape[1] - 1):
            label = str(sheet.iat[row_index, column_index]).lower()
            if "performance of the month" not in label:
                continue
            for value in sheet.iloc[row_index, column_index + 1:]:
                if pd.isna(value):
                    continue
                month = pd.to_datetime(value, errors="coerce")
                if pd.notna(month):
                    return month.to_period("M").to_timestamp()
    return None


def load_strategic_rca_actions(excel_path=None):
    if excel_path is None:
        excel_path = find_strategic_workbook()

    if not Path(excel_path).exists():
        return pd.DataFrame()

    try:
        file_bytes = read_file_safe_bytes(excel_path)
        xl = pd.ExcelFile(io.BytesIO(file_bytes), engine="openpyxl")
        sheet_name = "AOP Critical" if "AOP Critical" in xl.sheet_names else xl.sheet_names[0]
        df = xl.parse(sheet_name, header=None)
    except Exception as exc:
        print(f"[Strategic Loader Notice] {exc}")
        return pd.DataFrame()

    header_row_idx = None
    for r in range(min(12, len(df))):
        row_vals = [str(val).strip().lower() for val in df.iloc[r] if pd.notna(val)]
        if "jan" in row_vals or "feb" in row_vals:
            header_row_idx = r
            break

    if header_row_idx is None:
        header_row_idx = 4

    header_row = df.iloc[header_row_idx]
    month_prefixes = {
        "jan": 1, "feb": 2, "mar": 3, "q1": 3, "apr": 4, "may": 5, "jun": 6, "q2": 6,
        "jul": 7, "aug": 8, "sep": 9, "q3": 9, "oct": 10, "nov": 11, "dec": 12, "q4": 12,
    }
    month_cols = {}
    for c in range(len(header_row)):
        val = str(header_row[c]).strip().lower()
        for prefix, m_idx in month_prefixes.items():
            if val.startswith(prefix):
                month_cols[pd.Timestamp(year=2026, month=m_idx, day=1)] = c
                break

    def is_target_row(row_index):
        return str(df.iloc[row_index, 10] if df.shape[1] > 10 else "").strip().upper() == "T"

    # A KPI cell left blank continues the KPI above it (for example the (CS)
    # row under "Order Intake Growth"). Only names shared by several rows get
    # an (EQ)/(CS) suffix, so every other KPI keeps its name from the sheet.
    kpi_name_counts = {}
    carried_kpi_name = ""
    for r in range(header_row_idx + 1, len(df)):
        if not is_target_row(r):
            continue
        kpi_val = df.iloc[r, 1] if df.shape[1] > 1 else ""
        if pd.notna(kpi_val) and str(kpi_val).strip():
            carried_kpi_name = clean_text(kpi_val)
        kpi_name_counts[carried_kpi_name] = kpi_name_counts.get(carried_kpi_name, 0) + 1

    current_strat = ""
    current_kpi_name = ""
    records = []

    for r in range(header_row_idx + 1, len(df)):
        strat_val = df.iloc[r, 0] if df.shape[1] > 0 else None
        if pd.notna(strat_val) and str(strat_val).strip() and not str(strat_val).startswith("MoS") and not str(strat_val).startswith("Strategic"):
            raw_s = str(strat_val).strip()
            if "cusotomer" in raw_s.lower() or "customer" in raw_s.lower():
                current_strat = "Customer Focus"
            elif "deliverability" in raw_s.lower() or "profitability" in raw_s.lower():
                current_strat = "Improve Deliverability & Profitability"
            elif "safety" in raw_s.lower() or "quality" in raw_s.lower():
                current_strat = "Patient Safety & Quality"
            elif "commercial" in raw_s.lower() or "growth" in raw_s.lower():
                current_strat = "Drive growth through Commercial Excellence"
            elif "roadmap" in raw_s.lower() or "innovation" in raw_s.lower():
                current_strat = "Roadmap competitiveness & Innovation agility"
            elif "esg" in raw_s.lower() or "environmental" in raw_s.lower() or "governance" in raw_s.lower():
                current_strat = "Environmental, Social & Governance"
            else:
                current_strat = raw_s

        if is_target_row(r):
            kpi_val = df.iloc[r, 1] if df.shape[1] > 1 else ""
            if pd.notna(kpi_val) and str(kpi_val).strip():
                current_kpi_name = clean_text(kpi_val)

            kpi_def = clean_text(df.iloc[r, 3] if df.shape[1] > 3 else "")

            kpi_name = current_kpi_name
            if not kpi_name or kpi_name.lower().startswith("value lever") or kpi_name.lower().startswith("kpi "):
                continue

            if kpi_name_counts.get(current_kpi_name, 0) > 1:
                if "(EQ)" in kpi_def and "(EQ)" not in kpi_name:
                    kpi_name = f"{kpi_name} (EQ)"
                elif "(CS)" in kpi_def and "(CS)" not in kpi_name:
                    kpi_name = f"{kpi_name} (CS)"

            nature = str(df.iloc[r, 7] if df.shape[1] > 7 else "").strip().lower()
            lower_is_better = "lower" in nature
            higher_is_better = not lower_is_better

            cause = clean_text(df.iloc[r, 23] if df.shape[1] > 23 else "")
            action = clean_text(df.iloc[r, 24] if df.shape[1] > 24 else "")

            a_row = None
            if r + 1 < len(df) and str(df.iloc[r + 1, 10] if df.shape[1] > 10 else "").strip().upper() == "A":
                a_row = df.iloc[r + 1]
                a_cause = clean_text(df.iloc[r + 1, 23] if df.shape[1] > 23 else "")
                a_action = clean_text(df.iloc[r + 1, 24] if df.shape[1] > 24 else "")
                if not cause and a_cause:
                    cause = a_cause
                if not action and a_action:
                    action = a_action

            if cause.lower() in {"nan", "none"}:
                cause = ""
            if action.lower() in {"nan", "none"}:
                action = ""

            aop_target = pd.to_numeric(df.iloc[r, 4] if df.shape[1] > 4 else None, errors="coerce")
            unit_str = str(df.iloc[r, 8] if df.shape[1] > 8 else "").strip()
            if pd.notna(aop_target) and unit_str == "%" and aop_target > 1:
                aop_target = aop_target / 100.0

            for m_ts, c_idx in month_cols.items():
                t_val = df.iloc[r, c_idx]
                a_val = a_row[c_idx] if a_row is not None else None

                t_num = pd.to_numeric(t_val, errors="coerce")
                if pd.isna(t_num) and pd.notna(aop_target):
                    t_num = aop_target

                a_num = pd.to_numeric(a_val, errors="coerce")

                is_red = False
                if pd.notna(t_num) and pd.notna(a_num):
                    if lower_is_better:
                        is_red = bool(a_num > t_num)
                    else:
                        is_red = bool(a_num < t_num)

                records.append({
                    "reporting_month": m_ts,
                    "strategic_imperative": current_strat,
                    "kpi_name": kpi_name,
                    "target": t_num,
                    "actual": a_num,
                    "is_red": is_red,
                    "cause": cause,
                    "action": action,
                })

    result = pd.DataFrame(records)
    if result.empty:
        return result

    return result


def find_masterfile(data_dir=None):
    if data_dir is None:
        data_dir = BASE_DIR / "data"
    data_dir = Path(data_dir)

    preferred_file = data_dir / "Masterfile_DMB_Dashboard.xlsx"

    candidates = sorted(
        file_path
        for file_path in data_dir.glob("*.xlsx")
        if not file_path.name.startswith("~$")
    )

    if preferred_file.exists():
        candidates = [preferred_file] + [
            file_path
            for file_path in candidates
            if file_path != preferred_file
        ]

    fallback_file = None

    for file_path in candidates:
        try:
            file_bytes = read_file_safe_bytes(file_path)
            sheet_names = set(pd.ExcelFile(
                io.BytesIO(file_bytes),
                engine="openpyxl",
            ).sheet_names)

            if {
                "MPR Masterfile",
                "DMB Masterfile",
            }.issubset(sheet_names):
                if "RCA Actions" in sheet_names:
                    return file_path

                if fallback_file is None:
                    fallback_file = file_path
        except (OSError, ValueError):
            continue

    return fallback_file or preferred_file


def load_rca_data(excel_path=None):
    """Loads Cause and Actions of Red KPIs at MoS Level."""
    return load_red_kpis_rca(excel_path)


def load_dashboard_data(excel_path=None):
    if excel_path is None:
        excel_path = find_masterfile()
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    mpr_data = load_mpr_data(excel_path)
    dmb_data = load_dmb_data(excel_path)

    return mpr_data, dmb_data


import threading
import time


def data_folder_signature(data_dir=None):
    if data_dir is None:
        data_dir = BASE_DIR / "data"
    data_dir = Path(data_dir)
    signature = []
    for file_path in sorted(data_dir.glob("*.xlsx")):
        if file_path.name.startswith("~$"):
            continue
        try:
            file_stat = file_path.stat()
        except OSError:
            continue
        signature.append(
            (file_path.name, file_stat.st_mtime_ns, file_stat.st_size)
        )
    return tuple(signature)


class DMBDataStore:
    """Thread-safe in-memory store for all loaded and prepared DMB data."""
    def __init__(self):
        self._lock = threading.Lock()
        self._signature = None
        self._mpr_raw = pd.DataFrame()
        self._dmb_raw = pd.DataFrame()
        self._mpr_data = pd.DataFrame()
        self._dmb_data = pd.DataFrame()
        self._rca_actions = pd.DataFrame()
        self._strat_rca = pd.DataFrame()
        self._last_loaded = 0.0

    def reload(self, force=False):
        current_sig = data_folder_signature()
        with self._lock:
            if not force and self._signature == current_sig and not self._mpr_data.empty:
                return False

            try:
                from kpi_calculations import prepare_kpi_data
                master_file = find_masterfile()
                if master_file and master_file.exists():
                    mpr_raw, dmb_raw = load_dashboard_data(master_file)
                    self._mpr_raw = mpr_raw
                    self._dmb_raw = dmb_raw
                    self._mpr_data = prepare_kpi_data(mpr_raw, ["strategic_imperative", "kpi_name"])
                    self._dmb_data = prepare_kpi_data(dmb_raw, ["function", "kpi_name"])

                self._strat_rca = load_strategic_rca_actions()
                self._signature = current_sig
                self._last_loaded = time.time()
                return True
            except Exception as e:
                print(f"[DMBDataStore Reload Notice] {e}")
                return False

    def get_data(self):
        """Ensure data is loaded and return tuple of (mpr_data, dmb_data, strat_rca)."""
        current_sig = data_folder_signature()
        if self._signature != current_sig or self._mpr_data.empty:
            self.reload()
        with self._lock:
            return self._mpr_data.copy(), self._dmb_data.copy(), self._strat_rca.copy()

    def get_mpr_data(self):
        return self.get_data()[0]

    def get_dmb_data(self):
        return self.get_data()[1]

    def get_strat_rca(self):
        return self.get_data()[2]


_global_store = DMBDataStore()


def get_data_store() -> DMBDataStore:
    return _global_store


def reload_all_data(force=True) -> bool:
    return _global_store.reload(force=force)


def load_data():
    return _global_store.reload(force=True)


if __name__ == "__main__":
    mpr_data, dmb_data = load_dashboard_data()
    rca_data = load_rca_data()
    strat_rca = load_strategic_rca_actions()

    print("Excel connection successful.")
    print(f"MPR KPIs: {mpr_data['kpi_name'].nunique()}")
    print(f"Strategic imperatives: {mpr_data['strategic_imperative'].nunique()}")
    print(f"DMB KPIs: {dmb_data[['function', 'kpi_name']].drop_duplicates().shape[0]}")
    print(f"Functions: {dmb_data['function'].nunique()}")
    print(f"Strategic Red KPI RCA Records: {len(strat_rca)}")



